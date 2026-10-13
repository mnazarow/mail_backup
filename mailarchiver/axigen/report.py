"""
Выгрузка учётных записей сервера Axigen (по последней сверке) вместе с ящиками
архива — в Excel или CSV: для сверки «кто есть на сервере и что из этого в
архиве», отчёта о квотах и давно не входивших.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any, List, Tuple

from .common import load_settings
from .directory import overview

COLUMNS: List[Tuple[str, str]] = [
    ("email", "Адрес"), ("domain", "Домен"), ("full_name", "ФИО"), ("aliases", "Псевдонимы"),
    ("state", "На сервере"), ("services", "Службы"), ("last_login", "Последний вход"),
    ("quota_mb", "Квота, МБ"), ("used_mb", "Занято, МБ"), ("quota_percent", "Заполнено, %"),
    ("messages", "Писем на сервере"), ("mailbox", "Ящик в архиве"), ("copy", "Как копируется"),
    ("last_backup_at", "Последняя копия"), ("archived", "Писем в архиве"), ("archived_mb", "Архив, МБ"),
    ("errors", "Не прочитано при сверке"),
]
_COPY = {"axigen_ftp": "FTP Backup", "password": "IMAP, пароль", "master": "IMAP, администратор",
         "oauth2": "OAuth2"}


def _row(item) -> dict:
    box = item.get("mailbox") or {}
    if item.get("deleted_at"):
        state = "удалена"
    elif item.get("login_blocked"):
        state = "вход закрыт"
    elif item.get("login_blocked") is False:
        state = "вход открыт"
    else:
        state = "неизвестно"
    quota = item.get("quota_kb")
    used = item.get("used_bytes")
    return {
        "email": item["email"], "domain": item.get("domain", ""),
        "full_name": " ".join(x for x in (item.get("first_name") or "", item.get("last_name") or "") if x),
        "aliases": ", ".join(item.get("aliases") or []), "state": state,
        "services": ", ".join(item.get("services") or []) if item.get("services") is not None else "",
        "last_login": (item.get("last_login") or "")[:19].replace("T", " "),
        "quota_mb": round(quota / 1024, 1) if quota else None,
        "used_mb": round(used / 1024 / 1024, 1) if used is not None else None,
        "quota_percent": item.get("quota_percent"),
        "messages": item.get("messages"),
        "mailbox": box.get("name", "") if box else "нет",
        "copy": (_COPY.get(box.get("auth_type"), box.get("auth_type", "")) + ("" if box.get("enabled") else
                                                                             ", выключен")) if box else "",
        "last_backup_at": (box.get("last_backup_at") or "")[:19].replace("T", " ") if box else "",
        "archived": box.get("messages") if box else None,
        "archived_mb": round(box["bytes"] / 1024 / 1024, 1) if box and box.get("bytes") is not None else None,
        "errors": item.get("errors") or "",
    }


def _safe(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def export_accounts(svc, fmt: str = "xlsx") -> Tuple[bytes, str, str]:
    data = overview(svc, load_settings(svc))
    rows = [_row(item) for item in data["accounts"]]
    stamp = datetime.now(svc.local_tz()).strftime("%Y%m%d_%H%M")
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter=";")
        writer.writerow([title for _k, title in COLUMNS])
        for row in rows:
            writer.writerow([_safe(row[k]) for k, _t in COLUMNS])
        return ("﻿" + buf.getvalue()).encode("utf-8"), f"axigen_accounts_{stamp}.csv", "text/csv; charset=utf-8"
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Учётные записи Axigen"
    ws.append([title for _k, title in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6FED")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    def cell_value(value):
        # формулы и управляющие знаки из чужих данных (ФИО из кадровой выгрузки, имена
        # ящиков) не должны ни сработать в Excel, ни испортить файл
        if isinstance(value, str):
            value = _safe(ILLEGAL_CHARACTERS_RE.sub(" ", value))
        return value

    for row in rows:
        ws.append([cell_value(row[key]) for key, _title in COLUMNS])
    widths = {"email": 32, "full_name": 30, "aliases": 30, "services": 30, "mailbox": 28, "errors": 40,
              "last_login": 18, "last_backup_at": 18}
    for i, (key, _title) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 14)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows) + 1)}"
    if data.get("orphans"):
        ws2 = wb.create_sheet("Нет на сервере")
        ws2.append(["Ящик архива", "Логин", "Копирование"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        for orphan in data["orphans"]:
            ws2.append([cell_value(orphan["name"]), cell_value(orphan["username"]),
                        "включено" if orphan["enabled"] else "выключено"])
        ws2.column_dimensions["A"].width = 30
        ws2.column_dimensions["B"].width = 32
    out = io.BytesIO()
    wb.save(out)
    return (out.getvalue(), f"axigen_accounts_{stamp}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
