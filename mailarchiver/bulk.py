"""
Групповые действия над почтовыми ящиками.

Каждое действие — класс-обработчик: какие параметры оно принимает, почему
пропускает ящик (выключен, архив удерживается, по ящику уже идёт задание…),
что сделает с каждым ящиком и что сделало. Один и тот же код работает в двух
режимах:

* **предпросмотр** (``preview=True``) — ничего не меняет и возвращает по
  каждому ящику «будет сделано …» или причину пропуска, а также
  предупреждения (например, что на диске может не хватить места);
* **выполнение** — делает то же самое по-настоящему и сохраняет итог в
  истории групповых операций (таблица ``bulk_ops``) и в аудите.

Опасные действия (копия «с нуля», удаление ящиков, заливка в рабочие папки)
требуют подтверждения: в ``confirm`` передаётся число ящиков, которые
действительно будут обработаны, — администратор видит его в предпросмотре и
вводит руками. Так нельзя случайно запустить действие на 580 ящиков вместо 5.

Правила, общие с действиями над одним ящиком (срок хранения, параметры
восстановления, наличие пароля), берутся из :mod:`mailarchiver.accountops`,
чтобы две копии правил не разошлись.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .accountops import (after_account_retention_change, after_hold_set, credential_problem, hold_label,
                         validate_cron, validate_restore_options)
from .cronutil import describe_cron
from .errors import MailArchiverError, ValidationError
from .logging_setup import get_logger
from . import roles
from .models import Account, AuthType, JobStatus, JobType, ScheduleKind, Security
from .util import human_size

log = get_logger("bulk")

#: Больше ящиков за одну операцию не принимаем (защита от ошибки в запросе).
MAX_IDS = 10000
#: Статусы итога по ящику.
OK, SKIP, FAIL = "ok", "skip", "fail"

GROUPS: List[Tuple[str, str]] = [
    ("copy", "Копирование и проверка"),
    ("manage", "Управление ящиками"),
    ("schedule", "Расписания"),
    ("connect", "Подключение и папки"),
    ("data", "Выгрузка и восстановление"),
    ("storage", "Хранилище"),
    ("danger", "Удаление"),
]

AUTH_LABELS = {AuthType.PASSWORD: "по паролю ящика", AuthType.MASTER: "через администратора почты",
               AuthType.OAUTH2: "OAuth2"}


def retention_label(days: Optional[int]) -> str:
    days = -1 if days is None else int(days)
    if days < 0:
        return "как в общих настройках"
    if days == 0:
        return "хранить всё"
    return f"последние {days} дн."


_hold_label = hold_label


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _split_list(text: str) -> List[str]:
    """Список из текста: по строке, а если строка одна — через запятую или «;».

    Имена папок бывают и с запятой («Отчёты, 2024»): если список уже разбит по
    строкам, запятые внутри строки не трогаем.
    """
    out: List[str] = []
    text = text or ""
    pattern = r"\n+" if "\n" in text.strip() else r"[,;]+"
    for chunk in re.split(pattern, text):
        item = chunk.strip()
        if item and item not in out:
            out.append(item)
    return out


# =====================================================================
#  Описание параметров и действий
# =====================================================================
@dataclass
class Param:
    key: str
    label: str
    kind: str = "text"          # bool | int | text | textarea | select | date | time
    default: Any = None
    options: Sequence[Tuple[Any, str]] = ()
    help: str = ""
    placeholder: str = ""
    min: Optional[int] = None
    max: Optional[int] = None
    #: показывать поле, только если другой параметр имеет одно из значений
    show_if: Optional[Dict[str, Sequence[Any]]] = None
    #: показывать поле только когда проверка выдала предупреждение (например, о месте на диске)
    on_warning: bool = False

    def to_dict(self) -> Dict[str, Any]:
        return {"key": self.key, "label": self.label, "kind": self.kind, "default": self.default,
                "options": [{"value": v, "label": t} for v, t in self.options], "help": self.help,
                "placeholder": self.placeholder, "min": self.min, "max": self.max,
                "show_if": {k: list(v) for k, v in (self.show_if or {}).items()} or None,
                "on_warning": self.on_warning}

    def coerce(self, raw: Any) -> Any:
        if raw is None or (isinstance(raw, str) and raw == "" and self.kind in ("int", "select")):
            raw = self.default
        if self.kind == "bool":
            if isinstance(raw, str):
                return raw.strip().lower() in ("1", "true", "yes", "on", "да")
            return bool(raw)
        if self.kind == "int":
            if raw is None or raw == "":
                return None
            try:
                value = int(raw)
            except (TypeError, ValueError):
                raise ValidationError(f"«{self.label}» — целое число.")
            if (self.min is not None and value < self.min) or (self.max is not None and value > self.max):
                raise ValidationError(f"«{self.label}» — от {self.min} до {self.max}.")
            return value
        if self.kind == "select":
            for value, _title in self.options:
                if str(value) == str(raw):
                    return value
            raise ValidationError(f"Недопустимое значение «{self.label}»: {raw}.")
        value = "" if raw is None else str(raw)
        if len(value) > 20000:
            raise ValidationError(f"Слишком длинное значение «{self.label}».")
        return value if self.kind == "textarea" else value.strip()


class BulkContext:
    """Состояние одной групповой операции (предпросмотра или выполнения)."""

    def __init__(self, svc, user: dict, params: Dict[str, Any], preview: bool) -> None:
        self.svc = svc
        self.db = svc.db
        self.user = user
        self.username = (user or {}).get("username") or "admin"
        self.params = params
        self.preview = preview
        self.warnings: List[str] = []
        self.results: List[Dict[str, Any]] = []
        self.jobs: List[int] = []
        self.data: Dict[str, Any] = {}      # общее для действия (готовится в prepare)
        self.per: Dict[int, Any] = {}       # посчитанное по ящику в check (для plan/apply)
        self.extra: Dict[str, Any] = {}     # дополнительные сведения в ответ интерфейсу
        self._active: Optional[Dict[int, List[Any]]] = None
        self._totals: Optional[Dict[int, Dict[str, int]]] = None
        self._employees = None
        self._schedules = None

    # -- общие сведения (считаются один раз на операцию) ---------------------
    @property
    def active(self) -> Dict[int, List[Any]]:
        """Выполняющиеся и ждущие задания по ящикам."""
        if self._active is None:
            out: Dict[int, List[Any]] = {}
            for job in self.db.active_jobs():
                if job["account_id"] is not None:
                    out.setdefault(int(job["account_id"]), []).append(job)
            self._active = out
        return self._active

    def busy(self, account_id: int, types: Optional[Iterable[str]] = None) -> bool:
        """Идёт ли по ящику задание (любое или указанных типов).

        Учитывается и копирование всех ящиков по очереди: пока оно занимается
        ящиком, по нему как будто идёт своё задание.
        """
        wanted = set(types) if types else None
        for job in self.active.get(account_id, []):
            if wanted is None or job["type"] in wanted:
                return True
        queue = getattr(self.svc, "queue", None)
        if queue is not None and (wanted is None or JobType.BACKUP in wanted):
            try:
                if account_id in queue.held_accounts():
                    return True
            except Exception:  # noqa: BLE001
                pass
        return False

    @property
    def totals(self) -> Dict[int, Dict[str, int]]:
        if self._totals is None:
            self._totals = self.db.message_totals_by_account()
        return self._totals

    def messages(self, account_id: int) -> int:
        return int((self.totals.get(account_id) or {}).get("messages", 0))

    def size(self, account_id: int) -> int:
        return int((self.totals.get(account_id) or {}).get("bytes", 0))

    @property
    def employees(self):
        if self._employees is None:
            self._employees = self.db.employees_by_account()
        return self._employees

    @property
    def schedules(self):
        if self._schedules is None:
            self._schedules = self.db.schedules_by_account()
        return self._schedules

    @property
    def retry_attempts(self) -> int:
        try:
            return max(1, int(self.svc.rt("backup", "retry_attempts") or 1))
        except (TypeError, ValueError):
            return 1

    def enqueue(self, job_type: str, account_id: Optional[int], params: Dict[str, Any], *,
                priority: int = 5, max_attempts: int = 1) -> int:
        job_id = self.svc.queue.enqueue(job_type, account_id, params, priority=priority,
                                        max_attempts=max_attempts, created_by=self.username)
        self.jobs.append(job_id)
        return job_id


class Handler:
    """Групповое действие. Наследники переопределяют нужные шаги."""

    key = ""
    label = ""
    group = "manage"
    icon = "•"
    desc = ""
    #: 0 — обычное, 1 — внимательно (необратимо в мелочах), 2 — опасное (ввести число ящиков)
    danger = 0
    #: действие ставит задания в очередь (интерфейс покажет ссылку на «Очередь»)
    creates_jobs = False
    #: короткая пометка «что важно знать» для интерфейса
    note = ""
    params: Tuple[Param, ...] = ()

    def params_for(self, svc) -> Sequence[Param]:
        return self.params

    def danger_for(self, ctx: BulkContext) -> int:
        return self.danger

    def prepare(self, ctx: BulkContext) -> None:
        """Проверить параметры целиком (ValidationError — операция не начнётся)."""

    def check(self, ctx: BulkContext, acc: Account) -> str:
        """Почему ящик пропускается (пусто — будет обработан)."""
        return ""

    def validate_plan(self, ctx: BulkContext, accounts: List[Account]) -> None:
        """Проверка всего набора ящиков, прошедших check: предупреждения или отказ."""

    def plan(self, ctx: BulkContext, acc: Account) -> str:
        """Что будет сделано с ящиком (для предпросмотра)."""
        return "будет выполнено"

    def apply(self, ctx: BulkContext, acc: Account) -> Tuple[str, Optional[int]]:
        raise NotImplementedError

    def finish(self, ctx: BulkContext) -> None:
        """После обхода всех ящиков (одно общее задание, перечитать планировщик…)."""

    def describe(self, svc) -> Dict[str, Any]:
        return {"key": self.key, "label": self.label, "group": self.group, "icon": self.icon,
                "desc": self.desc, "danger": self.danger, "creates_jobs": self.creates_jobs,
                "note": self.note, "params": [p.to_dict() for p in self.params_for(svc)]}


# =====================================================================
#  Копирование и проверка
# =====================================================================
class BackupAction(Handler):
    key = "backup"
    label = "Сделать копию сейчас"
    group = "copy"
    icon = "💾"
    desc = ("Поставить в очередь обычное копирование: скачиваются только новые письма. "
            "Ящики копируются параллельно — сколько позволяет «Одновременных заданий».")
    creates_jobs = True
    rebuild = ""
    busy_types: Optional[Tuple[str, ...]] = (JobType.BACKUP,)
    busy_text = "копирование уже идёт или стоит в очереди"

    def check(self, ctx, acc):
        if not acc.enabled:
            return "копирование ящика выключено"
        problem = credential_problem(acc)
        if problem:
            return problem
        if ctx.busy(acc.id, self.busy_types):
            return self.busy_text
        return ""

    def plan(self, ctx, acc):
        return "копирование будет поставлено в очередь"

    def job_params(self, ctx) -> Dict[str, Any]:
        return {"rebuild": self.rebuild} if self.rebuild else {}

    def apply(self, ctx, acc):
        attempts = 1 if self.rebuild else ctx.retry_attempts
        job_id = ctx.enqueue(JobType.BACKUP, acc.id, self.job_params(ctx), max_attempts=attempts)
        return "копирование в очереди", job_id


class BackupSequenceAction(Handler):
    key = "backup_sequence"
    label = "Копировать по очереди, ящик за ящиком"
    group = "copy"
    icon = "🔁"
    desc = ("Одно задание копирует выбранные ящики строго по одному — сам проход никогда не копирует два "
            "ящика сразу. Сначала — ящики, которые копировались давнее всего.")
    creates_jobs = True
    params = (
        Param("pause_seconds", "Пауза между ящиками, секунд", "int", 0, min=0, max=3600,
              help="Дать почтовому серверу передышку между ящиками. 0 — без паузы."),
    )

    def prepare(self, ctx):
        running = next((j for j in ctx.db.active_jobs() if j["type"] == JobType.BACKUP_ALL), None)
        if running is not None:
            ctx.warnings.append(f"Копирование по очереди уже идёт (задание №{running['id']}) — новое "
                                f"начнётся после него.")

    def check(self, ctx, acc):
        if not acc.enabled:
            return "копирование ящика выключено"
        return credential_problem(acc)

    def plan(self, ctx, acc):
        return "войдёт в проход по очереди"

    def apply(self, ctx, acc):
        ctx.data.setdefault("ids", []).append(acc.id)
        return "в проходе по очереди", None

    def finish(self, ctx):
        ids = ctx.data.get("ids") or []
        if ctx.preview or not ids:
            return
        job_id = ctx.enqueue(JobType.BACKUP_ALL, None,
                             {"account_ids": ids, "pause_seconds": int(ctx.params.get("pause_seconds") or 0)},
                             max_attempts=ctx.retry_attempts)
        for item in ctx.results:
            if item["status"] == OK:
                item["job_id"] = job_id


class RebuildMissingAction(BackupAction):
    key = "rebuild_missing"
    label = "Докачать потерянные письма"
    icon = "🩹"
    desc = ("Сверить индекс с файлами на диске: письма, файлы которых пропали, скачиваются заново. "
            "Ничего не удаляется — безопасный режим.")
    rebuild = "missing"

    def plan(self, ctx, acc):
        return f"сверка {ctx.messages(acc.id)} писем с диском, потерянные скачаются заново"


class RebuildFullAction(BackupAction):
    key = "rebuild_full"
    label = "Скопировать с нуля"
    icon = "🧨"
    danger = 2
    desc = ("Прежняя копия каждого ящика переносится в карантин («Прежние копии»), а все письма "
            "скачиваются с сервера заново. Письма, которых на сервере уже нет, останутся только в "
            "карантине — в просмотре, поиске и экспорте их не будет. После копирования прежняя копия "
            "сама сравнивается с новой.")
    note = ("Карантин занимает место, пока его не удалить. Письма, которые остались только в нём, можно вернуть "
            "в архив («Вернуть письма из прежних копий»), затем удалить карантин. Ящики с удержанием архива "
            "пропускаются.")
    rebuild = "full"
    busy_types = None
    busy_text = "по ящику выполняется задание — дождитесь его окончания"
    params = (
        Param("force_space", "Запустить, даже если места на диске может не хватить", "bool", False,
              help="Прежние копии остаются на диске, пока вы не удалите карантин, — новая копия займёт "
                   "столько же места ещё раз. Если место кончится, копирование ящиков остановится само.",
              on_warning=True),
    )

    def check(self, ctx, acc):
        if acc.on_hold():
            return f"архив удерживается ({_hold_label(acc.hold_until)}) — «с нуля» не пересоздаём"
        return super().check(ctx, acc)

    def validate_plan(self, ctx, accounts):
        from .util import disk_free_bytes
        need = sum(ctx.size(a.id) for a in accounts)
        try:
            free = disk_free_bytes(ctx.svc.store.mail_root)
        except Exception:  # noqa: BLE001
            free = 0
        reserve = int(getattr(ctx.svc.store, "min_free_bytes", 0) or 0)
        ctx.extra["disk"] = {"need": need, "need_h": human_size(need), "free": free,
                             "free_h": human_size(free) if free else "неизвестно"}
        if free and need + reserve > free:
            text = (f"Места на диске может не хватить: прежние копии ({human_size(need)}) остаются в "
                    f"карантине, новые займут примерно столько же, а свободно {human_size(free)}.")
            ctx.warnings.append(text)
            if not ctx.preview and not ctx.params.get("force_space"):
                raise ValidationError(text, hint="Удалите прежние копии (карантин), запустите «с нуля» "
                                                 "для части ящиков или включите «Запустить, даже если места "
                                                 "на диске может не хватить».")

    def plan(self, ctx, acc):
        n = ctx.messages(acc.id)
        if not n:
            return "копии ещё нет — будет скачано всё"
        return f"прежняя копия ({n} писем, {human_size(ctx.size(acc.id))}) уйдёт в карантин, всё скачается заново"

    def apply(self, ctx, acc):
        ctx.db.add_audit(ctx.username, "backup_rebuild_full_request", f"{acc.name} — групповое действие")
        return super().apply(ctx, acc)


class FinalBackupAction(BackupAction):
    key = "final_backup"
    label = "Последняя копия, затем выключить"
    icon = "🏁"
    danger = 1
    desc = ("Сделать копию и сразу после неё выключить копирование ящика. Не получилась копия — ящик "
            "остаётся включённым, чтобы её можно было повторить.")
    busy_types = None
    busy_text = "по ящику выполняется задание"

    def plan(self, ctx, acc):
        return "копия, затем копирование ящика выключится"

    def job_params(self, ctx):
        return {"disable_after": True}


class VerifyAction(Handler):
    key = "verify"
    label = "Проверить целостность копии"
    group = "copy"
    icon = "🔍"
    desc = "Прочитать каждое письмо копии и сверить контрольную сумму — найти пропавшие и повреждённые файлы."
    creates_jobs = True

    def check(self, ctx, acc):
        if not ctx.messages(acc.id):
            return "в копии нет писем"
        if ctx.busy(acc.id, (JobType.VERIFY,)):
            return "проверка уже идёт или стоит в очереди"
        return ""

    def plan(self, ctx, acc):
        return f"будет проверено писем: {ctx.messages(acc.id)}"

    def apply(self, ctx, acc):
        return "проверка в очереди", ctx.enqueue(JobType.VERIFY, acc.id, {})


class CheckLoginsAction(Handler):
    key = "check_logins"
    label = "Проверить пароли (вход)"
    group = "copy"
    icon = "🔐"
    desc = ("Войти в каждый ящик и сразу выйти — письма не скачиваются. Итог появится в колонке «Вход». "
            "Неверные пароли — это неудачные входы на почтовом сервере (осторожно с fail2ban).")
    creates_jobs = True

    def prepare(self, ctx):
        running = next((j for j in ctx.db.active_jobs() if j["type"] == JobType.CHECK_LOGINS), None)
        ctx.data["running"] = running["id"] if running is not None else None

    def check(self, ctx, acc):
        if ctx.data.get("running"):
            return f"проверка паролей уже идёт (задание №{ctx.data['running']})"
        return credential_problem(acc)

    def plan(self, ctx, acc):
        return "вход будет проверен"

    def apply(self, ctx, acc):
        ctx.data.setdefault("ids", []).append(acc.id)
        return "в проверке паролей", None

    def finish(self, ctx):
        ids = ctx.data.get("ids") or []
        if ctx.preview or not ids:
            return
        job_id = ctx.enqueue(JobType.CHECK_LOGINS, None, {"account_ids": sorted(ids), "only_enabled": False},
                             priority=3)
        for item in ctx.results:
            if item["status"] == OK:
                item["job_id"] = job_id


class RetentionRunAction(Handler):
    key = "retention_run"
    label = "Очистить по сроку хранения сейчас"
    group = "copy"
    icon = "🧹"
    danger = 1
    desc = "Удалить из копии письма старше срока хранения ящика прямо сейчас, не дожидаясь ночной очистки."
    creates_jobs = True

    def check(self, ctx, acc):
        from .queue.jobs import _count_older_than, effective_retention_days, retention_cutoff_iso
        if acc.on_hold():
            return f"архив удерживается ({_hold_label(acc.hold_until)})"
        days = effective_retention_days(ctx.svc, acc)
        if days <= 0:
            return "срок хранения не задан — хранится всё"
        if ctx.busy(acc.id, (JobType.RETENTION,)):
            return "очистка уже идёт или стоит в очереди"
        count = _count_older_than(ctx.db, acc.id, retention_cutoff_iso(days))
        if not count:
            return f"писем старше {days} дн. нет"
        ctx.per[acc.id] = (days, count)
        return ""

    def plan(self, ctx, acc):
        days, count = ctx.per[acc.id]
        return f"будет удалено писем старше {days} дн.: {count}"

    def apply(self, ctx, acc):
        return "очистка в очереди", ctx.enqueue(JobType.RETENTION, acc.id, {})


class CancelJobsAction(Handler):
    key = "cancel_jobs"
    label = "Отменить задания"
    group = "copy"
    icon = "⏹️"
    desc = ("Снять задания выбранных ящиков: стоящие в очереди отменяются сразу, выполняющиеся — "
            "останавливаются на ближайшем шаге.")
    params = (
        Param("which", "Какие задания", "select", "all",
              options=(("all", "все задания ящиков"), ("backup", "только копирование"),
                       ("queued", "только стоящие в очереди (выполняющиеся не трогать)"))),
    )

    def _jobs(self, ctx, acc):
        which = ctx.params.get("which") or "all"
        jobs = list(ctx.active.get(acc.id, []))
        if roles.is_operator(ctx.user):
            # оператор снимает только копирование и проверки (не «с нуля»,
            # не последнюю копию перед выключением, не выгрузки администратора)
            jobs = [j for j in jobs if roles.operator_may_manage(j)]
        if which == "backup":
            jobs = [j for j in jobs if j["type"] == JobType.BACKUP]
        elif which == "queued":
            jobs = [j for j in jobs if j["status"] == JobStatus.QUEUED]
        return jobs

    def check(self, ctx, acc):
        jobs = self._jobs(ctx, acc)
        if not jobs:
            return "подходящих заданий нет"
        ctx.per[acc.id] = jobs
        return ""

    def plan(self, ctx, acc):
        return f"будет отменено заданий: {len(ctx.per[acc.id])}"

    def apply(self, ctx, acc):
        jobs = ctx.per[acc.id]
        for job in jobs:
            ctx.svc.queue.cancel(int(job["id"]))
        return f"отменено заданий: {len(jobs)}", None


# =====================================================================
#  Управление ящиками
# =====================================================================
class EnableAction(Handler):
    key = "enable"
    label = "Включить копирование"
    group = "manage"
    icon = "▶️"
    desc = "Включить ящики: они снова будут копироваться по расписаниям и кнопкам."

    def check(self, ctx, acc):
        if acc.enabled:
            return "уже включён"
        return ""

    def plan(self, ctx, acc):
        notes = []
        if credential_problem(acc):
            notes.append(credential_problem(acc) + " — копироваться не будет")
        if acc.dismissed_at:
            notes.append("сотрудник уволен")
        return "будет включён" + (f" ({'; '.join(notes)})" if notes else "")

    def apply(self, ctx, acc):
        ctx.db.execute("UPDATE accounts SET enabled=1, auto_disabled=0, updated_at=? WHERE id=?",
                       (datetime.now(timezone.utc).isoformat(), acc.id))
        return self.plan(ctx, acc).replace("будет включён", "включён"), None


class DisableAction(Handler):
    key = "disable"
    label = "Выключить копирование"
    group = "manage"
    icon = "⏸️"
    desc = "Выключить ящики: расписания и кнопки их больше не копируют. Архив писем остаётся как есть."
    params = (
        Param("cancel_queued", "Снять копирования этих ящиков, стоящие в очереди", "bool", True),
    )

    def check(self, ctx, acc):
        return "" if acc.enabled else "уже выключен"

    def plan(self, ctx, acc):
        queued = [j for j in ctx.active.get(acc.id, [])
                  if j["type"] == JobType.BACKUP and j["status"] == JobStatus.QUEUED]
        tail = f"; из очереди снимется копирований: {len(queued)}" if queued and ctx.params.get("cancel_queued") else ""
        return "будет выключен" + tail

    def apply(self, ctx, acc):
        ctx.db.set_account_enabled(acc.id, False)
        cancelled = 0
        if ctx.params.get("cancel_queued"):
            for job in ctx.active.get(acc.id, []):
                if job["type"] == JobType.BACKUP and job["status"] == JobStatus.QUEUED:
                    ctx.svc.queue.cancel(int(job["id"]))
                    cancelled += 1
        return "выключен" + (f"; снято из очереди: {cancelled}" if cancelled else ""), None


RETENTION_CHOICES = ((-1, "как в общих настройках"), (0, "хранить всё"), (3, "последние 3 дня"),
                     (7, "последняя неделя"), (14, "2 недели"), (30, "30 дней"), (90, "3 месяца"),
                     (180, "полгода"), (365, "год"), (730, "2 года"), (1095, "3 года"), (1825, "5 лет"),
                     ("custom", "свой срок…"))


class RetentionSetAction(Handler):
    key = "retention_set"
    label = "Срок хранения копий"
    group = "manage"
    icon = "🗓️"
    danger = 1
    desc = ("Сколько дней хранить письма в локальной копии. Письма старше срока удаляет ежедневная "
            "очистка (на сервере они остаются). Удерживаемые архивы очистка не трогает.")
    params = (
        Param("days", "Срок", "select", 30, options=RETENTION_CHOICES),
        Param("custom_days", "Свой срок, дней", "int", 60, min=1, max=36500, show_if={"days": ["custom"]}),
        Param("run_now", "Сразу запустить очистку (иначе — ночью по расписанию)", "bool", False),
    )

    def prepare(self, ctx):
        days = ctx.params.get("days")
        if days == "custom":
            days = ctx.params.get("custom_days")
            if not days:
                raise ValidationError("Укажите свой срок хранения в днях.")
        ctx.data["days"] = int(days)

    def check(self, ctx, acc):
        if int(acc.retention_days if acc.retention_days is not None else -1) == ctx.data["days"]:
            return "срок уже такой"
        return ""

    def plan(self, ctx, acc):
        text = f"{retention_label(acc.retention_days)} → {retention_label(ctx.data['days'])}"
        if acc.on_hold():
            text += f" (архив удерживается {_hold_label(acc.hold_until)} — до конца удержания не чистится)"
        return text

    def apply(self, ctx, acc):
        days = ctx.data["days"]
        ctx.db.set_account_retention(acc.id, days)
        after_account_retention_change(ctx.svc, acc)
        job_id = None
        if ctx.params.get("run_now") and days > 0 and not acc.on_hold() and not ctx.busy(acc.id, (JobType.RETENTION,)):
            job_id = ctx.enqueue(JobType.RETENTION, acc.id, {})
        return self.plan(ctx, acc) + ("; очистка в очереди" if job_id else ""), job_id


class HoldSetAction(Handler):
    key = "hold_set"
    label = "Удержание архива"
    group = "manage"
    icon = "🔒"
    desc = ("Пока действует удержание, письма ящика не удаляются очисткой по сроку хранения, "
            "а сам архив нельзя удалить или пересоздать «с нуля».")
    params = (
        Param("forever", "Бессрочно", "bool", False),
        Param("until", "Удерживать до", "date", "", show_if={"forever": [False]}),
    )

    def prepare(self, ctx):
        from .employees import HOLD_FOREVER
        from .util import parse_day
        if ctx.params.get("forever"):
            ctx.data["until"] = HOLD_FOREVER
            return
        raw = (ctx.params.get("until") or "").strip()
        if not raw:
            raise ValidationError("Укажите дату окончания удержания или «Бессрочно».")
        day = parse_day(raw)
        if day.isoformat() < datetime.now(timezone.utc).date().isoformat():
            raise ValidationError("Дата окончания удержания уже прошла.")
        ctx.data["until"] = day.isoformat()

    def check(self, ctx, acc):
        return "удержание уже такое" if acc.hold_until == ctx.data["until"] else ""

    def plan(self, ctx, acc):
        return f"удержание: {_hold_label(acc.hold_until)} → {_hold_label(ctx.data['until'])}"

    def apply(self, ctx, acc):
        text = self.plan(ctx, acc)
        ctx.db.set_account_hold(acc.id, ctx.data["until"], "manual")
        after_hold_set(ctx.svc, acc.id)
        return text, None


class HoldClearAction(Handler):
    key = "hold_clear"
    label = "Снять удержание архива"
    group = "manage"
    icon = "🔓"
    danger = 1
    desc = "Письма ящиков снова будут удаляться по сроку хранения (если он задан), архив можно будет удалить."

    def check(self, ctx, acc):
        return "" if acc.hold_until else "удержания нет"

    def plan(self, ctx, acc):
        return (f"удержание ({_hold_label(acc.hold_until)}) будет снято"
                + ("; сотрудник уволен" if acc.dismissed_at else ""))

    def apply(self, ctx, acc):
        ctx.db.set_account_hold(acc.id, "", "")
        ctx.db.add_audit(ctx.username, "account_hold", f"{acc.name}: снято — групповое действие")
        return f"удержание ({_hold_label(acc.hold_until)}) снято", None


class LogoutSessionsAction(Handler):
    key = "logout_sessions"
    label = "Завершить сеансы сотрудников"
    group = "manage"
    icon = "🚪"
    desc = "Разлогинить сотрудников, вошедших в веб-интерфейс по паролю своего ящика."

    def check(self, ctx, acc):
        count = ctx.db.count_active_sessions(acc.id)
        if not count:
            return "открытых сеансов нет"
        ctx.per[acc.id] = count
        return ""

    def plan(self, ctx, acc):
        return f"будет завершено сеансов: {ctx.per[acc.id]}"

    def apply(self, ctx, acc):
        closed = ctx.db.delete_mailbox_sessions(acc.id)
        return f"завершено сеансов: {closed}", None


class RenameAction(Handler):
    key = "rename"
    label = "Переименовать по шаблону"
    group = "manage"
    icon = "🏷️"
    desc = ("Задать названия ящиков по шаблону. Подстановки: {employee} — ФИО сотрудника, {department} — "
            "отдел, {position} — должность, {username} — адрес ящика, {login} — адрес до «@», "
            "{domain} — домен, {name} — прежнее название.")
    TOKENS = ("employee", "department", "position", "username", "login", "domain", "name")
    params = (
        Param("template", "Шаблон названия", "text", "{employee}", placeholder="{employee} ({login})"),
        Param("fallback", "Если у ящика нет сотрудника", "select", "skip",
              options=(("skip", "не переименовывать"), ("username", "назвать адресом ящика"))),
    )

    def prepare(self, ctx):
        template = (ctx.params.get("template") or "").strip()
        if not template:
            raise ValidationError("Укажите шаблон названия.", hint="Например «{employee}» или «{employee} ({login})».")
        unknown = [t for t in re.findall(r"\{([^{}]*)\}", template) if t not in self.TOKENS]
        if unknown:
            raise ValidationError(f"Неизвестная подстановка: {{{unknown[0]}}}.",
                                  hint="Допустимо: " + ", ".join("{" + t + "}" for t in self.TOKENS) + ".")
        ctx.data["template"] = template

    def _render(self, ctx, acc) -> Tuple[str, str]:
        emp = ctx.employees.get(acc.id)
        template = ctx.data["template"]
        uses_emp = any(f"{{{t}}}" in template for t in ("employee", "department", "position"))
        if uses_emp and emp is None:
            if ctx.params.get("fallback") == "username":
                return acc.username, ""
            return "", "у ящика нет сотрудника в справочнике"
        login, _, domain = (acc.username or "").partition("@")
        values = {"employee": (emp["full_name"] if emp else "") or "", "department": (emp["department"] if emp else "") or "",
                  "position": (emp["position"] if emp else "") or "", "username": acc.username or "",
                  "login": login, "domain": domain, "name": acc.name or ""}
        name = re.sub(r"\{(\w+)\}", lambda m: values.get(m.group(1), ""), template)
        # пустые подстановки не должны оставлять «()» и висящие разделители
        name = re.sub(r"\(\s*\)|\[\s*\]", "", name)
        name = re.sub(r"\s+", " ", name).strip(" -—,;")
        return name[:200], ""

    def check(self, ctx, acc):
        name, problem = self._render(ctx, acc)
        if problem:
            return problem
        if not name:
            return "по шаблону получилось пустое название"
        if name == acc.name:
            return "название уже такое"
        ctx.per[acc.id] = name
        return ""

    def plan(self, ctx, acc):
        return f"«{acc.name}» → «{ctx.per[acc.id]}»"

    def apply(self, ctx, acc):
        ctx.db.execute("UPDATE accounts SET name=?, updated_at=? WHERE id=?",
                       (ctx.per[acc.id], datetime.now(timezone.utc).isoformat(), acc.id))
        return self.plan(ctx, acc), None


class NotesAction(Handler):
    key = "notes"
    label = "Заметка к ящикам"
    group = "manage"
    icon = "📝"
    desc = "Дописать, заменить или стереть заметку администратора в карточках ящиков."
    params = (
        Param("op", "Что сделать", "select", "append",
              options=(("append", "дописать с новой строки"), ("replace", "заменить заметку"),
                       ("clear", "стереть заметку"))),
        Param("text", "Текст", "textarea", "", show_if={"op": ["append", "replace"]}),
    )

    def prepare(self, ctx):
        if ctx.params.get("op") != "clear" and not (ctx.params.get("text") or "").strip():
            raise ValidationError("Введите текст заметки.")

    def _new(self, ctx, acc) -> str:
        op = ctx.params.get("op")
        text = (ctx.params.get("text") or "").strip()
        old = acc.notes or ""
        if op == "clear":
            return ""
        if op == "replace":
            return text
        return (old.rstrip() + "\n" + text).strip() if old.strip() else text

    def check(self, ctx, acc):
        new = self._new(ctx, acc)
        if new == (acc.notes or ""):
            return "заметка не изменится"
        if len(new) > 5000:
            return "заметка получилась длиннее 5000 символов"
        ctx.per[acc.id] = new
        return ""

    def plan(self, ctx, acc):
        new = ctx.per[acc.id]
        return "заметка будет стёрта" if not new else f"заметка: «{new[-120:]}»"

    def apply(self, ctx, acc):
        ctx.db.execute("UPDATE accounts SET notes=?, updated_at=? WHERE id=?",
                       (ctx.per[acc.id], datetime.now(timezone.utc).isoformat(), acc.id))
        return "заметка сохранена" if ctx.per[acc.id] else "заметка стёрта", None


# =====================================================================
#  Расписания
# =====================================================================
def _parse_hhmm(raw: str, what: str) -> Tuple[int, int]:
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", (raw or "").strip())
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        raise ValidationError(f"{what}: время указывается как ЧЧ:ММ, например 02:30.")
    return int(m.group(1)), int(m.group(2))


class ScheduleSetAction(Handler):
    key = "schedule_set"
    label = "Назначить расписание копирования"
    group = "schedule"
    icon = "⏰"
    desc = ("Задать ящикам расписание резервного копирования. Чтобы сотни ящиков не стартовали "
            "в одну минуту, запуск можно разнести по времени.")
    note = "Если нужно копировать все ящики строго по очереди — заведите в «Расписаниях» задание «Копирование всех ящиков по очереди»."
    params = (
        Param("when", "Когда", "select", "daily",
              options=(("daily", "каждый день"), ("weekdays", "по будням (пн–пт)"),
                       ("cron", "своё cron-выражение"), ("interval", "через интервал"))),
        Param("time", "Время запуска", "time", "02:00", show_if={"when": ["daily", "weekdays"]}),
        Param("spread_minutes", "Разнести запуск ящиков на, минут", "int", 0, min=0, max=720,
              show_if={"when": ["daily", "weekdays"]},
              help="0 — все ящики в одно время. Например, 120: запуски равномерно распределятся с 02:00 до 04:00."),
        Param("cron_expr", "Cron-выражение", "text", "0 2 * * *", show_if={"when": ["cron"]},
              placeholder="0 2 * * *"),
        Param("interval_hours", "Интервал, часов", "int", 24, min=1, max=8760, show_if={"when": ["interval"]}),
        Param("mode", "Если у ящика уже есть расписание копирования", "select", "replace",
              options=(("replace", "заменить его новым"), ("missing", "не трогать ящик"),
                       ("add", "добавить ещё одно"))),
        Param("enabled", "Расписание включено", "bool", True),
    )

    def prepare(self, ctx):
        when = ctx.params.get("when")
        if when in ("daily", "weekdays"):
            hh, mm = _parse_hhmm(ctx.params.get("time") or "", "Время запуска")
            ctx.data.update(kind=ScheduleKind.CRON, base=hh * 60 + mm,
                            dow="1-5" if when == "weekdays" else "*",
                            spread=int(ctx.params.get("spread_minutes") or 0))
        elif when == "cron":
            expr = (ctx.params.get("cron_expr") or "").strip()
            validate_cron(expr)
            ctx.data.update(kind=ScheduleKind.CRON, cron=expr)
        else:
            hours = int(ctx.params.get("interval_hours") or 0)
            if hours < 1:
                raise ValidationError("Интервал — от 1 часа.")
            ctx.data.update(kind=ScheduleKind.INTERVAL, seconds=hours * 3600)

    def validate_plan(self, ctx, accounts):
        ctx.data["order"] = {a.id: i for i, a in enumerate(accounts)}
        ctx.data["count"] = len(accounts)

    def _cron_for(self, ctx, acc) -> str:
        if "cron" in ctx.data:
            return ctx.data["cron"]
        count = max(1, ctx.data.get("count") or 1)
        offset = (ctx.data["order"].get(acc.id, 0) * ctx.data["spread"]) // count if ctx.data["spread"] else 0
        total = ctx.data["base"] + offset
        minute_of_day = total % 1440
        dow = ctx.data["dow"]
        if total >= 1440 and dow == "1-5":
            dow = "2-6"          # «по будням в 23:30» + разнос за полночь: запуск уже на следующий день
        return f"{minute_of_day % 60} {minute_of_day // 60} * * {dow}"

    def _backup_schedules(self, ctx, acc):
        return [s for s in ctx.schedules.get(acc.id, []) if s["job_type"] == JobType.BACKUP]

    def check(self, ctx, acc):
        existing = self._backup_schedules(ctx, acc)
        mode = ctx.params.get("mode")
        if existing and mode == "missing":
            return "у ящика уже есть расписание копирования"
        return ""

    def _describe(self, ctx, acc) -> str:
        if ctx.data["kind"] == ScheduleKind.INTERVAL:
            return f"каждые {ctx.data['seconds'] // 3600} ч"
        return describe_cron(self._cron_for(ctx, acc))

    def _text(self, ctx, acc, done: bool) -> str:
        existing = self._backup_schedules(ctx, acc)
        text = ("назначено" if done else "расписание") + f": {self._describe(ctx, acc)}"
        if not ctx.params.get("enabled"):
            text += " (выключено)"
        if existing and ctx.params.get("mode") == "replace":
            text += (f"; заменено прежних: {len(existing)}" if done
                     else f"; прежних расписаний копирования заменится: {len(existing)}")
        elif existing and ctx.params.get("mode") == "add":
            text += f"; прежних остаётся: {len(existing)}"
        return text

    def plan(self, ctx, acc):
        return self._text(ctx, acc, done=False)

    def apply(self, ctx, acc):
        text = self._text(ctx, acc, done=True)
        if ctx.params.get("mode") == "replace":
            for sched in self._backup_schedules(ctx, acc):
                ctx.db.delete_schedule(int(sched["id"]))
        if ctx.data["kind"] == ScheduleKind.INTERVAL:
            ctx.db.create_schedule(acc.id, ScheduleKind.INTERVAL, JobType.BACKUP, "",
                                   ctx.data["seconds"], bool(ctx.params.get("enabled")), {})
        else:
            ctx.db.create_schedule(acc.id, ScheduleKind.CRON, JobType.BACKUP, self._cron_for(ctx, acc), 0,
                                   bool(ctx.params.get("enabled")), {})
        return text, None

    def finish(self, ctx):
        if not ctx.preview:
            _reload_scheduler(ctx)


def _reload_scheduler(ctx) -> None:
    scheduler = getattr(ctx.svc, "scheduler", None)
    if scheduler is not None:
        try:
            scheduler.reload()
        except Exception:  # noqa: BLE001
            log.exception("Не удалось перечитать расписания после группового действия")


_SCHEDULE_WHICH = (("backup", "только расписания копирования"), ("all", "все расписания ящиков"))


class ScheduleRemoveAction(Handler):
    key = "schedule_remove"
    label = "Удалить расписания"
    group = "schedule"
    icon = "🚫"
    danger = 1
    desc = ("Удалить расписания выбранных ящиков — например, перед переходом на «Копирование всех ящиков "
            "по очереди», чтобы ящики не копировались дважды.")
    params = (Param("which", "Какие расписания", "select", "backup", options=_SCHEDULE_WHICH),)

    def _list(self, ctx, acc):
        items = ctx.schedules.get(acc.id, [])
        if ctx.params.get("which") != "all":
            items = [s for s in items if s["job_type"] == JobType.BACKUP]
        return items

    def check(self, ctx, acc):
        return "" if self._list(ctx, acc) else "расписаний нет"

    def plan(self, ctx, acc):
        return f"будет удалено расписаний: {len(self._list(ctx, acc))}"

    def apply(self, ctx, acc):
        items = self._list(ctx, acc)
        for sched in items:
            ctx.db.delete_schedule(int(sched["id"]))
        return f"удалено расписаний: {len(items)}", None

    def finish(self, ctx):
        if not ctx.preview:
            _reload_scheduler(ctx)


class ScheduleToggleAction(ScheduleRemoveAction):
    key = "schedule_toggle"
    label = "Включить или выключить расписания"
    icon = "⏯️"
    danger = 0
    desc = "Временно выключить (или снова включить) расписания выбранных ящиков, не удаляя их."
    params = (
        Param("enabled", "Расписания включены", "bool", False),
        Param("which", "Какие расписания", "select", "backup", options=_SCHEDULE_WHICH),
    )

    def check(self, ctx, acc):
        items = self._list(ctx, acc)
        if not items:
            return "расписаний нет"
        want = bool(ctx.params.get("enabled"))
        if all(bool(s["enabled"]) == want for s in items):
            return "уже включены" if want else "уже выключены"
        return ""

    def plan(self, ctx, acc):
        want = bool(ctx.params.get("enabled"))
        changed = [s for s in self._list(ctx, acc) if bool(s["enabled"]) != want]
        return f"будет {'включено' if want else 'выключено'} расписаний: {len(changed)}"

    def apply(self, ctx, acc):
        want = bool(ctx.params.get("enabled"))
        changed = [s for s in self._list(ctx, acc) if bool(s["enabled"]) != want]
        for sched in changed:
            ctx.db.update_schedule(int(sched["id"]), enabled=want)
        return f"{'включено' if want else 'выключено'} расписаний: {len(changed)}", None


# =====================================================================
#  Подключение и папки
# =====================================================================
class SetServerAction(Handler):
    key = "set_server"
    label = "Сменить сервер IMAP"
    group = "connect"
    icon = "🖥️"
    danger = 1
    desc = ("Поменять адрес почтового сервера, порт или шифрование сразу у многих ящиков — например, "
            "при переезде почты. Пустое поле — не менять.")
    note = ("Если это другой почтовый сервер (переезд), у папок там другой UIDVALIDITY: следующее копирование "
            "скачает письма заново, а прежние останутся в архиве как история — проверьте место на диске. "
            "Новое имя того же сервера перекачки не вызывает.")
    params = (
        Param("host", "IMAP-сервер", "text", "", placeholder="mx.example.ru"),
        Param("port", "Порт (пусто — не менять)", "int", None, min=1, max=65535),
        Param("security", "Шифрование", "select", "",
              options=(("", "не менять"), (Security.SSL, Security.LABELS[Security.SSL]),
                       (Security.STARTTLS, Security.LABELS[Security.STARTTLS]),
                       (Security.PLAIN, Security.LABELS[Security.PLAIN]))),
    )

    def prepare(self, ctx):
        host = (ctx.params.get("host") or "").strip()
        if host and (len(host) > 255 or re.search(r"\s", host)):
            raise ValidationError("Адрес сервера указан неверно.")
        if not host and not ctx.params.get("port") and not ctx.params.get("security"):
            raise ValidationError("Укажите, что поменять: сервер, порт или шифрование.")
        ctx.data.update(host=host, port=ctx.params.get("port"), security=ctx.params.get("security") or "")

    def _changes(self, ctx, acc) -> List[str]:
        out = []
        if ctx.data["host"] and ctx.data["host"].lower() != (acc.host or "").lower():
            out.append(f"сервер {acc.host} → {ctx.data['host']}")
        if ctx.data["port"] and int(ctx.data["port"]) != int(acc.port or 0):
            out.append(f"порт {acc.port} → {ctx.data['port']}")
        if ctx.data["security"] and ctx.data["security"] != acc.security:
            out.append(f"шифрование {acc.security} → {ctx.data['security']}")
        return out

    def check(self, ctx, acc):
        return "" if self._changes(ctx, acc) else "уже так"

    def plan(self, ctx, acc):
        return "; ".join(self._changes(ctx, acc))

    def apply(self, ctx, acc):
        text = self.plan(ctx, acc)
        # Меняем только эти поля: запись всей карточки из снимка затёрла бы то,
        # что успело измениться с момента проверки (например, выключение ящика).
        ctx.db.execute(
            "UPDATE accounts SET host=?, port=?, security=?, login_status='', login_error='', updated_at=? "
            "WHERE id=?",
            (ctx.data["host"] or acc.host, int(ctx.data["port"] or acc.port), ctx.data["security"] or acc.security,
             _now_iso(), acc.id))
        return text, None


class SetAuthAction(Handler):
    key = "set_auth"
    label = "Способ входа в ящик"
    group = "connect"
    icon = "🔑"
    danger = 1
    desc = ("Переключить ящики на вход через учётную запись администратора почты (пароли ящиков не нужны) "
            "или обратно — на пароли самих ящиков. Сначала проверьте вход администратора в «Настройках».")
    params = (
        Param("to", "Входить", "select", AuthType.MASTER,
              options=((AuthType.MASTER, AUTH_LABELS[AuthType.MASTER]),
                       (AuthType.PASSWORD, AUTH_LABELS[AuthType.PASSWORD]))),
    )

    def prepare(self, ctx):
        if ctx.params.get("to") == AuthType.MASTER:
            master = ctx.svc.master_credentials()
            if master is None:
                raise ValidationError("Вход через администратора почты выключен.",
                                      hint="Включите и проверьте его в «Настройки → Вход через администратора почты».")
            ctx.data["master_host"] = (master.get("host") or "").strip().lower()

    def check(self, ctx, acc):
        to = ctx.params.get("to")
        if acc.auth_type == to:
            return "уже так"
        if acc.auth_type == AuthType.OAUTH2:
            return "вход OAuth2 меняется в карточке ящика"
        host = ctx.data.get("master_host")
        if to == AuthType.MASTER and host and (acc.host or "").strip().lower() != host:
            return f"ящик на другом сервере ({acc.host}), а администратор почты настроен для {host}"
        return ""

    def plan(self, ctx, acc):
        to = ctx.params.get("to")
        text = f"{AUTH_LABELS.get(acc.auth_type, acc.auth_type)} → {AUTH_LABELS.get(to, to)}"
        if to == AuthType.PASSWORD and not acc.password:
            text += " (пароль ящика не задан — загрузите пароли)"
        return text

    def apply(self, ctx, acc):
        ctx.db.execute("UPDATE accounts SET auth_type=?, login_status='', login_error='', updated_at=? WHERE id=?",
                       (ctx.params.get("to"), datetime.now(timezone.utc).isoformat(), acc.id))
        return self.plan(ctx, acc), None


class FoldersCheckAction(Handler):
    key = "folders_check"
    label = "Проверить папки на сервере"
    group = "connect"
    icon = "🩺"
    desc = ("Одно задание по очереди открывает каждую папку выбранных ящиков на почтовом сервере (письма не "
            "скачиваются) и показывает, какие папки не открываются и сколько в них писем, — то есть из-за чего "
            "копия ящика неполная. То же, что «Проверить папки» в карточке ящика, только сразу для многих.")
    creates_jobs = True

    def prepare(self, ctx):
        running = next((j for j in ctx.db.active_jobs() if j["type"] == JobType.FOLDERS_CHECK), None)
        ctx.data["running"] = running["id"] if running is not None else None

    def check(self, ctx, acc):
        if ctx.data.get("running"):
            return f"проверка папок уже идёт (задание №{ctx.data['running']})"
        if not (acc.host or "").strip():
            return "не указан IMAP-сервер"
        return credential_problem(acc)

    def plan(self, ctx, acc):
        return "папки будут проверены на сервере"

    def apply(self, ctx, acc):
        ctx.data.setdefault("ids", []).append(acc.id)
        return "в проверке папок", None

    def finish(self, ctx):
        ids = ctx.data.get("ids") or []
        if ctx.preview or not ids:
            return
        job_id = ctx.enqueue(JobType.FOLDERS_CHECK, None, {"account_ids": sorted(ids)}, priority=4)
        for item in ctx.results:
            if item["status"] == OK:
                item["job_id"] = job_id


class FoldersAction(Handler):
    key = "folders"
    label = "Какие папки копировать"
    group = "connect"
    icon = "🗂️"
    desc = ("Добавить папки в «Пропускать папки» (например, «Спам» или «Корзина»), убрать их оттуда или "
            "ограничить копирование списком папок. Можно шаблоны: «Архив/*».")
    params = (
        Param("op", "Что сделать", "select", "exclude_add",
              options=(("exclude_add", "добавить в «Пропускать папки»"),
                       ("exclude_remove", "убрать из «Пропускать папки»"),
                       ("exclude_set", "заменить список «Пропускать папки»"),
                       ("include_set", "копировать только эти папки"),
                       ("include_clear", "копировать все папки (снять ограничение «только эти»)"))),
        Param("folders", "Папки — по одной в строке или через запятую", "textarea", "",
              placeholder="Спам\nКорзина", show_if={"op": ["exclude_add", "exclude_remove", "exclude_set", "include_set"]}),
    )

    def prepare(self, ctx):
        op = ctx.params.get("op")
        folders = _split_list(ctx.params.get("folders") or "")
        if op in ("exclude_add", "exclude_remove", "include_set") and not folders:
            raise ValidationError("Укажите хотя бы одну папку.")
        if len(folders) > 1000:
            raise ValidationError("Слишком длинный список папок (больше 1000).")
        ctx.data["folders"] = folders

    def _new(self, ctx, acc) -> Tuple[List[str], List[str]]:
        op = ctx.params.get("op")
        folders = ctx.data["folders"]
        include, exclude = list(acc.folder_include or []), list(acc.folder_exclude or [])
        if op == "exclude_add":
            exclude = exclude + [f for f in folders if f not in exclude]
        elif op == "exclude_remove":
            exclude = [f for f in exclude if f not in folders]
        elif op == "exclude_set":
            exclude = list(folders)
        elif op == "include_set":
            include = list(folders)
        elif op == "include_clear":
            include = []
        return include, exclude

    def check(self, ctx, acc):
        include, exclude = self._new(ctx, acc)
        if include == list(acc.folder_include or []) and exclude == list(acc.folder_exclude or []):
            return "без изменений"
        ctx.per[acc.id] = (include, exclude)
        return ""

    def plan(self, ctx, acc):
        include, exclude = ctx.per[acc.id]
        parts = []
        if exclude != list(acc.folder_exclude or []):
            parts.append("пропускать: " + (", ".join(exclude) or "ничего"))
        if include != list(acc.folder_include or []):
            parts.append("только: " + (", ".join(include) or "все папки"))
        return "; ".join(parts)[:300]

    def apply(self, ctx, acc):
        text = self.plan(ctx, acc)
        include, exclude = ctx.per[acc.id]
        ctx.db.execute("UPDATE accounts SET folder_include=?, folder_exclude=?, updated_at=? WHERE id=?",
                       (json.dumps(include, ensure_ascii=False), json.dumps(exclude, ensure_ascii=False),
                        _now_iso(), acc.id))
        return text, None


class CopySettingsAction(Handler):
    key = "copy_settings"
    label = "Настройки как у ящика-образца"
    group = "connect"
    icon = "🧬"
    danger = 1
    desc = ("Скопировать выбранным ящикам настройки одного ящика-образца: сервер и шифрование, способ входа, "
            "какие папки копировать, срок хранения, расписания копирования. Пароли не копируются.")

    def params_for(self, svc):
        options: Tuple[Tuple[Any, str], ...] = ()
        if svc is not None:
            options = tuple((a.id, f"{a.name} — {a.username}") for a in svc.db.list_accounts())
        return (
            Param("source", "Ящик-образец", "select", options[0][0] if options else "", options=options),
            Param("copy_server", "Сервер, порт и шифрование", "bool", True),
            Param("copy_auth", "Способ входа (пароль / через администратора почты)", "bool", False),
            Param("copy_folders", "Какие папки копировать и какие пропускать", "bool", True),
            Param("copy_retention", "Срок хранения копий", "bool", True),
            Param("copy_schedules", "Расписания копирования (заменят свои)", "bool", False),
        )

    def prepare(self, ctx):
        src = ctx.db.get_account(int(ctx.params.get("source") or 0))
        if src is None:
            raise ValidationError("Выберите ящик-образец.")
        if not any(ctx.params.get(k) for k in ("copy_server", "copy_auth", "copy_folders", "copy_retention",
                                                "copy_schedules")):
            raise ValidationError("Отметьте, какие настройки копировать.")
        ctx.data["src"] = src
        ctx.data["src_schedules"] = [s for s in ctx.db.list_schedules(src.id) if s["job_type"] == JobType.BACKUP]

    def _changes(self, ctx, acc) -> List[str]:
        src = ctx.data["src"]
        out = []
        if ctx.params.get("copy_server") and (acc.host, acc.port, acc.security) != (src.host, src.port, src.security):
            out.append(f"сервер {acc.host}:{acc.port} {acc.security} → {src.host}:{src.port} {src.security}")
        if ctx.params.get("copy_auth") and acc.auth_type != src.auth_type and AuthType.OAUTH2 not in (acc.auth_type, src.auth_type):
            out.append(f"вход {AUTH_LABELS.get(acc.auth_type, acc.auth_type)} → {AUTH_LABELS.get(src.auth_type, src.auth_type)}")
        if ctx.params.get("copy_folders") and (list(acc.folder_include or []), list(acc.folder_exclude or [])) != \
                (list(src.folder_include or []), list(src.folder_exclude or [])):
            out.append("папки: пропускать " + (", ".join(src.folder_exclude or []) or "ничего")
                       + (f"; только {', '.join(src.folder_include)}" if src.folder_include else ""))
        if ctx.params.get("copy_retention") and int(acc.retention_days if acc.retention_days is not None else -1) != \
                int(src.retention_days if src.retention_days is not None else -1):
            out.append(f"срок хранения {retention_label(acc.retention_days)} → {retention_label(src.retention_days)}")
        if ctx.params.get("copy_schedules"):
            mine = sorted((s["kind"], s["cron_expr"], s["interval_seconds"], bool(s["enabled"]))
                          for s in ctx.schedules.get(acc.id, []) if s["job_type"] == JobType.BACKUP)
            theirs = sorted((s["kind"], s["cron_expr"], s["interval_seconds"], bool(s["enabled"]))
                            for s in ctx.data["src_schedules"])
            if mine != theirs:
                out.append(f"расписаний копирования: {len(mine)} → {len(theirs)}")
        return out

    def check(self, ctx, acc):
        if acc.id == ctx.data["src"].id:
            return "это и есть ящик-образец"
        changes = self._changes(ctx, acc)
        if not changes:
            return "настройки уже как у образца"
        ctx.per[acc.id] = changes
        return ""

    def plan(self, ctx, acc):
        return "; ".join(ctx.per[acc.id])[:400]

    def apply(self, ctx, acc):
        src = ctx.data["src"]
        text = self.plan(ctx, acc)
        sets: List[str] = []
        values: List[Any] = []
        if ctx.params.get("copy_server"):
            sets += ["host=?", "port=?", "security=?"]
            values += [src.host, src.port, src.security]
        if ctx.params.get("copy_auth") and AuthType.OAUTH2 not in (acc.auth_type, src.auth_type):
            sets.append("auth_type=?")
            values.append(src.auth_type)
        if ctx.params.get("copy_server") or ctx.params.get("copy_auth"):
            sets += ["login_status=''", "login_error=''"]
        if ctx.params.get("copy_folders"):
            sets += ["folder_include=?", "folder_exclude=?"]
            values += [json.dumps(list(src.folder_include or []), ensure_ascii=False),
                       json.dumps(list(src.folder_exclude or []), ensure_ascii=False)]
        before = ctx.db.get_account(acc.id)
        retention_changed = ctx.params.get("copy_retention") and acc.retention_days != src.retention_days
        if ctx.params.get("copy_retention"):
            sets.append("retention_days=?")
            values.append(src.retention_days)
        if sets:
            ctx.db.execute(f"UPDATE accounts SET {', '.join(sets)}, updated_at=? WHERE id=?",
                           (*values, _now_iso(), acc.id))
        if retention_changed and before is not None:
            after_account_retention_change(ctx.svc, before)
        if ctx.params.get("copy_schedules"):
            for sched in ctx.schedules.get(acc.id, []):
                if sched["job_type"] == JobType.BACKUP:
                    ctx.db.delete_schedule(int(sched["id"]))
            for sched in ctx.data["src_schedules"]:
                ctx.db.create_schedule(acc.id, sched["kind"], JobType.BACKUP, sched["cron_expr"] or "",
                                       int(sched["interval_seconds"] or 0), bool(sched["enabled"]), {})
            ctx.data["reload"] = True
        return text, None

    def finish(self, ctx):
        if not ctx.preview and ctx.data.get("reload"):
            _reload_scheduler(ctx)


# =====================================================================
#  Выгрузка и восстановление
# =====================================================================
class ExportAction(Handler):
    key = "export"
    label = "Экспорт (PST / EML / MBOX)"
    group = "data"
    icon = "📤"
    desc = "Выгрузить копии ящиков — по файлу на ящик. Готовые файлы появятся в разделе «Экспорт»."
    creates_jobs = True

    def params_for(self, svc):
        from .export import list_engines
        engines = [(e["name"], e["title"] + ("" if e["available"] else " — недоступен"))
                   for e in list_engines() if e["available"]]
        default = str(svc.rt("export", "default_engine") or "eml") if svc is not None else "eml"
        if default not in [name for name, _t in engines]:
            default = engines[0][0] if engines else "eml"
        return (
            Param("engine", "Формат / движок", "select", default, options=tuple(engines)),
            Param("pst_format", "Формат PST", "select", "unicode",
                  options=(("unicode", "Unicode — Outlook 2003 и новее"), ("ansi", "ANSI — Outlook 97–2002 (до 2 ГБ)")),
                  show_if={"engine": ["aspose", "native"]}),
            Param("date_from", "Письма с даты", "date", ""),
            Param("date_to", "Письма по дату", "date", ""),
            Param("folders", "Только папки (пусто — все)", "textarea", "", placeholder="INBOX\nОтправленные"),
        )

    def prepare(self, ctx):
        from .export import list_engines
        engines = {e["name"]: e for e in list_engines()}
        engine = engines.get(ctx.params.get("engine"))
        if engine is None or not engine["available"]:
            raise ValidationError("Выбранный движок экспорта недоступен.")
        df, dt = ctx.params.get("date_from") or None, ctx.params.get("date_to") or None
        if df or dt:
            ctx.svc.day_bounds_utc(df, dt)       # проверка формата и порядка дат
        ctx.data.update(engine=engine["name"], fmt=engine["fmt"], date_from=df, date_to=dt,
                        folders=_split_list(ctx.params.get("folders") or "") or None)

    def check(self, ctx, acc):
        if ctx.busy(acc.id, (JobType.EXPORT,)):
            return "выгрузка уже идёт или стоит в очереди"
        if not ctx.messages(acc.id):
            return "в копии нет писем"
        d = ctx.data
        count = ctx.svc.count_mail_items(acc.id, d["folders"], d["date_from"], d["date_to"])
        if not count:
            return "нет писем под выбранные папки и даты"
        ctx.per[acc.id] = count
        return ""

    def validate_plan(self, ctx, accounts):
        from .util import disk_free_bytes
        need = sum(ctx.size(a.id) for a in accounts)
        try:
            free = disk_free_bytes(ctx.svc.cfg.exports_dir)
        except Exception:  # noqa: BLE001
            free = 0
        ctx.extra["disk"] = {"need": need, "need_h": human_size(need), "free": free,
                             "free_h": human_size(free) if free else "неизвестно"}
        if free and need > free:
            ctx.warnings.append(f"Выгрузки могут занять до {human_size(need)}, а свободно {human_size(free)}. "
                                f"Выгружайте частями и скачивайте готовые файлы.")

    def plan(self, ctx, acc):
        return f"будет выгружено писем: {ctx.per[acc.id]} ({ctx.data['engine']})"

    def apply(self, ctx, acc):
        d = ctx.data
        params = {"engine": d["engine"], "format": d["fmt"], "folders": d["folders"],
                  "date_from": d["date_from"], "date_to": d["date_to"], "limit": 0}
        if d["fmt"] == "pst":
            params["pst_format"] = ctx.params.get("pst_format") or "unicode"
        job_id = ctx.enqueue(JobType.EXPORT, acc.id, params)
        ctx.db.add_audit(ctx.username, "export_start", f"account={acc.id} engine={d['engine']} — групповое действие")
        return "выгрузка в очереди", job_id


class RestoreAction(Handler):
    key = "restore"
    label = "Восстановить на сервер"
    group = "data"
    icon = "♻️"
    danger = 1
    desc = ("Залить письма из копии обратно на почтовый сервер — например, после переезда на новый сервер "
            "(сначала «Сменить сервер IMAP»). Безопаснее — в папки с префиксом и с пробным прогоном.")
    creates_jobs = True
    params = (
        Param("target_mode", "Куда восстанавливать", "select", "prefixed",
              options=(("prefixed", "в папки с префиксом (безопасно)"), ("single", "всё в одну папку"),
                       ("original", "в исходные папки"))),
        Param("target_prefix", "Префикс папок", "text", "Восстановлено", show_if={"target_mode": ["prefixed"]}),
        Param("target_folder", "Имя папки", "text", "Восстановлено", show_if={"target_mode": ["single"]}),
        Param("folders", "Только папки (пусто — все)", "textarea", "", placeholder="INBOX"),
        Param("check_duplicates", "Пропускать письма, которые уже есть на сервере", "bool", True),
        Param("dry_run", "Пробный прогон (ничего не заливать, только посчитать)", "bool", True),
    )

    def danger_for(self, ctx):
        opts = ctx.data.get("opts") or {}
        return 2 if opts.get("target_mode") == "original" and not opts.get("dry_run") else 1

    def prepare(self, ctx):
        opts = validate_restore_options({
            "target_mode": ctx.params.get("target_mode"), "target_prefix": ctx.params.get("target_prefix"),
            "target_folder": ctx.params.get("target_folder"),
            "folders": _split_list(ctx.params.get("folders") or "") or None,
            "check_duplicates": bool(ctx.params.get("check_duplicates")), "dry_run": bool(ctx.params.get("dry_run")),
            "limit": 0})
        ctx.data["opts"] = opts
        if opts["target_mode"] == "original" and not opts["dry_run"]:
            ctx.warnings.append("Письма будут залиты прямо в рабочие папки ящиков. Отменить это нельзя.")

    def check(self, ctx, acc):
        if not ctx.messages(acc.id):
            return "в копии нет писем"
        problem = credential_problem(acc)
        if problem:
            return problem
        if ctx.busy(acc.id, (JobType.RESTORE,)):
            return "восстановление уже идёт или стоит в очереди"
        return ""

    def plan(self, ctx, acc):
        opts = ctx.data["opts"]
        where = {"prefixed": f"в папки «{opts['target_prefix']}/…»", "single": f"в папку «{opts['target_folder']}»",
                 "original": "в исходные папки"}[opts["target_mode"]]
        return (("пробный прогон: " if opts["dry_run"] else "") + f"{ctx.messages(acc.id)} писем {where}"
                + f" на {acc.host}")

    def apply(self, ctx, acc):
        opts = dict(ctx.data["opts"])
        job_id = ctx.enqueue(JobType.RESTORE, acc.id, opts)
        if opts["target_mode"] == "original" and not opts["dry_run"]:
            ctx.db.add_audit(ctx.username, "restore_original", f"account={acc.id} — групповое действие")
        return ("пробный прогон" if opts["dry_run"] else "восстановление") + " в очереди", job_id


# =====================================================================
#  Хранилище и удаление
# =====================================================================
class SearchReindexAction(Handler):
    key = "search_reindex"
    label = "Переиндексировать поиск"
    group = "storage"
    icon = "🔎"
    desc = ("Заново прочитать письма выбранных ящиков и обновить их в индексе поиска: тема, адреса, имена "
            "вложений и текст. Нужно, если поиск не находит письма ящика — например, текст писем начали "
            "индексировать уже после первого прохода или письма не читались без ключа шифрования. Остальной "
            "индекс не трогается; идёт фоновым заданием.")
    creates_jobs = True

    def prepare(self, ctx):
        from . import search as search_mod
        if not search_mod.fts_available(ctx.db):
            raise ValidationError("Полнотекстовый поиск недоступен: SQLite собран без FTS5 — индексировать нечего.")
        running = next((j for j in ctx.db.active_jobs() if j["type"] == JobType.SEARCH_REINDEX), None)
        if running is not None:
            ctx.warnings.append(f"Переиндексация уже идёт (задание №{running['id']}) — новая начнётся после неё.")
        if not search_mod.bodies_enabled(ctx.svc):
            ctx.warnings.append("Текст писем сейчас не индексируется (так настроено в «Настройки → Поиск») — "
                                "обновятся только тема, адреса и имена вложений.")

    def check(self, ctx, acc):
        if not ctx.messages(acc.id):
            return "в копии нет писем"
        return ""

    def plan(self, ctx, acc):
        return f"будет переиндексировано писем: {ctx.messages(acc.id)}"

    def apply(self, ctx, acc):
        ctx.data.setdefault("ids", []).append(acc.id)
        return "в переиндексации", None

    def finish(self, ctx):
        ids = ctx.data.get("ids") or []
        if ctx.preview or not ids:
            return
        job_id = ctx.enqueue(JobType.SEARCH_REINDEX, None, {"account_ids": sorted(ids)}, priority=7)
        for item in ctx.results:
            if item["status"] == OK:
                item["job_id"] = job_id


class QuarantineCheckAction(Handler):
    key = "quarantine_check"
    label = "Сравнить прежние копии с новыми"
    group = "storage"
    icon = "🔍"
    desc = ("Для каждой прежней копии (карантина после копии «с нуля») выяснить, какие её письма есть в новой "
            "копии, а какие остались только в ней (их уже нет на сервере). Ничего не меняет; идёт фоновыми "
            "заданиями, итог — в «Прежних копиях» ящика и в проверке перед удалением.")
    creates_jobs = True

    def check(self, ctx, acc):
        paths = ctx.svc.store.quarantine_paths(acc.id)
        if not paths:
            return "прежних копий нет"
        if ctx.busy(acc.id, (JobType.QUARANTINE_CHECK, JobType.QUARANTINE_RESCUE)):
            return "сравнение или возврат писем уже идёт"
        ctx.per[acc.id] = paths
        return ""

    def plan(self, ctx, acc):
        from . import quarantine as qmod
        parts = [f"«{os.path.basename(p)}»: {qmod.check_label(qmod.stored_check(ctx.db, p))}"
                 for p in ctx.per[acc.id]]
        return "будет сравнено — " + "; ".join(parts)

    def apply(self, ctx, acc):
        job_id = ctx.enqueue(JobType.QUARANTINE_CHECK, acc.id, {"paths": list(ctx.per[acc.id])}, priority=6)
        return "сравнение поставлено в очередь", job_id


class QuarantineRescueAction(Handler):
    key = "quarantine_rescue"
    label = "Вернуть письма из прежних копий"
    group = "storage"
    icon = "🛟"
    danger = 1
    desc = ("Письма, которые есть только в прежней копии (на сервере их уже нет), вернуть в архив ящика — в те "
            "же папки. После этого прежнюю копию можно удалять без потерь. Письма старше срока хранения ящика "
            "не возвращаются: их удалила бы ночная очистка.")
    note = ("Возврат — только после того, как копирование хотя бы раз прочитало все папки ящика: иначе письма ещё "
            "не прочитанных папок попали бы в архив дважды.")
    creates_jobs = True

    def check(self, ctx, acc):
        from . import quarantine as qmod
        paths = ctx.svc.store.quarantine_paths(acc.id)
        if not paths:
            return "прежних копий нет"
        if ctx.busy(acc.id):
            return "по ящику выполняется задание — дождитесь его окончания"
        checks = [qmod.stored_check(ctx.db, p) for p in paths]
        if all(c is not None and not c.get("unique") for c in checks):
            return "все письма прежних копий уже есть в новой"
        if not any(qmod.new_copy_complete(ctx.svc, acc.id, p) for p in paths):
            return "после копии «с нуля» копирование ещё не прочитало все папки ящика"
        if all(c is not None and int(c.get("unique_unreadable") or 0) >= int(c.get("unique") or 0) for c in checks):
            return "письма есть только в папках, которые сейчас не открываются на сервере"
        items = [(p, 0, 0) for p in paths]
        ctx.per[acc.id] = (items, checks)
        return ""

    def plan(self, ctx, acc):
        items, checks = ctx.per[acc.id]
        unique = sum(int(c.get("unique") or 0) for c in checks if c)
        unchecked = sum(1 for c in checks if c is None)
        text = f"будет возвращено писем: {unique}" if unique else "письма будут найдены сравнением"
        if unchecked:
            text += f" (прежних копий без сравнения: {unchecked} — сравнятся перед возвратом)"
        return text

    def apply(self, ctx, acc):
        items, _checks = ctx.per[acc.id]
        job_id = ctx.enqueue(JobType.QUARANTINE_RESCUE, acc.id, {"paths": [p for p, _n, _s in items]}, priority=6)
        return "возврат писем поставлен в очередь", job_id


class QuarantineDeleteAction(Handler):
    key = "quarantine_delete"
    label = "Удалить прежние копии (карантин)"
    group = "storage"
    icon = "🧺"
    # Необратимо: в прежней копии могут быть письма, которых нет ни на сервере,
    # ни в новой копии, — подтверждение числом ящиков, как у удаления ящиков.
    danger = 2
    desc = ("После копии «с нуля» прежние письма лежат в карантине. Когда новая копия проверена, "
            "карантин можно удалить и освободить место. Удаление идёт фоновым заданием и необратимо.")
    note = ("Сначала «Сравнить прежние копии с новыми»: по умолчанию удаляются только прежние копии, все письма "
            "которых есть в новой копии.")
    creates_jobs = True
    params = (Param("only_safe", "Только прежние копии, все письма которых есть в новой копии (по сравнению)",
                    "bool", True,
                    help="Выключите, чтобы удалить и несравнённые прежние копии, и те, где есть письма, которых "
                         "нет в новой копии, — такие письма будут потеряны."),)

    def check(self, ctx, acc):
        from . import quarantine as qmod
        if acc.on_hold():
            return f"архив удерживается ({_hold_label(acc.hold_until)}) — прежние копии не удаляем"
        if ctx.busy(acc.id):
            # копия «с нуля» в очереди прямо сейчас создаёт НОВЫЙ карантин —
            # удалять что-либо, пока по ящику идёт работа, нельзя
            return "по ящику выполняется задание — дождитесь его окончания"
        paths = ctx.svc.store.quarantine_paths(acc.id)
        if not paths:
            return "прежних копий нет"
        checked = []
        reasons = []
        from .queue.jobs import effective_retention_days
        days = effective_retention_days(ctx.svc, acc)
        for path in paths:
            check = qmod.stored_check(ctx.db, path)
            safe, why = qmod.is_safe_to_delete(check, days)
            if safe or not ctx.params.get("only_safe"):
                # Размер берём из итога сравнения: обходить прежнюю копию в сотни
                # тысяч файлов ради предпросмотра — минуты на каждый ящик.
                if check and check.get("disk_bytes") is not None:
                    n, size = int(check.get("files") or 0), int(check.get("disk_bytes") or 0)
                else:
                    n, size = ctx.svc.store.quarantine_usage(path)
                checked.append((path, n, size, safe, why))
            else:
                reasons.append(why)
        if not checked:
            return reasons[0]
        ctx.per[acc.id] = checked
        ctx.data.setdefault("kept", {})[acc.id] = reasons
        return ""

    def validate_plan(self, ctx, accounts):
        total = sum(size for a in accounts for _p, _n, size, _ok, _w in ctx.per.get(a.id, []))
        ctx.extra["freed"] = {"bytes": total, "h": human_size(total)}
        risky = sum(1 for a in accounts for _p, _n, _s, ok, _w in ctx.per.get(a.id, []) if not ok)
        if risky:
            ctx.warnings.append(f"Прежних копий, которые не сравнивались с новой или содержат письма, которых в ней "
                                f"нет: {risky}. Такие письма будут потеряны.")

    def plan(self, ctx, acc):
        items = ctx.per[acc.id]
        text = f"каталогов: {len(items)}, файлов: {sum(n for _p, n, _s, _o, _w in items)}, " \
               f"{human_size(sum(s for _p, _n, s, _o, _w in items))}"
        risky = [w for _p, _n, _s, ok, w in items if not ok]
        if risky:
            text += "; ВНИМАНИЕ: " + "; ".join(risky)
        kept = (ctx.data.get("kept") or {}).get(acc.id) or []
        if kept:
            text += f"; не удаляются ({len(kept)}): " + "; ".join(kept)
        return text

    def apply(self, ctx, acc):
        # Удаляются ровно те каталоги, что были показаны при проверке: карантин,
        # появившийся позже (новая копия «с нуля»), задание не тронет.
        ctx.data.setdefault("items", []).append({"account_id": acc.id, "name": acc.name, "what": "quarantine",
                                                 "paths": [path for path, _n, _s, _o, _w in ctx.per[acc.id]],
                                                 "only_safe": bool(ctx.params.get("only_safe"))})
        return "удаление в очереди", None

    def finish(self, ctx):
        _enqueue_cleanup(ctx)


def _enqueue_cleanup(ctx) -> None:
    items = ctx.data.get("items") or []
    if ctx.preview or not items:
        return
    job_id = ctx.enqueue(JobType.CLEANUP, None, {"items": items}, priority=6)
    for item in ctx.results:
        if item["status"] == OK:
            item["job_id"] = job_id


class DeleteAction(Handler):
    key = "delete"
    label = "Удалить ящики (архив остаётся на диске)"
    group = "danger"
    icon = "🗑️"
    danger = 2
    desc = ("Удалить ящики и их настройки, расписания и индекс писем. Файлы писем на диске не трогаются "
            "(каталог account_<номер ящика> в хранилище) и занимают место, пока их не удалить вручную. "
            "Ящик, заведённый заново, начнёт архив с чистого листа — старые файлы он не подхватит.")

    def check(self, ctx, acc):
        if acc.on_hold():
            return f"архив удерживается ({_hold_label(acc.hold_until)}) — ящик не удаляем"
        if ctx.busy(acc.id):
            return "по ящику выполняется задание — отмените его или дождитесь окончания"
        return ""

    def plan(self, ctx, acc):
        n = ctx.messages(acc.id)
        return "ящик будет удалён" + (f"; файлы {n} писем ({human_size(ctx.size(acc.id))}) останутся на диске" if n else "")

    def apply(self, ctx, acc):
        ctx.db.delete_account(acc.id)
        ctx.db.add_audit(ctx.username, "account_delete", f"{acc.name} — групповое действие")
        return "ящик удалён", None


class PurgeAction(Handler):
    key = "purge"
    label = "Удалить ящики вместе с архивом"
    group = "danger"
    icon = "🔥"
    danger = 2
    desc = ("Удалить ящики, их индекс и ВСЕ файлы писем на диске. Необратимо: письма, которых уже нет на "
            "почтовом сервере, не вернуть. Удерживаемые архивы пропускаются.")
    creates_jobs = True
    params = (Param("with_quarantine", "Удалить и прежние копии (карантин)", "bool", True),)

    def check(self, ctx, acc):
        if acc.on_hold():
            return f"архив удерживается ({_hold_label(acc.hold_until)}) — удалить нельзя"
        if ctx.busy(acc.id):
            return "по ящику выполняется задание — отмените его или дождитесь окончания"
        return ""

    def plan(self, ctx, acc):
        n = ctx.messages(acc.id)
        return f"будут удалены ящик и {n} писем ({human_size(ctx.size(acc.id))})" if n else "будет удалён ящик (писем нет)"

    def apply(self, ctx, acc):
        # Выключаем сразу: пока задание удаления ждёт очереди, ящик не должен
        # начать копироваться по расписанию.
        ctx.db.set_account_enabled(acc.id, False)
        ctx.data.setdefault("items", []).append({"account_id": acc.id, "name": acc.name, "what": "purge",
                                                 "with_quarantine": bool(ctx.params.get("with_quarantine"))})
        return "ящик выключен, удаление в очереди", None

    def finish(self, ctx):
        _enqueue_cleanup(ctx)


# =====================================================================
#  Реестр и выполнение
# =====================================================================
HANDLERS: Dict[str, Handler] = {h.key: h for h in (
    BackupAction(), BackupSequenceAction(), RebuildMissingAction(), RebuildFullAction(), FinalBackupAction(),
    VerifyAction(), CheckLoginsAction(), RetentionRunAction(), CancelJobsAction(),
    EnableAction(), DisableAction(), RetentionSetAction(), HoldSetAction(), HoldClearAction(),
    LogoutSessionsAction(), RenameAction(), NotesAction(),
    ScheduleSetAction(), ScheduleRemoveAction(), ScheduleToggleAction(),
    SetServerAction(), SetAuthAction(), FoldersCheckAction(), FoldersAction(), CopySettingsAction(),
    ExportAction(), RestoreAction(),
    SearchReindexAction(), QuarantineCheckAction(), QuarantineRescueAction(), QuarantineDeleteAction(),
    DeleteAction(), PurgeAction(),
)}


def describe_actions(svc) -> Dict[str, Any]:
    return {"groups": [{"key": k, "label": t} for k, t in GROUPS],
            "actions": [h.describe(svc) for h in HANDLERS.values()],
            "create_defaults": create_defaults(svc)}


def create_defaults(svc) -> Dict[str, Any]:
    """Сервер для новых ящиков по умолчанию — тот, на котором больше всего уже заведённых."""
    from collections import Counter
    accounts = svc.db.list_accounts()
    hosts = Counter(((a.host or "").strip().lower(), a.port, a.security) for a in accounts if a.host)
    if hosts:
        (host, port, security), _n = hosts.most_common(1)[0]
    else:
        host, port, security = str(svc.rt("mailadmin", "host") or "").strip(), 993, Security.SSL
    has_global = any(s["job_type"] == JobType.BACKUP_ALL and s["enabled"] for s in svc.db.list_schedules())
    return {"host": host, "port": port, "security": security, "master": svc.master_credentials() is not None,
            "global_schedule": has_global}


def _clean_ids(ids: Optional[Sequence[Any]]) -> List[int]:
    out: List[int] = []
    seen = set()
    for raw in ids or []:
        try:
            value = int(raw)
        except (TypeError, ValueError):
            continue
        if value > 0 and value not in seen:
            seen.add(value)
            out.append(value)
    if len(out) > MAX_IDS:
        raise ValidationError(f"Слишком много ящиков за одну операцию (больше {MAX_IDS}).")
    return out


def _safe_params(params: Dict[str, Any]) -> Dict[str, Any]:
    """Параметры для истории: длинные тексты укорачиваем."""
    out = {}
    for key, value in (params or {}).items():
        if isinstance(value, str) and len(value) > 300:
            value = value[:300] + "…"
        out[key] = value
    return out


def run(svc, user: dict, action: str, ids: Sequence[Any], params: Optional[Dict[str, Any]] = None, *,
        preview: bool = True, confirm: str = "") -> Dict[str, Any]:
    """Выполнить (или просчитать) групповое действие над ящиками ``ids``."""
    handler = HANDLERS.get(action or "")
    if handler is None:
        raise ValidationError(f"Неизвестное групповое действие «{action}».")
    ids = _clean_ids(ids)
    if not ids:
        raise ValidationError("Не выбрано ни одного ящика.", hint="Отметьте ящики в списке или выберите их отбором.")
    raw = dict(params or {})
    values: Dict[str, Any] = {}
    for param in handler.params_for(svc):
        values[param.key] = param.coerce(raw.get(param.key))
    ctx = BulkContext(svc, user, values, preview)
    handler.prepare(ctx)

    by_id = {a.id: a for a in svc.db.list_accounts()}
    to_apply: List[Account] = []
    for account_id in ids:
        acc = by_id.get(account_id)
        if acc is None:
            ctx.results.append({"id": account_id, "name": f"№{account_id}", "username": "", "status": FAIL,
                                "detail": "ящик не найден (удалён?)"})
            continue
        try:
            reason = handler.check(ctx, acc)
        except MailArchiverError as exc:
            reason = exc.message
        if reason:
            ctx.results.append({"id": acc.id, "name": acc.name, "username": acc.username, "status": SKIP,
                                "detail": reason})
        else:
            to_apply.append(acc)
    handler.validate_plan(ctx, to_apply)
    danger = handler.danger_for(ctx)
    if not preview and to_apply and danger >= 2 and str(confirm or "").strip() != str(len(to_apply)):
        raise ValidationError(f"Для подтверждения введите число ящиков, которые будут обработаны: {len(to_apply)}.",
                              hint="Число видно в окне проверки перед запуском.")
    op_id = None
    if not preview:
        op_id = svc.db.create_bulk_op(ctx.username, handler.key, handler.label, _safe_params(values), len(ids))
    for acc in to_apply:
        if preview:
            try:
                detail = handler.plan(ctx, acc)
                status = OK
            except MailArchiverError as exc:
                detail, status = exc.message, FAIL
            ctx.results.append({"id": acc.id, "name": acc.name, "username": acc.username, "status": status,
                                "detail": detail})
            continue
        try:
            detail, job_id = handler.apply(ctx, acc)
            item = {"id": acc.id, "name": acc.name, "username": acc.username, "status": OK, "detail": detail}
            if job_id:
                item["job_id"] = job_id
        except MailArchiverError as exc:
            item = {"id": acc.id, "name": acc.name, "username": acc.username, "status": FAIL,
                    "detail": exc.message + (f" {exc.hint}" if exc.hint else "")}
        except Exception as exc:  # noqa: BLE001
            log.exception("Групповое действие %s: ящик «%s» не обработан", handler.key, acc.name)
            item = {"id": acc.id, "name": acc.name, "username": acc.username, "status": FAIL,
                    "detail": f"{type(exc).__name__}: {exc}"}
        ctx.results.append(item)
    handler.finish(ctx)

    # порядок строк — как в запросе (пропуски и ошибки не уезжают в конец)
    order = {account_id: i for i, account_id in enumerate(ids)}
    ctx.results.sort(key=lambda item: order.get(item["id"], 0))
    counts = {OK: 0, SKIP: 0, FAIL: 0}
    for item in ctx.results:
        counts[item["status"]] = counts.get(item["status"], 0) + 1
    reasons: Dict[str, int] = {}
    for item in ctx.results:
        if item["status"] == SKIP:
            reasons[item["detail"]] = reasons.get(item["detail"], 0) + 1
    verb = "будет обработано" if preview else "обработано"
    summary = f"{handler.label}: {verb} {counts[OK]} из {len(ids)}"
    if counts[SKIP]:
        summary += f", пропущено {counts[SKIP]}"
    if counts[FAIL]:
        summary += f", ошибок {counts[FAIL]}"
    summary += "."
    if ctx.jobs:
        summary += f" Заданий в очереди: {len(set(ctx.jobs))}."
    if op_id is not None:
        svc.db.finish_bulk_op(op_id, ok=counts[OK], skipped=counts[SKIP], failed=counts[FAIL],
                              jobs=len(set(ctx.jobs)), summary=summary, results=ctx.results)
        names = ", ".join(item["name"] for item in ctx.results if item["status"] == OK)
        svc.db.add_audit(ctx.username, f"bulk_{handler.key}",
                         f"операция №{op_id}: {summary} {names}"[:900])
    return {"ok": True, "preview": preview, "op_id": op_id, "action": handler.key, "label": handler.label,
            "danger": danger, "confirm_value": str(counts[OK]) if danger >= 2 else "",
            "total": len(ids), "counts": counts, "reasons": reasons, "warnings": ctx.warnings,
            "summary": summary, "jobs": sorted(set(ctx.jobs)), "extra": ctx.extra, "results": ctx.results}


# =====================================================================
#  Добавление ящиков списком
# =====================================================================
#: Сколько строк списка принимаем за раз.
MAX_CREATE_LINES = 5000
_ANGLE_EMAIL = re.compile(r"<([^<>\s]+@[^<>\s]+)>")


def parse_create_lines(text: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Строки «адрес[;пароль]» (или «Имя <адрес>;пароль», или две колонки из Excel через табуляцию).

    Пароль необязателен: ящик без пароля заводится выключенным, пароль можно
    загрузить потом файлом. Возвращает ``(строки, проблемы)``.
    """
    from .employees import looks_like_email, normalize_email
    rows: List[Dict[str, Any]] = []
    problems: List[Dict[str, Any]] = []
    lines = (text or "").splitlines()
    if len(lines) > MAX_CREATE_LINES:
        raise ValidationError(f"Слишком длинный список: больше {MAX_CREATE_LINES} строк.",
                              hint="Добавляйте ящики частями.")
    for number, raw in enumerate(lines, start=1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        name = ""
        raw_line = raw.rstrip("\r\n")
        if "\t" in raw_line:
            # Колонки из Excel: адрес, пароль и (необязательно) ФИО. Пароль берём
            # как есть — пробелы в нём могут быть частью пароля.
            cols = raw_line.split("\t")
            who, password = cols[0].strip(), (cols[1] if len(cols) > 1 else "")
            if len(cols) > 2:
                name = cols[2].strip()
        else:
            parts = re.split(r";", line, maxsplit=1)
            if len(parts) == 1 and "," in line:
                parts = line.split(",", 1)
            who, password = parts[0].strip(), (parts[1].strip() if len(parts) > 1 else "")
        m = _ANGLE_EMAIL.search(who)
        if m:
            name = who[:m.start()].strip().strip('"') or name
            who = m.group(1)
        email = normalize_email(who)
        if not looks_like_email(email):
            if number == 1 and any(w in line.lower() for w in ("email", "почта", "адрес", "логин")):
                continue                                  # строка заголовков
            # Содержимое строки не показываем: вместо адреса там мог оказаться
            # пароль (перепутаны колонки, разделитель — пробел).
            problems.append({"row": number, "reason": "первое поле не похоже на адрес почты — строка пропущена"})
            continue
        rows.append({"row": number, "email": email, "password": password, "name": name})
    return rows, problems


def bulk_create(svc, user: dict, text: str, *, host: str, port: int = 993, security: str = Security.SSL,
                auth_type: str = AuthType.PASSWORD, enable: bool = True, schedule_time: str = "",
                preview: bool = True) -> Dict[str, Any]:
    """Завести много ящиков сразу из списка адресов (с паролями или без).

    Название ящика — ФИО сотрудника из справочника (если он там есть), иначе
    имя из строки «Имя <адрес>», иначе сам адрес. Ящик сразу связывается с
    карточкой сотрудника. Уже заведённые адреса пропускаются. Пароли никуда,
    кроме зашифрованного поля ящика, не попадают — ни в историю, ни в аудит.
    """
    username = (user or {}).get("username") or "admin"
    host = (host or "").strip()
    if not host or len(host) > 255 or re.search(r"\s", host):
        raise ValidationError("Укажите адрес IMAP-сервера для новых ящиков.")
    if not 1 <= int(port or 0) <= 65535:
        raise ValidationError("Порт IMAP — число от 1 до 65535.")
    if security not in Security.ALL:
        raise ValidationError("Некорректный режим шифрования.")
    if auth_type not in (AuthType.PASSWORD, AuthType.MASTER):
        raise ValidationError("Способ входа — по паролю ящика или через администратора почты.")
    if auth_type == AuthType.MASTER and svc.master_credentials() is None:
        raise ValidationError("Вход через администратора почты выключен.",
                              hint="Включите его в «Настройки → Вход через администратора почты».")
    cron = ""
    if schedule_time:
        hh, mm = _parse_hhmm(schedule_time, "Время расписания")
        cron = f"{mm} {hh} * * *"
    rows, problems = parse_create_lines(text)
    if not rows and not problems:
        raise ValidationError("Список пуст.", hint="По одному адресу в строке; пароль — через «;» или табуляцию.")
    existing = {(a.username or "").lower(): a for a in svc.db.list_accounts()}
    employees = {}
    for row in svc.db.query("SELECT id, full_name, email, account_id FROM employees WHERE email IS NOT NULL AND email<>''"):
        employees.setdefault((row["email"] or "").lower(), row)
    results: List[Dict[str, Any]] = []
    seen = set()
    for item in rows:
        email = item["email"]
        if email in seen:
            results.append({"id": 0, "name": email, "username": email, "status": SKIP, "detail": "повтор в списке"})
            continue
        seen.add(email)
        if email in existing:
            results.append({"id": existing[email].id, "name": existing[email].name, "username": email,
                            "status": SKIP, "detail": "ящик уже есть"})
            continue
        emp = employees.get(email)
        name = (emp["full_name"] if emp else "") or item["name"] or email
        has_secret = bool(item["password"]) or auth_type == AuthType.MASTER
        on = bool(enable and has_secret)
        detail = ("будет заведён" if preview else "заведён") + (" и включён" if on else ", выключен")
        if not has_secret:
            detail += " (без пароля — загрузите пароли файлом)"
        if emp is not None:
            detail += "; связан с сотрудником"
        if cron:
            detail += f"; расписание в {schedule_time}"
        entry = {"id": 0, "name": name[:200], "username": email, "status": OK, "detail": detail}
        if not preview:
            acc_id = svc.db.create_account(Account(
                name=name[:200], host=host, port=int(port), username=email, password=item["password"],
                auth_type=auth_type, security=security, enabled=on))
            entry["id"] = acc_id
            if emp is not None and not emp["account_id"]:
                svc.db.set_employee_account(int(emp["id"]), acc_id)
            if cron:
                svc.db.create_schedule(acc_id, ScheduleKind.CRON, JobType.BACKUP, cron, 0, True, {})
        results.append(entry)
    for problem in problems:
        results.append({"id": 0, "name": f"строка {problem['row']}", "username": "", "status": FAIL,
                        "detail": problem["reason"]})
    counts = {OK: 0, SKIP: 0, FAIL: 0}
    for item in results:
        counts[item["status"]] += 1
    reasons: Dict[str, int] = {}
    for item in results:
        if item["status"] == SKIP:
            reasons[item["detail"]] = reasons.get(item["detail"], 0) + 1
    label = "Добавить ящики списком"
    summary = f"{label}: {'будет заведено' if preview else 'заведено'} {counts[OK]} из {len(rows)}"
    if counts[SKIP]:
        summary += f", пропущено {counts[SKIP]}"
    if counts[FAIL]:
        summary += f", непонятных строк {counts[FAIL]}"
    summary += "."
    op_id = None
    if not preview:
        params = {"host": host, "port": int(port), "security": security, "auth_type": auth_type,
                  "enable": bool(enable), "schedule_time": schedule_time}
        op_id = svc.db.create_bulk_op(username, "create", label, params, len(rows) + len(problems))
        svc.db.finish_bulk_op(op_id, ok=counts[OK], skipped=counts[SKIP], failed=counts[FAIL], jobs=0,
                              summary=summary, results=results)
        svc.db.add_audit(username, "accounts_bulk_create", f"операция №{op_id}: {summary}"[:900])
        if cron:
            _reload_scheduler(BulkContext(svc, user, {}, False))
    return {"ok": True, "preview": preview, "op_id": op_id, "action": "create", "label": label, "danger": 0,
            "confirm_value": "", "total": len(rows) + len(problems), "counts": counts, "reasons": reasons,
            "warnings": [], "summary": summary, "jobs": [], "extra": {}, "results": results}


# =====================================================================
#  История операций
# =====================================================================
def history(svc, limit: int = 50) -> List[Dict[str, Any]]:
    out = []
    for row in svc.db.list_bulk_ops(limit=max(1, min(300, int(limit)))):
        try:
            params = json.loads(row["params"] or "{}")
        except (TypeError, ValueError):
            params = {}
        out.append({"id": row["id"], "created_at": row["created_at"], "finished_at": row["finished_at"],
                    "user": row["user"], "action": row["action"], "label": row["label"], "params": params,
                    "total": row["total"], "ok": row["ok"], "skipped": row["skipped"], "failed": row["failed"],
                    "jobs": row["jobs"], "summary": row["summary"]})
    return out


def history_item(svc, op_id: int) -> Optional[Dict[str, Any]]:
    row = svc.db.get_bulk_op(op_id)
    if row is None:
        return None
    try:
        results = json.loads(row["results"] or "[]")
    except (TypeError, ValueError):
        results = []
    try:
        params = json.loads(row["params"] or "{}")
    except (TypeError, ValueError):
        params = {}
    job_ids = sorted({int(r["job_id"]) for r in results if r.get("job_id")})
    statuses = svc.db.job_statuses(job_ids)
    progress: Dict[str, int] = {}
    for job_id in job_ids:
        status = statuses.get(job_id, "deleted")
        progress[status] = progress.get(status, 0) + 1
    for item in results:
        if item.get("job_id"):
            item["job_status"] = statuses.get(int(item["job_id"]), "deleted")
    return {"id": row["id"], "created_at": row["created_at"], "finished_at": row["finished_at"],
            "user": row["user"], "action": row["action"], "label": row["label"], "params": params,
            "total": row["total"], "ok": row["ok"], "skipped": row["skipped"], "failed": row["failed"],
            "summary": row["summary"], "results": results, "job_progress": progress,
            "jobs_active": sum(progress.get(s, 0) for s in JobStatus.ACTIVE)}


def cancel_op_jobs(svc, op_id: int, only_operator: bool = False) -> int:
    """Отменить незавершённые задания групповой операции.

    ``only_operator`` — только задания, которые может снимать оператор.
    """
    item = history_item(svc, op_id)
    if item is None:
        raise ValidationError("Групповая операция не найдена.")
    cancelled = 0
    for job_id in sorted({int(r["job_id"]) for r in item["results"] if r.get("job_id")}):
        row = svc.db.get_job(job_id)
        if row is not None and only_operator and not roles.operator_may_manage(row):
            continue
        if row is not None and row["status"] in JobStatus.ACTIVE:
            svc.queue.cancel(job_id)
            cancelled += 1
    return cancelled


# =====================================================================
#  Выгрузка списка ящиков (Excel / CSV)
# =====================================================================
LIST_COLUMNS = (
    ("name", "Название"), ("username", "Логин (адрес)"), ("host", "Сервер"), ("port", "Порт"),
    ("security", "Шифрование"), ("auth", "Способ входа"), ("enabled", "Копирование"),
    ("login", "Проверка входа"), ("login_checked_at", "Вход проверен"), ("messages", "Писем в копии"),
    ("bytes", "Размер копии, байт"), ("size", "Размер копии"), ("first_backup_at", "Первая копия"),
    ("last_backup_at", "Последняя удачная копия"), ("last_backup_status", "Итог последней копии"),
    ("retention", "Срок хранения"), ("hold", "Удержание архива"), ("dismissed_at", "Уволен"),
    ("employee", "Сотрудник"), ("department", "Отдел"), ("position", "Должность"),
    ("schedules", "Расписаний"), ("notes", "Заметка"),
)
_LOGIN_LABELS = {"ok": "пароль верный", "auth_error": "неверный пароль", "conn_error": "нет связи",
                 "no_password": "нет пароля", "secret_broken": "пароль не читается"}
_STATUS_LABELS = {"success": "успешно", "partial": "частично", "failed": "ошибка", "cancelled": "отменено"}


def _local(svc, iso: Optional[str]) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso).astimezone(svc.local_tz()).strftime("%d.%m.%Y %H:%M")
    except (TypeError, ValueError):
        return str(iso)


def account_rows(svc, ids: Optional[Sequence[Any]] = None) -> List[Dict[str, Any]]:
    """Строки списка ящиков для выгрузки: всё, что видно в разделе, плюс сотрудник."""
    wanted = set(_clean_ids(ids)) if ids else None
    totals = svc.db.message_totals_by_account()
    employees = svc.db.employees_by_account()
    schedules = svc.db.schedules_by_account()
    rows = []
    for acc in svc.db.list_accounts():
        if wanted is not None and acc.id not in wanted:
            continue
        t = totals.get(acc.id) or {}
        emp = employees.get(acc.id)
        rows.append({
            "name": acc.name, "username": acc.username, "host": acc.host, "port": acc.port,
            "security": Security.LABELS.get(acc.security, acc.security),
            "auth": AUTH_LABELS.get(acc.auth_type, acc.auth_type),
            "enabled": "включено" if acc.enabled else "выключено",
            "login": _LOGIN_LABELS.get(acc.login_status, "не проверялся" if not acc.login_status else acc.login_status),
            "login_checked_at": _local(svc, acc.login_checked_at),
            "messages": int(t.get("messages", 0)), "bytes": int(t.get("bytes", 0)),
            "size": human_size(int(t.get("bytes", 0))),
            "first_backup_at": _local(svc, acc.first_backup_at), "last_backup_at": _local(svc, acc.last_backup_at),
            "last_backup_status": _STATUS_LABELS.get(acc.last_backup_status, acc.last_backup_status or ""),
            "retention": retention_label(acc.retention_days),
            "hold": _hold_label(acc.hold_until) if acc.hold_until else "",
            "dismissed_at": _local(svc, acc.dismissed_at),
            "employee": (emp["full_name"] if emp else "") or "", "department": (emp["department"] if emp else "") or "",
            "position": (emp["position"] if emp else "") or "",
            "schedules": len(schedules.get(acc.id, [])), "notes": acc.notes or "",
        })
    return rows


def export_list(svc, ids: Optional[Sequence[Any]] = None, fmt: str = "xlsx",
                with_notes: bool = True) -> Tuple[bytes, str, str]:
    """Список ящиков файлом: ``(содержимое, имя файла, тип)``.

    ``with_notes=False`` — без заметок администратора (выгрузка оператора).
    """
    rows = account_rows(svc, ids)
    columns = [c for c in LIST_COLUMNS if with_notes or c[0] != "notes"]
    stamp = datetime.now(svc.local_tz()).strftime("%Y%m%d_%H%M")
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter=";")
        writer.writerow([title for _k, title in columns])
        for row in rows:
            writer.writerow([_csv_safe(row[k]) for k, _t in columns])
        # BOM — чтобы Excel сразу открыл файл в UTF-8, а не «кракозябрами»
        return ("﻿" + buf.getvalue()).encode("utf-8"), f"mailboxes_{stamp}.csv", "text/csv; charset=utf-8"
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Почтовые ящики"
    ws.append([title for _k, title in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6FED")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_xlsx_safe(row[k]) for k, _t in columns])
    widths = {"name": 30, "username": 32, "host": 22, "notes": 40, "employee": 30, "department": 24,
              "position": 24, "login": 18, "last_backup_at": 18, "first_backup_at": 18, "login_checked_at": 18}
    for i, (key, _title) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 14)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    out = io.BytesIO()
    wb.save(out)
    return (out.getvalue(), f"mailboxes_{stamp}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _csv_safe(value: Any) -> Any:
    """Текст, начинающийся с = + - @, Excel принял бы за формулу (CSV-инъекция)."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _xlsx_safe(value: Any) -> Any:
    """Строка без управляющих символов (Excel их не принимает) и не формула."""
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        value = ILLEGAL_CHARACTERS_RE.sub(" ", value)
        if value[:1] == "=":
            return "'" + value
    return value


__all__ = ["HANDLERS", "GROUPS", "Param", "Handler", "BulkContext", "run", "describe_actions", "history",
           "history_item", "cancel_op_jobs", "account_rows", "export_list", "retention_label"]
