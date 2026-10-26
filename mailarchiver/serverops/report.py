"""
Отчёт о ящиках на сервере: сколько писем и места, квота, непрочитанные, самые
большие папки — и файлы Excel (отчёт, список ящиков раздела, выданные пароли).

Учётные записи Axigen читаются через CLI (папки учётной записи с числом писем и
размером — как их показывает сервер), остальные ящики — по IMAP: STATUS по
каждой папке и квота (GETQUOTAROOT), если сервер её поддерживает. На сервере
ничего не меняется.
"""
from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple

from ..bulk import _xlsx_safe
from ..errors import MailArchiverError, ValidationError
from ..util import human_size
from . import base
from . import targets as tg
from .base import Ctx, ServerAction

#: Сколько самых больших папок называть в отчёте.
TOP_FOLDERS = 5


class ReportAction(ServerAction):
    key = "report"
    label = "Отчёт: место, квота, папки"
    group = "reports"
    icon = "📊"
    needs = "any"
    read_only = True
    desc = ("Прочитать с сервера, сколько у каждого ящика писем и места, какая квота и насколько она заполнена, "
            "сколько непрочитанных и какие папки самые большие. Учётные записи Axigen читаются через CLI, остальные "
            "ящики — по IMAP. Итог — таблица в задании и файл Excel в «Истории». На сервере ничего не меняется.")

    def check(self, ctx, t):
        if t.axigen is not None:
            if t.deleted:
                return "учётная запись удалена на сервере"
            st = ctx.axigen()
            if st.enabled and st.configured:
                return ""
        problem = tg.imap_problem(ctx.svc, t.mailbox)
        if t.server is not None:
            if t.deleted:
                return "учётная запись удалена на сервере"
            if problem and self._by_api(t):
                return ""
        return problem

    @staticmethod
    def _by_api(t) -> bool:
        from ..mailservers.base import Cap
        from . import server_accounts
        return server_accounts.has(t, Cap.USAGE)

    def plan(self, ctx, t):
        st = ctx.axigen()
        if t.axigen is not None and st.enabled and st.configured:
            return "через CLI Axigen"
        if t.server is not None and tg.imap_problem(ctx.svc, t.mailbox):
            from . import server_accounts
            return f"через API {server_accounts.label(t)} (без папок и непрочитанных)"
        return "по IMAP"

    def execute(self, ctx, t):
        st = ctx.axigen()
        if t.axigen is not None and st.enabled and st.configured:
            data = self._axigen(ctx, t)
        elif t.server is not None and tg.imap_problem(ctx.svc, t.mailbox):
            data = self._server(ctx, t)
        else:
            data = self._imap(ctx, t)
            if t.server is not None and data.get("quota") is None and self._by_api(t):
                # IMAP квоту не показал — берём из API сервера (последняя сверка)
                quota = (t.server.get("quota_kb") or 0) * 1024 or None
                used = t.server.get("used_bytes")
                data["quota"] = quota
                if data.get("size") is None:
                    data["size"] = used
                if quota and data.get("size") is not None:
                    data["percent"] = round(100.0 * data["size"] / quota, 1)
        data.update({"email": t.email, "name": t.name, "server": t.host})
        from .usage import record, today
        record(ctx.db, t.key, data.get("size"), data.get("quota"), data.get("messages"), day=today(ctx.svc))
        ctx.data.setdefault("item_extra", {})[t.key] = {"report": data}
        parts = [f"писем {data.get('messages') if data.get('messages') is not None else '?'}",
                 f"занято {human_size(data['size']) if data.get('size') is not None else '?'}"]
        if data.get("quota"):
            parts.append(f"квота {human_size(data['quota'])} ({data.get('percent')} %)")
        if data.get("unseen"):
            parts.append(f"непрочитанных {data['unseen']}")
        if data.get("top"):
            parts.append("больше всего: " + data["top"])
        return "ok", ", ".join(parts)

    def _axigen(self, ctx, t) -> Dict[str, Any]:
        from ..axigen import actions, directory
        from ..axigen.service import logged_session
        local, domain = actions.parse_email(t.email)
        with logged_session(ctx.svc, username=ctx.user, kind="read", summary=f"Отчёт {t.email}",
                            st=ctx.axigen()) as cli:
            directory.enter_account(cli, domain, local)
            try:
                folders, _table = directory.account_folders(cli)
            finally:
                cli.to_root()
            quota_kb = actions.read_quota_kb(cli, t.email)
        sizes = [f.size for f in folders if f.size is not None]
        counts = [f.messages for f in folders if f.messages is not None]
        size = sum(sizes) if sizes else None
        quota = quota_kb * 1024 if quota_kb else None
        top = sorted((f for f in folders if f.size), key=lambda f: -(f.size or 0))[:TOP_FOLDERS]
        return {"source": "Axigen CLI", "messages": sum(counts) if counts else None, "size": size,
                "quota": quota, "percent": round(100.0 * size / quota, 1) if quota and size is not None else None,
                "unseen": None, "folders": len(folders),
                "top": "; ".join(f"{f.name} {human_size(f.size)}" for f in top)}

    def _server(self, ctx, t) -> Dict[str, Any]:
        from ..mailservers import service as mss
        from ..mailservers import store as msstore
        from ..mailservers.base import normalize_account
        from . import server_accounts
        with mss.session(ctx.svc, t.conn, user=ctx.user, kind="read", summary=f"Отчёт {t.email}") as drv:
            raw = drv.get_account(t.email)
        if raw is None:
            raise ValidationError("учётной записи на сервере больше нет")
        acc = normalize_account(raw)
        msstore.upsert_account(ctx.db, t.conn.id, acc)
        size = acc.get("used_bytes")
        quota = acc.get("quota_bytes") or None
        messages = acc.get("messages")
        source = f"API {server_accounts.label(t)}"
        if size is None and t.server.get("used_bytes") is not None:
            # одна учётная запись занятого места не сообщает (Microsoft 365 — только отчётом) — из сверки
            size, messages = t.server.get("used_bytes"), t.server.get("messages")
            quota = quota or ((t.server.get("quota_kb") or 0) * 1024 or None)
            msstore.update_fields(ctx.db, t.conn.id, t.email, used_bytes=size, messages=messages,
                                  quota_kb=(quota // 1024) if quota else None)
            source += " (по последней сверке)"
        return {"source": source, "messages": messages, "size": size,
                "quota": quota, "percent": round(100.0 * size / quota, 1) if quota and size is not None else None,
                "unseen": None, "folders": None, "top": ""}

    def _imap(self, ctx, t) -> Dict[str, Any]:
        from .imap_ops import connect
        conn = connect(ctx, t.mailbox)
        try:
            folders = [f for f in conn.list_folders() if f.selectable]
            total = unseen = 0
            per: List[Tuple[str, int]] = []
            known = True
            for f in folders:
                counts = conn.message_counts(f.name)
                if counts is None:
                    known = False
                    continue
                total += counts["messages"]
                unseen += counts["unseen"]
                per.append((f.name, counts["messages"]))
            quotas = conn.quota("INBOX")
        finally:
            conn.close()
        storage = next((q for q in quotas if q["resource"] == "STORAGE"), None)
        size = storage["usage"] * 1024 if storage else None
        quota = storage["limit"] * 1024 if storage and storage["limit"] else None
        top = sorted(per, key=lambda x: -x[1])[:TOP_FOLDERS]
        return {"source": "IMAP", "messages": total if (known or per) else None, "size": size, "quota": quota,
                "percent": round(100.0 * size / quota, 1) if quota and size is not None else None,
                "unseen": unseen, "folders": len(folders),
                "top": "; ".join(f"{name} {n} писем" for name, n in top if n)}

    def job_finished(self, ctx, result):
        result["report"] = True


base.register(ReportAction())


# ---------------------------------------------------------------------------
#  Файлы Excel
# ---------------------------------------------------------------------------
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(title: str, columns: Sequence[Tuple[str, str, int]], rows: Sequence[Dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([c[1] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6FED")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_xlsx_safe(row.get(c[0])) for c in columns])
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = col[2]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M")


def report_file(svc, op_id: int) -> Tuple[bytes, str, str]:
    item = base.history_item(svc, op_id)
    if item is None or item.get("action") != ReportAction.key or not item.get("job_id"):
        raise ValidationError("Отчёт не найден.")
    row = svc.db.get_job(item["job_id"])
    try:
        result = json.loads(row["result"] or "{}") if row is not None else {}
    except (TypeError, ValueError):
        result = {}
    rows = []
    for it in (result or {}).get("items") or []:
        rep = it.get("report")
        if not rep:
            rows.append({"email": it.get("name"), "error": it.get("detail")})
            continue
        rows.append({**rep, "size_mb": round(rep["size"] / 1048576, 1) if rep.get("size") is not None else None,
                     "quota_mb": round(rep["quota"] / 1048576, 1) if rep.get("quota") else None})
    if not rows:
        raise ValidationError("В отчёте пока нет строк — задание ещё идёт или не прочитало ни одного ящика.")
    columns = (("email", "Адрес", 32), ("name", "Имя", 28), ("server", "Сервер", 22), ("source", "Откуда", 12),
               ("messages", "Писем", 10), ("size_mb", "Занято, МБ", 12), ("quota_mb", "Квота, МБ", 12),
               ("percent", "Заполнено, %", 12), ("unseen", "Непрочитанных", 13), ("folders", "Папок", 8),
               ("top", "Самые большие папки", 60), ("error", "Не прочитано", 40))
    return _workbook("Отчёт", columns, rows), f"server_report_{op_id}_{_stamp()}.xlsx", _XLSX


def passwords_file(svc, op_id: int) -> Tuple[bytes, str, str]:
    """Новые пароли операции — файлом, один раз."""
    from . import secrets
    if base.history_item(svc, op_id) is None:
        raise ValidationError("Операция не найдена.")
    columns = (("email", "Адрес", 34), ("name", "Имя", 30), ("password", "Пароль", 24), ("note", "Примечание", 40))
    # файл собирается до того, как пароли стираются: сбой сборки не должен их потерять
    data = secrets.take(svc.db, op_id, build=lambda rows: _workbook("Пароли", columns, rows) if rows else None)
    if data is None:
        raise ValidationError("Паролей для выдачи нет: файл уже скачан, прошло больше суток или задание ещё не "
                              "дошло до этих учётных записей.")
    return data, f"passwords_{op_id}_{_stamp()}.xlsx", _XLSX


def targets_file(svc, keys: Optional[Sequence[str]] = None) -> Tuple[bytes, str, str]:
    """Список ящиков раздела (все или перечисленные) — для Excel."""
    items = tg.load(svc)
    own = tg.own_addresses(items)
    if keys:
        wanted = set(keys)
        items = [t for t in items if t.key in wanted]
    totals = svc.db.message_totals_by_account()
    keeps = tg.keeps_everything_map(svc, [t.mailbox for t in items if t.mailbox is not None])
    from .usage import latest
    from .usage import trends as usage_trends
    dismissed = tg.dismissed_map(svc)
    trends = usage_trends(svc.db, [t.key for t in items])
    last = latest(svc.db)
    rows = []
    for t in items:
        r = tg.row_dict(svc, t, totals, keeps, dismissed, trends, own, last)
        mb = r.get("mailbox") or {}
        trend = r.get("trend") or {}
        if r["kind"] != "axigen" and r.get("usage"):
            # ящик не на Axigen: место — по последнему «Отчёту»
            r["used_bytes"] = r["usage"]["used"]
            r["quota_kb"] = (r["usage"]["quota"] or 0) // 1024 or None
            if r["quota_kb"] and r["used_bytes"] is not None:
                r["quota_percent"] = round(100.0 * r["used_bytes"] / (r["quota_kb"] * 1024), 1)
        rows.append({
            "email": r["email"], "name": r["name"], "server": r["host"],
            "kind": "Axigen" if r["kind"] == "axigen" else "IMAP",
            "login": {True: "закрыт", False: "открыт"}.get(r.get("login_blocked"), ""),
            "quota_mb": round(r["quota_kb"] / 1024, 1) if r.get("quota_kb") else None,
            "used_mb": round(r["used_bytes"] / 1048576, 1) if r.get("used_bytes") is not None else None,
            "percent": r.get("quota_percent"), "last_login": (r.get("last_login") or "")[:16].replace("T", " "),
            "aliases": ", ".join(r.get("aliases") or []),
            "mailbox": mb.get("name") or "", "auth": mb.get("auth") or "",
            "imap": ("да" if mb.get("imap") else (mb.get("imap_problem") or "")) if mb else "",
            "archive_messages": mb.get("messages") if mb else None,
            "keeps_all": ("да" if mb.get("keeps_all") else "нет") if mb else "",
            "growth_mb": round(trend["per_month"] / 1048576, 1) if trend else None,
            "days_to_quota": trend.get("days_to_quota") if trend else None,
            "dismissed": (r.get("dismissed_at") or "")[:10],
            "deleted": "удалена на сервере" if r.get("deleted") else ""})
    columns = (("email", "Адрес", 32), ("name", "Имя", 28), ("server", "Сервер", 22), ("kind", "Вид", 8),
               ("login", "Вход", 9), ("quota_mb", "Квота, МБ", 11), ("used_mb", "Занято, МБ", 11),
               ("percent", "Заполнено, %", 11), ("last_login", "Последний вход", 17), ("aliases", "Псевдонимы", 30),
               ("mailbox", "Ящик в архиве", 26), ("auth", "Как копируется", 18), ("imap", "Вход по IMAP", 24),
               ("archive_messages", "Писем в архиве", 13), ("keeps_all", "Архив хранит всё", 12),
               ("growth_mb", "Рост в месяц, МБ", 12), ("days_to_quota", "До квоты, дней", 11),
               ("dismissed", "Уволен", 11), ("deleted", "Состояние", 16))
    return _workbook("Ящики на сервере", columns, rows), f"server_mailboxes_{_stamp()}.xlsx", _XLSX


__all__ = ["ReportAction", "report_file", "passwords_file", "targets_file", "MailArchiverError"]
