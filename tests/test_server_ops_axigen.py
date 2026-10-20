"""Раздел «Ящики на сервере»: действия с учётными записями Axigen (учебные CLI и FTP Backup).

Каждое изменение проверяется по состоянию учебного сервера, а не по ответу клиента.
"""
import io

import pytest

from axigen_fake import FakeAccount, FakeAxigenCLI, FakeAxigenFTP, FtpWorld, World, message
from mailarchiver import models
from mailarchiver.axigen import store as axstore
from mailarchiver.errors import ValidationError
from mailarchiver.models import AuthType, JobStatus, JobType
from mailarchiver.queue.jobs import JobContext
from mailarchiver.serverops import base, create, report, secrets
from mailarchiver.serverops import targets as tg
from mailarchiver.serverops.axigen_ops import render_alias, split_full_name, translit

ADMIN = {"username": "tester", "role": "admin"}
IV, PE = "ax:ivanov@company.ru", "ax:petrov@company.ru"


@pytest.fixture()
def ax(services):
    world = World()
    fworld = FtpWorld({"domains": {"company.ru": {"ivanov": {"Inbox": {"1.eml": message("Первое")}},
                                                  "petrov": {"Inbox": {}}, "postmaster": {"Inbox": {}}}}})
    with FakeAxigenCLI(world) as cli, FakeAxigenFTP(fworld) as ftp:
        services.db.set_settings_many({
            "axigen.enabled": True, "axigen.host": "127.0.0.1", "axigen.cli_port": cli.port,
            "axigen.ftp_port": ftp.port, "axigen.admin_user": "admin", "axigen.admin_password": "Adm1nPass",
            "axigen.timeout_s": 5, "axigen.imap_host": "mx.company.ru", "axigen.allow_changes": True})
        services.apply_runtime_settings()
        box = type("Ax", (), {"world": world, "svc": services, "cli": cli})
        sync(services)
        yield box


def run_job(svc, job_type, params):
    job_id = svc.db.enqueue_job(job_type, None, params, 5, 1, "tester")
    ctx = JobContext(svc, job_id, job_type, None, params)
    from mailarchiver.queue.jobs import HANDLERS
    return HANDLERS[job_type](ctx)


def sync(svc):
    return run_job(svc, JobType.AXIGEN_SYNC, {"requested_by": "tester"})


def act(svc, action, keys, op_id=None, **options):
    res = run_job(svc, JobType.SERVER_ACTION, {"action": action, "keys": keys, "options": options,
                                               "requested_by": "tester", "op_id": op_id})
    return res, {i["key"]: i for i in res["items"]}


def acc(world, name):
    return world.domains["company.ru"].accounts[name]


def mailbox(svc, username, **kw):
    data = dict(name=username, host="mx.company.ru", username=username, password="",
                auth_type=AuthType.AXIGEN_FTP, retention_days=0)
    data.update(kw)
    return svc.db.create_account(models.Account(**data))


# ---------------------------------------------------------------------------
def test_targets_combine_axigen_accounts_and_other_mailboxes(ax):
    linked = mailbox(ax.svc, "ivanov@company.ru")
    other = mailbox(ax.svc, "someone@other.org", host="imap.other.org", auth_type=AuthType.PASSWORD, password="x")
    keys = {t.key: t for t in tg.load(ax.svc)}
    assert {IV, PE, "ax:postmaster@company.ru", f"mb:{other}"} == set(keys)
    assert keys[IV].mailbox.id == linked and keys[IV].kind == "axigen" and keys[f"mb:{other}"].kind == "imap"
    rows = {r["key"]: r for r in tg.overview(ax.svc)["rows"]}
    assert rows[IV]["first_name"] == "Иван" and rows[IV]["mailbox"]["auth"] == "FTP Backup Axigen"
    assert rows[f"mb:{other}"]["mailbox"]["imap"] is True


def test_changes_need_permission(ax):
    ax.svc.db.set_settings_many({"axigen.allow_changes": False})
    ax.svc.apply_runtime_settings()
    with pytest.raises(ValidationError, match="не разрешены"):
        base.run(ax.svc, ADMIN, "ax_block", [IV], {}, preview=True)
    # чтение — можно
    assert base.run(ax.svc, ADMIN, "ax_refresh", [IV], {}, preview=True)["counts"]["ok"] == 1


def test_block_and_unblock_account_without_archive_mailbox(ax):
    res, items = act(ax.svc, "ax_block", [PE], final_copy=True, keep_receiving=True, random_password=True)
    assert items[PE]["status"] == "success", items[PE]
    p = acc(ax.world, "petrov")
    assert p.services == ["localDelivery"] and p.password != "Secret123"
    row = axstore.get_account(ax.svc.db, "petrov@company.ru")
    assert row["login_blocked"] is True and "pop3" in row["saved_services"]
    pv = base.run(ax.svc, ADMIN, "ax_block", [PE], {}, preview=True)
    assert pv["results"][0]["detail"] == "вход уже закрыт"
    res, items = act(ax.svc, "ax_unblock", [PE])
    assert items[PE]["status"] == "success" and "imap" in acc(ax.world, "petrov").services


def test_passwords_file_is_given_once(ax):
    linked = mailbox(ax.svc, "petrov@company.ru", auth_type=AuthType.PASSWORD, password="old")
    op = ax.svc.db.create_bulk_op("tester", "srv_ax_password", "Сменить пароли", {}, 2)
    res, items = act(ax.svc, "ax_password", [IV, PE], op_id=op, mode="random_file", must_change=True)
    assert all(i["status"] == "success" for i in items.values()), items
    iv, pe = acc(ax.world, "ivanov"), acc(ax.world, "petrov")
    assert iv.must_change_pass == "yes" and iv.password != "Secret123"
    assert ax.svc.db.get_account(linked).password == pe.password            # ящик по IMAP продолжит копироваться
    assert secrets.count(ax.svc.db, op) == 2
    data, name, _ = report.passwords_file(ax.svc, op)
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data)).active
    given = {ws.cell(r, 1).value: ws.cell(r, 3).value for r in range(2, 4)}
    assert given == {"ivanov@company.ru": iv.password, "petrov@company.ru": pe.password}
    with pytest.raises(ValidationError, match="уже скачан"):
        report.passwords_file(ax.svc, op)
    # в журнал команд пароли не попадают
    logs = " ".join(r["transcript"] for r in ax.svc.db.query("SELECT transcript FROM axigen_log"))
    assert iv.password not in logs and pe.password not in logs


def test_common_password_is_stored_encrypted_in_job(ax, monkeypatch):
    queued = []
    monkeypatch.setattr(ax.svc.queue, "enqueue", lambda *a, **k: queued.append(a) or 5)
    base.run(ax.svc, ADMIN, "ax_password", [IV], {"mode": "same", "password": "Temp12345"}, preview=False)
    options = queued[0][2]["options"]
    assert "password" not in options and "Temp12345" not in str(queued[0])
    # ни в истории операций, ни в аудите общего пароля нет
    stored = str([tuple(r) for r in ax.svc.db.query("SELECT params, results, summary FROM bulk_ops")])
    stored += str([tuple(r) for r in ax.svc.db.query("SELECT * FROM audit")])
    assert "Temp12345" not in stored and "задан (не сохраняется)" in stored
    res, items = act(ax.svc, "ax_password", [IV], **options)
    assert acc(ax.world, "ivanov").password == "Temp12345"
    with pytest.raises(ValidationError, match="короче"):
        base.run(ax.svc, ADMIN, "ax_password", [IV], {"mode": "same", "password": "1"}, preview=True)


def test_services_off_on_and_block_through_services(ax):
    res, items = act(ax.svc, "ax_services_off", [IV], svc_pop3=True, svc_mobile=True)
    assert items[IV]["status"] == "success"
    assert "pop3" not in acc(ax.world, "ivanov").services and "imap" in acc(ax.world, "ivanov").services
    assert axstore.get_account(ax.svc.db, "ivanov@company.ru")["login_blocked"] is False
    res, items = act(ax.svc, "ax_services_on", [IV], svc_mobile=True)
    assert items[IV]["status"] == "success" and "activesync" in acc(ax.world, "ivanov").services
    # домен не разрешает CalDAV — так и сказано
    res, items = act(ax.svc, "ax_services_on", [IV], svc_dav=True)
    assert items[IV]["status"] == "skipped" and "домен не разрешает" in items[IV]["detail"]
    # отключить всё, кроме приёма писем, — это блокировка, и её можно снять
    res, items = act(ax.svc, "ax_services_off", [PE], svc_pop3=True, svc_imap=True, svc_webmail=True,
                     svc_smtp=True)
    row = axstore.get_account(ax.svc.db, "petrov@company.ru")
    assert row["login_blocked"] is True and row["saved_services"]
    res, items = act(ax.svc, "ax_unblock", [PE])
    assert "webmail" in acc(ax.world, "petrov").services


def test_quota_raise_and_skip_when_too_small(ax):
    acc(ax.world, "ivanov").folders = {"INBOX": (3, 3 * 1024 * 1024 * 1024)}       # занято 3 ГБ
    sync(ax.svc)
    pv = base.run(ax.svc, ADMIN, "ax_quota", [IV, PE], {"quota_mb": 2048, "only_raise": False}, preview=True)
    by = {r["key"]: r for r in pv["results"]}
    assert by[IV]["status"] == "skip" and "больше новой квоты" in by[IV]["detail"]
    res, items = act(ax.svc, "ax_quota", [PE], quota_mb=2048, only_raise=False, threshold=90)
    assert items[PE]["status"] == "success", items[PE]
    p = acc(ax.world, "petrov")
    assert p.quota_kb == 2048 * 1024 and p.quota_threshold == 90
    assert axstore.get_account(ax.svc.db, "petrov@company.ru")["quota_kb"] == 2048 * 1024
    pv = base.run(ax.svc, ADMIN, "ax_quota", [PE], {"quota_mb": 1024, "only_raise": True}, preview=True)
    assert "не меньше новой" in pv["results"][0]["detail"]


def test_alias_by_template_with_collisions(ax):
    assert translit("Пётр Щукин") == "petr shchukin"
    assert render_alias("{f}.{last}", {"first_name": "Иван", "last_name": "Иванов", "name": "ivanov"}) == "i.ivanov"
    w = ax.world
    w.domains["company.ru"].accounts["ivanov2"] = FakeAccount(first_name="Иван", last_name="Иванов")
    w.domains["company.ru"].accounts["sidorov"] = FakeAccount(first_name="Ivan", last_name="Petrov")
    acc(w, "petrov").aliases = ["ivan.petrov"]
    sync(ax.svc)
    keys = [IV, "ax:ivanov2@company.ru", "ax:sidorov@company.ru", "ax:postmaster@company.ru"]
    pv = base.run(ax.svc, ADMIN, "ax_alias_add", keys, {"template": "{first}.{last}"}, preview=True)
    by = {r["key"]: r for r in pv["results"]}
    assert by[IV]["status"] == "ok" and "ivan.ivanov@company.ru" in by[IV]["detail"]
    assert "у другой выбранной учётной записи" in by["ax:ivanov2@company.ru"]["detail"]
    assert "уже занят" in by["ax:sidorov@company.ru"]["detail"]
    assert "не заполнены имя и фамилия" in by["ax:postmaster@company.ru"]["detail"]
    res, items = act(ax.svc, "ax_alias_add", [IV], template="{first}.{last}")
    assert items[IV]["status"] == "success" and "ivan.ivanov" in acc(w, "ivanov").aliases
    row = axstore.get_account(ax.svc.db, "ivanov@company.ru")
    assert "ivan.ivanov@company.ru" in row["aliases"]
    res, items = act(ax.svc, "ax_alias_remove", [IV], which="list", aliases="i.ivanov\nnone@company.ru")
    assert items[IV]["status"] == "success" and acc(w, "ivanov").aliases == ["ivan.ivanov"]
    with pytest.raises(ValidationError, match="без «@домен»"):
        base.run(ax.svc, ADMIN, "ax_alias_add", [IV], {"template": "{first}@x.ru"}, preview=True)


def test_contact_from_employees_fills_only_empty_fields(ax):
    assert split_full_name("Петров Пётр Петрович", "last_first") == {"lastName": "Петров", "firstName": "Пётр",
                                                                      "middleName": "Петрович"}
    ax.svc.db.create_employee(full_name="Петров Пётр Петрович", email="petrov@company.ru", department="Продажи",
                              position="Менеджер", phone="+7 900 000-00-00")
    acc(ax.world, "petrov").contact = {"department": "Старый отдел"}
    pv = base.run(ax.svc, ADMIN, "ax_contact", [PE, IV], {"source": "employees"}, preview=True)
    by = {r["key"]: r for r in pv["results"]}
    assert by[IV]["status"] == "skip" and "нет в разделе «Сотрудники»" in by[IV]["detail"]
    res, items = act(ax.svc, "ax_contact", [PE], source="employees", name_order="last_first", phone=True,
                     company="ООО Ромашка", overwrite=False)
    assert items[PE]["status"] == "success", items[PE]
    c = acc(ax.world, "petrov").contact
    assert c["department"] == "Старый отдел"                    # заполненное не трогаем
    assert c["position"] == "Менеджер" and c["company"] == "ООО Ромашка" and c["businessPhone"].startswith("+7")
    res, items = act(ax.svc, "ax_contact", [PE], source="employees", name_order="last_first", phone=True,
                     company="", overwrite=True)
    assert acc(ax.world, "petrov").contact["department"] == "Продажи"


def test_flags_account_type_and_publish(ax):
    res, items = act(ax.svc, "ax_account_type", [IV], value="premium")
    assert items[IV]["status"] == "success" and acc(ax.world, "ivanov").account_type == "premium"
    res, items = act(ax.svc, "ax_account_type", [IV], value="premium")
    assert items[IV]["status"] == "skipped"
    res, items = act(ax.svc, "ax_publish", [IV], value="no")
    assert acc(ax.world, "ivanov").publish == "no"
    res, items = act(ax.svc, "ax_must_change", [IV], value="yes")
    assert acc(ax.world, "ivanov").must_change_pass == "yes"


def test_refresh_rereads_and_notices_deleted_accounts(ax):
    acc(ax.world, "ivanov").quota_kb = 5 * 1024 * 1024
    del ax.world.domains["company.ru"].accounts["petrov"]
    res, items = act(ax.svc, "ax_refresh", [IV, PE])
    assert axstore.get_account(ax.svc.db, "ivanov@company.ru")["quota_kb"] == 5 * 1024 * 1024
    assert axstore.get_account(ax.svc.db, "petrov@company.ru")["deleted_at"]
    assert "помечена удалённой" in items[PE]["detail"]


def test_delete_and_purge_need_archive(ax):
    pv = base.run(ax.svc, ADMIN, "ax_delete", [PE], {}, preview=True)
    assert "ящика нет в архиве" in pv["results"][0]["detail"]
    pv = base.run(ax.svc, ADMIN, "ax_purge", [PE], {"days": 365}, preview=True)
    assert "ящика нет в архиве" in pv["results"][0]["detail"]
    mailbox(ax.svc, "petrov@company.ru", retention_days=30)
    pv = base.run(ax.svc, ADMIN, "ax_purge", [PE], {"days": 365}, preview=True)
    assert "хранит только последние 30 дн." in pv["results"][0]["detail"]
    pv = base.run(ax.svc, ADMIN, "ax_delete", [PE], {}, preview=True)
    assert "не заблокирован" in pv["results"][0]["detail"]


def test_add_mailboxes_immediately(ax):
    res = base.run(ax.svc, ADMIN, "ax_add_mailboxes", [IV, PE], {"auth": AuthType.AXIGEN_FTP, "enable": False},
                   preview=False)
    assert res["counts"]["ok"] == 2 and not res["jobs"]
    boxes = {a.username: a for a in ax.svc.db.list_accounts()}
    assert boxes["ivanov@company.ru"].auth_type == AuthType.AXIGEN_FTP and not boxes["ivanov@company.ru"].enabled
    again = base.run(ax.svc, ADMIN, "ax_add_mailboxes", [IV], {}, preview=True)
    assert "уже копируется" in again["results"][0]["detail"]


def test_report_through_cli(ax):
    res, items = act(ax.svc, "report", [IV])
    rep = items[IV]["report"]
    assert rep["source"] == "Axigen CLI" and rep["messages"] == 4 and rep["quota"] == 1048576 * 1024
    assert "INBOX" in rep["top"]


def test_create_accounts_from_list(ax, monkeypatch):
    text = ("new1@company.ru; Анна; Смирнова\nСидоров Олег <new2@company.ru>\nivanov@company.ru\n"
            "bad-line\nx@unknown.org")
    pv = create.run(ax.svc, ADMIN, {"text": text, "password_mode": "random_file"}, preview=True)
    by = {i["email"]: i for i in pv["items"]}
    assert by["new1@company.ru"]["status"] == "ok" and by["new2@company.ru"]["last"] == "Сидоров"
    assert "уже есть учётная запись" in by["ivanov@company.ru"]["text"]
    assert by["bad-line"]["status"] == "skip" and by["x@unknown.org"]["status"] == "skip"
    queued = []
    monkeypatch.setattr(ax.svc.queue, "enqueue", lambda *a, **k: queued.append(a) or 9)
    res = create.run(ax.svc, ADMIN, {"text": text, "password_mode": "random_file", "must_change": True,
                                     "quota_mb": 512, "add_to_archive": True}, preview=False)
    assert res["op_id"] and res["jobs"] == [9]
    params = queued[0][2]
    out = run_job(ax.svc, JobType.SERVER_ACTION, params)
    assert out["ok"] == 2 and out["final_status"] == JobStatus.SUCCESS, out
    n1 = acc(ax.world, "new1")
    assert n1.first_name == "Анна" and n1.quota_kb == 512 * 1024 and n1.must_change_pass == "yes"
    assert axstore.get_account(ax.svc.db, "new1@company.ru") is not None
    assert {a.username for a in ax.svc.db.list_accounts()} >= {"new1@company.ru", "new2@company.ru"}
    assert secrets.count(ax.svc.db, res["op_id"]) == 2
    with pytest.raises(ValidationError, match="Пароль короче"):
        create.run(ax.svc, ADMIN, {"text": text, "password_mode": "same", "password": "1"}, preview=True)


def test_create_from_employees_lists_only_missing_accounts(ax):
    ax.svc.db.create_employee(full_name="Петров Пётр", email="petrov@company.ru")
    ax.svc.db.create_employee(full_name="Новикова Анна Сергеевна", email="Novikova@company.ru")
    pv = create.run(ax.svc, ADMIN, {"source": "employees", "password_mode": "random_file"}, preview=True)
    assert [i["email"] for i in pv["items"]] == ["novikova@company.ru"]
    assert pv["items"][0]["first"] == "Анна" and pv["items"][0]["last"] == "Новикова"
    assert any("уже есть" in w and w.endswith("1.") for w in pv["warnings"])


def test_secrets_expire(ax):
    secrets.save(ax.svc.db, 5, "a@company.ru", "Pw123456")
    ax.svc.db.execute("UPDATE server_secrets SET created_at='2000-01-01T00:00:00+00:00'")
    assert secrets.count(ax.svc.db, 5) == 0
    # и без открытия «Истории»: их стирает служебное обслуживание
    secrets.save(ax.svc.db, 6, "b@company.ru", "Pw123456")
    secrets.save(ax.svc.db, 7, "c@company.ru", "Pw123456")
    ax.svc.db.execute("UPDATE server_secrets SET created_at='2000-01-01T00:00:00+00:00' WHERE op_id=6")
    from mailarchiver.maintenance import run_maintenance
    assert run_maintenance(ax.svc)["server_secrets"] == 1
    left = [r["op_id"] for r in ax.svc.db.query("SELECT op_id FROM server_secrets")]
    assert left == [7]


def test_cancel_keeps_results_of_processed_accounts(ax, monkeypatch):
    params = {"action": "ax_must_change", "keys": [IV, PE, "ax:postmaster@company.ru"], "options": {"value": "yes"},
              "requested_by": "tester", "op_id": None}
    job_id = ax.svc.db.enqueue_job(JobType.SERVER_ACTION, None, params, 5, 1, "tester")
    ctx = JobContext(ax.svc, job_id, JobType.SERVER_ACTION, None, params)
    calls = []
    # администратор нажал «Отменить» после первого ящика
    monkeypatch.setattr(ctx, "is_cancelled", lambda: calls.append(1) or len(calls) > 1)
    monkeypatch.setattr(ax.svc.db, "is_cancel_requested", lambda _id: True)
    from mailarchiver.queue.jobs import HANDLERS
    res = HANDLERS[JobType.SERVER_ACTION](ctx)
    assert res["final_status"] == JobStatus.CANCELLED and res["ok"] == 1
    assert [i["key"] for i in res["items"]] == [IV] and "не обработано 2" in res["summary"]
    assert acc(ax.world, "ivanov").must_change_pass == "yes" and acc(ax.world, "petrov").must_change_pass == "no"
    # остановка службы — не отмена: задание прерывается как прежде (очередь пометит его сама)
    monkeypatch.setattr(ax.svc.db, "is_cancel_requested", lambda _id: False)
    calls.clear()
    from mailarchiver.errors import JobCancelled
    with pytest.raises(JobCancelled):
        HANDLERS[JobType.SERVER_ACTION](ctx)


def test_password_file_failure_keeps_passwords(ax, monkeypatch):
    secrets.save(ax.svc.db, 11, "a@company.ru", "Pw123456")
    monkeypatch.setattr(report, "_workbook", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("диск")))
    monkeypatch.setattr(base, "history_item", lambda svc, op_id: {"id": op_id})
    with pytest.raises(RuntimeError):
        report.passwords_file(ax.svc, 11)
    assert secrets.count(ax.svc.db, 11) == 1                  # сбой сборки файла пароли не стёр
    monkeypatch.undo()
    monkeypatch.setattr(base, "history_item", lambda svc, op_id: {"id": op_id})
    data, name, _ = report.passwords_file(ax.svc, 11)
    assert data and secrets.count(ax.svc.db, 11) == 0


def test_server_jobs_are_not_retried_blindly(ax):
    job_id = ax.svc.db.enqueue_job(JobType.SERVER_ACTION, None, {"action": "ax_password", "keys": [IV]}, 5, 1,
                                   "tester")
    ax.svc.db.finish_job(job_id, JobStatus.PARTIAL, {"summary": "x"})
    with pytest.raises(ValidationError, match="Повторить"):
        ax.svc.queue.retry(job_id)
    assert ax.svc.db.get_job(job_id)["status"] == JobStatus.PARTIAL


def test_created_account_keeps_its_password_when_a_later_step_fails(ax, monkeypatch):
    from mailarchiver.axigen import actions as axactions
    queued = []
    monkeypatch.setattr(ax.svc.queue, "enqueue", lambda *a, **k: queued.append(a) or 9)
    res = create.run(ax.svc, ADMIN, {"text": "late@company.ru", "password_mode": "random_file"}, preview=False)
    monkeypatch.setattr(axactions, "get_services", lambda cli, email: (_ for _ in ()).throw(ValidationError("сбой")))
    out = run_job(ax.svc, JobType.SERVER_ACTION, queued[0][2])
    it = out["items"][0]
    assert it["status"] == "failed" and "создана" in it["detail"] and "затем ошибка" in it["detail"]
    assert "late" in ax.world.domains["company.ru"].accounts
    assert secrets.count(ax.svc.db, res["op_id"]) == 1          # пароль новой учётной записи не потерян
    audit = [r["detail"] for r in ax.svc.db.query("SELECT detail FROM audit WHERE action='srv_ax_create'")]
    assert "late@company.ru" in audit


def test_create_cancel_keeps_results_of_created_accounts(ax, monkeypatch):
    queued = []
    monkeypatch.setattr(ax.svc.queue, "enqueue", lambda *a, **k: queued.append(a) or 9)
    create.run(ax.svc, ADMIN, {"text": "c1@company.ru\nc2@company.ru", "password_mode": "random_file"}, preview=False)
    params = queued[0][2]
    job_id = ax.svc.db.enqueue_job(JobType.SERVER_ACTION, None, params, 5, 1, "tester")
    ctx = JobContext(ax.svc, job_id, JobType.SERVER_ACTION, None, params)
    calls = []
    monkeypatch.setattr(ctx, "is_cancelled", lambda: calls.append(1) or len(calls) > 1)
    monkeypatch.setattr(ax.svc.db, "is_cancel_requested", lambda _id: True)
    from mailarchiver.queue.jobs import HANDLERS
    res = HANDLERS[JobType.SERVER_ACTION](ctx)
    assert res["final_status"] == JobStatus.CANCELLED and res["ok"] == 1 and len(res["items"]) == 1
    assert "c1" in ax.world.domains["company.ru"].accounts and "c2" not in ax.world.domains["company.ru"].accounts


def test_unconfirmed_password_change_keeps_the_password_with_a_note(ax, monkeypatch):
    from mailarchiver.axigen import actions as axactions
    from mailarchiver.axigen.common import AxigenError
    op = ax.svc.db.create_bulk_op("tester", "srv_ax_password", "Сменить пароли", {}, 1)
    monkeypatch.setattr(axactions, "set_password",
                        lambda cli, email, pw: (_ for _ in ()).throw(AxigenError("Axigen не подтвердил команду")))
    res, items = act(ax.svc, "ax_password", [PE], op_id=op, mode="random_file")
    assert items[PE]["status"] == "failed"
    rows = secrets.take(ax.svc.db, op)
    assert rows[0]["email"] == "petrov@company.ru" and "не подтверждена" in rows[0]["note"]
