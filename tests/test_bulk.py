"""Групповые действия над ящиками: предпросмотр, пропуски, подтверждение, история."""
import io
import json
import os

import pytest

from mailarchiver import bulk, models
from mailarchiver.models import JobStatus, JobType

PW = "Sw0rdfish!"


def _login(client):
    client.post("/api/setup", json={"username": "admin", "password": PW})
    client.post("/api/login", json={"username": "admin", "password": PW})


@pytest.fixture()
def admin(client):
    """Клиент администратора; очередь остановлена — задания только ставятся, не выполняются."""
    _login(client)
    svc = client.app.state.services
    svc.queue.stop(timeout=2)
    # очередь стоит, но «служба останавливается» — неправда: обработчики,
    # вызванные в тесте напрямую, не должны считать себя прерванными
    svc.queue._shutting_down.clear()
    return client, svc


def _acc(svc, name, **kw):
    data = dict(name=name, host="mx.example.ru", port=993, username=f"{name.lower()}@example.ru", password="p")
    data.update(kw)
    return svc.db.create_account(models.Account(**data))


def _bulk(client, action, ids, params=None, preview=True, confirm=""):
    return client.post("/api/accounts/bulk", json={"action": action, "ids": ids, "params": params or {},
                                                   "preview": preview, "confirm": confirm})


def _jobs(svc, job_type=None):
    rows = svc.db.query("SELECT * FROM jobs ORDER BY id")
    return [r for r in rows if job_type is None or r["type"] == job_type]


def test_actions_are_described_for_the_interface(admin):
    client, _svc = admin
    data = client.get("/api/accounts/bulk/actions").json()
    keys = {a["key"] for a in data["actions"]}
    assert {"backup", "backup_sequence", "rebuild_full", "enable", "disable", "retention_set", "hold_set",
            "schedule_set", "set_server", "set_auth", "folders", "export", "restore", "quarantine_delete",
            "delete", "purge", "rename", "notes", "cancel_jobs", "check_logins"} <= keys
    groups = {g["key"] for g in data["groups"]}
    assert all(a["group"] in groups for a in data["actions"])
    export = next(a for a in data["actions"] if a["key"] == "export")
    assert any(p["key"] == "engine" and p["options"] for p in export["params"])


def test_only_admin_can_use_bulk_actions(client):
    r = client.post("/api/accounts/bulk", json={"action": "enable", "ids": [1]})
    assert r.status_code == 401


def test_preview_changes_nothing_and_run_records_history(admin):
    client, svc = admin
    a, b, off = _acc(svc, "A"), _acc(svc, "B"), _acc(svc, "Off", enabled=False)
    pv = _bulk(client, "disable", [a, b, off]).json()
    assert pv["preview"] is True and pv["op_id"] is None
    assert pv["counts"] == {"ok": 2, "skip": 1, "fail": 0}
    assert pv["reasons"] == {"уже выключен": 1}
    assert svc.db.get_account(a).enabled                        # предпросмотр ничего не менял
    run = _bulk(client, "disable", [a, b, off], preview=False).json()
    assert run["op_id"] and run["counts"]["ok"] == 2
    assert not svc.db.get_account(a).enabled and not svc.db.get_account(b).enabled
    hist = client.get("/api/accounts/bulk/history").json()["items"]
    assert hist[0]["id"] == run["op_id"] and hist[0]["ok"] == 2 and hist[0]["skipped"] == 1
    item = client.get(f"/api/accounts/bulk/history/{run['op_id']}").json()
    assert [r["status"] for r in item["results"]] == ["ok", "ok", "skip"]      # порядок как в запросе
    assert any(r["action"] == "bulk_disable" for r in svc.db.list_audit(20))


def test_unknown_ids_and_empty_selection(admin):
    client, svc = admin
    a = _acc(svc, "A", enabled=False)
    r = _bulk(client, "enable", [a, 999999]).json()
    assert r["counts"] == {"ok": 1, "skip": 0, "fail": 1}
    assert _bulk(client, "enable", []).status_code == 400
    assert _bulk(client, "no_such_action", [a]).status_code == 400


def test_enable_clears_auto_disabled_flag(admin):
    client, svc = admin
    a = _acc(svc, "A")
    svc.db.set_account_auto_disabled(a, True)
    _bulk(client, "enable", [a], preview=False)
    acc = svc.db.get_account(a)
    assert acc.enabled and not acc.auto_disabled


def test_disable_cancels_queued_backups(admin):
    client, svc = admin
    a = _acc(svc, "A")
    job = svc.queue.enqueue(JobType.BACKUP, a, {})
    _bulk(client, "disable", [a], {"cancel_queued": True}, preview=False)
    assert svc.db.get_job(job)["status"] == JobStatus.CANCELLED


def test_backup_skips_disabled_passwordless_and_busy(admin):
    client, svc = admin
    ok = _acc(svc, "Ok")
    off = _acc(svc, "Off", enabled=False)
    nopw = _acc(svc, "NoPw", password="")
    busy = _acc(svc, "Busy")
    svc.queue.enqueue(JobType.BACKUP, busy, {})
    r = _bulk(client, "backup", [ok, off, nopw, busy], preview=False).json()
    assert r["counts"] == {"ok": 1, "skip": 3, "fail": 0}
    assert set(r["reasons"]) == {"копирование ящика выключено", "не задан пароль",
                                 "копирование уже идёт или стоит в очереди"}
    backups = [j for j in _jobs(svc, JobType.BACKUP) if j["account_id"] == ok]
    assert len(backups) == 1 and backups[0]["created_by"] == "admin"


def test_rebuild_full_requires_typed_confirmation_and_skips_held(admin):
    client, svc = admin
    a, b = _acc(svc, "A"), _acc(svc, "B")
    held = _acc(svc, "Held")
    svc.db.set_account_hold(held, "9999-12-31", "manual")
    pv = _bulk(client, "rebuild_full", [a, b, held]).json()
    assert pv["danger"] == 2 and pv["confirm_value"] == "2"
    assert pv["counts"]["skip"] == 1 and "удерживается" in next(iter(pv["reasons"]))
    assert "disk" in pv["extra"]
    refused = _bulk(client, "rebuild_full", [a, b, held], preview=False)
    assert refused.status_code == 400 and "введите число ящиков" in refused.json()["message"]
    assert not _jobs(svc, JobType.BACKUP)
    done = _bulk(client, "rebuild_full", [a, b, held], preview=False, confirm="2").json()
    jobs = _jobs(svc, JobType.BACKUP)
    assert done["counts"]["ok"] == 2 and len(jobs) == 2
    assert all(json.loads(j["params"]) == {"rebuild": "full"} and j["max_attempts"] == 1 for j in jobs)
    assert sum(1 for r in svc.db.list_audit(50) if r["action"] == "backup_rebuild_full_request") == 2


def test_rebuild_full_refuses_when_disk_may_run_out(admin, monkeypatch):
    client, svc = admin
    a = _acc(svc, "A")
    svc.db.add_message_index(a, "INBOX", 1, 1, "<m@x>", 50 * 1024 * 1024, "2026-01-01T00:00:00+00:00", "",
                             "INBOX/cur/1", "h")
    import mailarchiver.util as util
    monkeypatch.setattr(util, "disk_free_bytes", lambda path: 10 * 1024 * 1024)
    pv = _bulk(client, "rebuild_full", [a]).json()
    assert pv["warnings"] and "Места на диске может не хватить" in pv["warnings"][0]
    assert _bulk(client, "rebuild_full", [a], preview=False, confirm="1").status_code == 400
    forced = _bulk(client, "rebuild_full", [a], {"force_space": True}, preview=False, confirm="1")
    assert forced.status_code == 200 and forced.json()["counts"]["ok"] == 1


def test_single_rebuild_full_is_refused_for_held_archive(admin):
    client, svc = admin
    a = _acc(svc, "A")
    svc.db.set_account_hold(a, "9999-12-31", "manual")
    r = client.post(f"/api/accounts/{a}/backup", json={"rebuild": "full"})
    assert r.status_code == 400 and "удерживается" in r.json()["message"]


def test_backup_sequence_creates_one_job_for_selection(admin):
    client, svc = admin
    a, b, off = _acc(svc, "A"), _acc(svc, "B"), _acc(svc, "Off", enabled=False)
    r = _bulk(client, "backup_sequence", [a, b, off], {"pause_seconds": 3}, preview=False).json()
    jobs = _jobs(svc, JobType.BACKUP_ALL)
    assert len(jobs) == 1 and r["counts"]["ok"] == 2
    params = json.loads(jobs[0]["params"])
    assert params == {"account_ids": [a, b], "pause_seconds": 3}
    assert all(x.get("job_id") == jobs[0]["id"] for x in r["results"] if x["status"] == "ok")


def test_retention_set_and_run_now(admin):
    client, svc = admin
    a, b = _acc(svc, "A"), _acc(svc, "B", retention_days=7)
    svc.db.add_message_index(a, "INBOX", 1, 1, "<m@x>", 10, "2001-01-01T00:00:00+00:00", "", "INBOX/cur/1", "h")
    r = _bulk(client, "retention_set", [a, b], {"days": 7, "run_now": True}, preview=False).json()
    assert r["counts"] == {"ok": 1, "skip": 1, "fail": 0}
    assert svc.db.get_account(a).retention_days == 7
    assert [j["account_id"] for j in _jobs(svc, JobType.RETENTION)] == [a]
    custom = _bulk(client, "retention_set", [a], {"days": "custom", "custom_days": 45}, preview=False).json()
    assert custom["counts"]["ok"] == 1 and svc.db.get_account(a).retention_days == 45
    assert _bulk(client, "retention_set", [a], {"days": "custom", "custom_days": 0}).status_code == 400


def test_hold_set_and_clear(admin):
    client, svc = admin
    a = _acc(svc, "A")
    assert _bulk(client, "hold_set", [a], {"until": "2001-01-01"}).status_code == 400
    _bulk(client, "hold_set", [a], {"forever": True}, preview=False)
    assert svc.db.get_account(a).on_hold()
    r = _bulk(client, "hold_clear", [a], preview=False).json()
    assert r["counts"]["ok"] == 1 and not svc.db.get_account(a).hold_until


def test_schedule_set_spreads_start_times_and_replaces(admin):
    client, svc = admin
    ids = [_acc(svc, f"A{i}") for i in range(4)]
    svc.db.create_schedule(ids[0], "cron", JobType.BACKUP, "0 3 * * *")
    r = _bulk(client, "schedule_set", ids, {"when": "daily", "time": "02:00", "spread_minutes": 60,
                                             "mode": "replace"}, preview=False).json()
    assert r["counts"]["ok"] == 4
    crons = {row["account_id"]: row["cron_expr"] for row in svc.db.list_schedules()}
    assert [crons[i] for i in ids] == ["0 2 * * *", "15 2 * * *", "30 2 * * *", "45 2 * * *"]
    assert len(svc.db.list_schedules()) == 4                     # прежнее расписание заменено
    missing = _bulk(client, "schedule_set", ids, {"when": "weekdays", "time": "23:30", "mode": "missing"}).json()
    assert missing["counts"]["skip"] == 4
    assert _bulk(client, "schedule_set", ids, {"when": "cron", "cron_expr": "bad"}).status_code == 400
    off = _bulk(client, "schedule_toggle", ids, {"enabled": False}, preview=False).json()
    assert off["counts"]["ok"] == 4 and not any(r["enabled"] for r in svc.db.list_schedules())
    gone = _bulk(client, "schedule_remove", ids, preview=False).json()
    assert gone["counts"]["ok"] == 4 and not svc.db.list_schedules()


def test_server_auth_folders_notes_rename(admin):
    client, svc = admin
    a, b = _acc(svc, "A", notes="старое"), _acc(svc, "B", folder_exclude=["Спам"])
    r = _bulk(client, "set_server", [a, b], {"host": "mx2.example.ru", "port": 143, "security": "starttls"},
              preview=False).json()
    acc = svc.db.get_account(a)
    assert r["counts"]["ok"] == 2 and (acc.host, acc.port, acc.security) == ("mx2.example.ru", 143, "starttls")
    assert acc.password == "p"                                   # пароль не затёрт
    assert _bulk(client, "set_server", [a], {}).status_code == 400
    # вход через администратора почты требует включённой настройки
    assert _bulk(client, "set_auth", [a], {"to": "master"}).status_code == 400
    svc.db.set_settings_many({"mailadmin.enabled": True, "mailadmin.host": "mx2.example.ru",
                              "mailadmin.user": "admin@example.ru", "mailadmin.password": "x"})
    svc.apply_runtime_settings()
    r = _bulk(client, "set_auth", [a, b], {"to": "master"}, preview=False).json()
    assert r["counts"]["ok"] == 2 and svc.db.get_account(b).auth_type == "master"
    r = _bulk(client, "folders", [a, b], {"op": "exclude_add", "folders": "Спам\nКорзина"}, preview=False).json()
    assert svc.db.get_account(a).folder_exclude == ["Спам", "Корзина"]
    assert svc.db.get_account(b).folder_exclude == ["Спам", "Корзина"]
    _bulk(client, "folders", [a], {"op": "exclude_remove", "folders": "Спам"}, preview=False)
    assert svc.db.get_account(a).folder_exclude == ["Корзина"]
    _bulk(client, "notes", [a, b], {"op": "append", "text": "переезд 2026"}, preview=False)
    assert svc.db.get_account(a).notes == "старое\nпереезд 2026" and svc.db.get_account(b).notes == "переезд 2026"
    emp = svc.db.create_employee(full_name="Иванов Иван", email="a@example.ru", department="Бухгалтерия")
    svc.db.set_employee_account(emp, a)
    pv = _bulk(client, "rename", [a, b], {"template": "{employee} ({login})", "fallback": "skip"}).json()
    assert pv["counts"] == {"ok": 1, "skip": 1, "fail": 0}
    assert pv["results"][0]["detail"] == "«A» → «Иванов Иван (a)»"
    _bulk(client, "rename", [a, b], {"template": "{employee} ({login})", "fallback": "username"}, preview=False)
    assert svc.db.get_account(a).name == "Иванов Иван (a)" and svc.db.get_account(b).name == "b@example.ru"
    assert _bulk(client, "rename", [a], {"template": "{bad}"}).status_code == 400


def test_export_and_restore_enqueue_jobs(admin):
    client, svc = admin
    a, empty = _acc(svc, "A"), _acc(svc, "Empty")
    svc.db.add_message_index(a, "INBOX", 1, 1, "<m@x>", 10, "2026-01-01T00:00:00+00:00", "", "INBOX/cur/1", "h")
    r = _bulk(client, "export", [a, empty], {"engine": "eml"}, preview=False).json()
    assert r["counts"] == {"ok": 1, "skip": 1, "fail": 0}
    params = json.loads(_jobs(svc, JobType.EXPORT)[0]["params"])
    assert params["engine"] == "eml" and params["format"] == "eml"
    pv = _bulk(client, "restore", [a], {"target_mode": "original", "dry_run": False}).json()
    assert pv["danger"] == 2 and pv["warnings"]
    assert _bulk(client, "restore", [a], {"target_mode": "prefixed", "target_prefix": ""}).status_code == 400
    safe = _bulk(client, "restore", [a], {"target_mode": "prefixed", "target_prefix": "Архив", "dry_run": True},
                 preview=False).json()
    assert safe["danger"] == 1 and safe["counts"]["ok"] == 1
    opts = json.loads(_jobs(svc, JobType.RESTORE)[0]["params"])
    assert opts["target_prefix"] == "Архив" and opts["dry_run"] is True


def test_cancel_jobs_and_cancel_operation(admin):
    client, svc = admin
    a, b = _acc(svc, "A"), _acc(svc, "B")
    r = _bulk(client, "backup", [a, b], preview=False).json()
    assert len(r["jobs"]) == 2
    item = client.get(f"/api/accounts/bulk/history/{r['op_id']}").json()
    assert item["jobs_active"] == 2 and item["job_progress"] == {"queued": 2}
    cancelled = client.post(f"/api/accounts/bulk/history/{r['op_id']}/cancel").json()
    assert cancelled["cancelled"] == 2
    assert all(j["status"] == JobStatus.CANCELLED for j in _jobs(svc, JobType.BACKUP))
    svc.queue.enqueue(JobType.VERIFY, a, {})
    c = _bulk(client, "cancel_jobs", [a, b], {"which": "all"}, preview=False).json()
    assert c["counts"] == {"ok": 1, "skip": 1, "fail": 0}


def test_purge_and_quarantine_cleanup_job(admin):
    from mailarchiver.queue.jobs import JobContext, handle_cleanup
    client, svc = admin
    a, held, other = _acc(svc, "A"), _acc(svc, "Held"), _acc(svc, "Other")
    for acc_id in (a, held, other):
        svc.store.store_message(acc_id, "INBOX", "/", 1, b"Subject: x\r\n\r\nbody\r\n")
    # прежняя копия (карантин) у ящика Other
    q, _n, _b = svc.store.quarantine_account_files(other)
    assert q and os.path.isdir(q)
    svc.db.set_account_hold(held, "9999-12-31", "manual")
    pv = _bulk(client, "purge", [a, held]).json()
    assert pv["counts"] == {"ok": 1, "skip": 1, "fail": 0} and pv["confirm_value"] == "1"
    r = _bulk(client, "purge", [a, held], preview=False, confirm="1").json()
    assert not svc.db.get_account(a).enabled                    # выключен сразу, удалит задание
    job = _jobs(svc, JobType.CLEANUP)[-1]
    ctx = JobContext(svc, job["id"], JobType.CLEANUP, None, json.loads(job["params"]))
    res = handle_cleanup(ctx)
    assert res["purged"] == 1 and svc.db.get_account(a) is None
    assert not os.path.exists(svc.store.account_dir(a))
    assert svc.db.get_account(held) is not None                  # удерживаемый архив не тронут
    qr = _bulk(client, "quarantine_delete", [other, held], {"only_safe": False}, preview=False, confirm="1").json()
    assert qr["counts"]["ok"] == 1
    job = _jobs(svc, JobType.CLEANUP)[-1]
    res = handle_cleanup(JobContext(svc, job["id"], JobType.CLEANUP, None, json.loads(job["params"])))
    assert res["quarantines"] == 1 and not os.path.exists(q)
    assert r["op_id"]


def test_delete_keeps_files(admin):
    client, svc = admin
    a = _acc(svc, "A")
    svc.store.store_message(a, "INBOX", "/", 1, b"Subject: x\r\n\r\nbody\r\n")
    assert _bulk(client, "delete", [a], preview=False).status_code == 400      # без подтверждения
    _bulk(client, "delete", [a], preview=False, confirm="1")
    assert svc.db.get_account(a) is None and os.path.isdir(svc.store.account_dir(a))


def test_export_list_xlsx_and_csv(admin):
    client, svc = admin
    a = _acc(svc, "=HYPERLINK(\"x\")")
    _acc(svc, "B")
    r = client.post("/api/accounts/export-list", json={"ids": [a], "format": "xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(1, 1).value == "Название" and ws.max_row == 2
    assert str(ws.cell(2, 1).value).startswith("'=")             # не формула
    csv_resp = client.post("/api/accounts/export-list", json={"format": "csv"})
    text = csv_resp.content.decode("utf-8")
    assert text.startswith("﻿Название;") and "'=HYPERLINK" in text and text.count("\n") == 3


def test_describe_and_run_do_not_touch_other_accounts(services):
    """Ядро без веба: run() с preview не меняет базу, а список действий описан полностью."""
    a = services.db.create_account(models.Account(name="A", host="h", username="u", password="p", enabled=False))
    user = {"username": "tester", "role": "admin"}
    res = bulk.run(services, user, "enable", [a], preview=True)
    assert res["counts"]["ok"] == 1 and not services.db.get_account(a).enabled
    desc = bulk.describe_actions(services)
    assert len(desc["actions"]) == len(bulk.HANDLERS)


def test_copy_settings_from_template_mailbox(admin):
    client, svc = admin
    src = _acc(svc, "Src", host="mx2.example.ru", port=143, security="starttls", folder_exclude=["Спам"],
               retention_days=30)
    svc.db.create_schedule(src, "cron", JobType.BACKUP, "10 1 * * *")
    a, b = _acc(svc, "A"), _acc(svc, "B", password="")
    svc.db.create_schedule(a, "cron", JobType.BACKUP, "0 3 * * *")
    pv = _bulk(client, "copy_settings", [src, a, b], {"source": src, "copy_schedules": True}).json()
    assert pv["counts"] == {"ok": 2, "skip": 1, "fail": 0} and pv["reasons"] == {"это и есть ящик-образец": 1}
    _bulk(client, "copy_settings", [a, b], {"source": src, "copy_schedules": True}, preview=False)
    for acc_id in (a, b):
        acc = svc.db.get_account(acc_id)
        assert (acc.host, acc.port, acc.security) == ("mx2.example.ru", 143, "starttls")
        assert acc.folder_exclude == ["Спам"] and acc.retention_days == 30
        assert [s["cron_expr"] for s in svc.db.list_schedules(acc_id)] == ["10 1 * * *"]
    assert svc.db.get_account(a).password == "p" and svc.db.get_account(b).password == ""
    again = _bulk(client, "copy_settings", [a], {"source": src, "copy_schedules": True}).json()
    assert again["reasons"] == {"настройки уже как у образца": 1}
    assert _bulk(client, "copy_settings", [a], {"source": src, "copy_server": False, "copy_folders": False,
                                                "copy_retention": False}).status_code == 400


def test_create_mailboxes_from_list(admin):
    client, svc = admin
    _acc(svc, "Old", username="old@example.ru")
    emp = svc.db.create_employee(full_name="Петров Пётр", email="petrov@example.ru")
    text = ("email;пароль\n"
            "petrov@example.ru;Secret1\n"
            "Сидорова Анна <sidorova@example.ru>\tSecret2\n"
            "nopass@example.ru\n"
            "OLD@example.ru;x\n"
            "petrov@example.ru;dup\n"
            "не адрес;x\n")
    body = {"text": text, "host": "mx.example.ru", "port": 993, "security": "ssl", "auth_type": "password",
            "enable": True, "schedule_time": "03:30"}
    pv = client.post("/api/accounts/bulk-create", json={**body, "preview": True}).json()
    assert pv["counts"] == {"ok": 3, "skip": 2, "fail": 1}
    assert svc.db.get_account_by_username("petrov@example.ru") is None           # предпросмотр ничего не завёл
    r = client.post("/api/accounts/bulk-create", json={**body, "preview": False}).json()
    assert r["counts"]["ok"] == 3 and r["op_id"]
    petrov = svc.db.get_account_by_username("petrov@example.ru")
    assert petrov.name == "Петров Пётр" and petrov.password == "Secret1" and petrov.enabled
    assert svc.db.get_employee(emp)["account_id"] == petrov.id
    sid = svc.db.get_account_by_username("sidorova@example.ru")
    assert sid.name == "Сидорова Анна" and sid.password == "Secret2"
    nopass = svc.db.get_account_by_username("nopass@example.ru")
    assert not nopass.enabled and not nopass.password
    assert [s["cron_expr"] for s in svc.db.list_schedules(petrov.id)] == ["30 3 * * *"]
    item = client.get(f"/api/accounts/bulk/history/{r['op_id']}").json()
    assert "Secret1" not in json.dumps(item, ensure_ascii=False)                  # пароли не в истории
    assert all("Secret" not in (row["detail"] or "") for row in svc.db.list_audit(50))
    assert client.post("/api/accounts/bulk-create", json={**body, "host": ""}).status_code == 400
    assert client.post("/api/accounts/bulk-create", json={**body, "text": ""}).status_code == 400
    assert client.post("/api/accounts/bulk-create", json={**body, "auth_type": "master"}).status_code == 400


def test_quarantine_cleanup_deletes_only_previewed_dirs_and_skips_busy(admin):
    from mailarchiver.queue.jobs import JobContext, handle_cleanup
    client, svc = admin
    a, busy = _acc(svc, "A"), _acc(svc, "Busy")
    for acc_id in (a, busy):
        svc.store.store_message(acc_id, "INBOX", "/", 1, b"Subject: x\r\n\r\nbody\r\n")
        svc.store.quarantine_account_files(acc_id)
    svc.queue.enqueue(JobType.BACKUP, busy, {"rebuild": "full"})       # «с нуля» стоит в очереди
    # удаление прежних копий необратимо — без подтверждения числом ящиков не выполняется
    assert _bulk(client, "quarantine_delete", [a, busy], {"only_safe": False}, preview=False).status_code == 400
    r = _bulk(client, "quarantine_delete", [a, busy], {"only_safe": False}, preview=False, confirm="1").json()
    assert r["counts"] == {"ok": 1, "skip": 1, "fail": 0}
    job = _jobs(svc, JobType.CLEANUP)[-1]
    params = json.loads(job["params"])
    old_paths = params["items"][0]["paths"]
    assert len(old_paths) == 1
    # после проверки появилась НОВАЯ прежняя копия (прошла копия «с нуля»)
    svc.store.store_message(a, "INBOX", "/", 2, b"Subject: y\r\n\r\nbody\r\n")
    import time as _t
    _t.sleep(1.1)                                        # другое имя каталога карантина
    newer, _n, _b = svc.store.quarantine_account_files(a)
    res = handle_cleanup(JobContext(svc, job["id"], JobType.CLEANUP, None, params))
    assert res["quarantines"] == 1
    assert not os.path.exists(old_paths[0]) and os.path.isdir(newer)       # новую не тронули


def test_cleanup_skips_account_taken_by_another_job(admin):
    from mailarchiver.queue.jobs import JobContext, handle_cleanup
    client, svc = admin
    a = _acc(svc, "A")
    _bulk(client, "purge", [a], preview=False, confirm="1")
    job = _jobs(svc, JobType.CLEANUP)[-1]
    assert svc.queue.acquire_account(a, 999999)           # ящик занят проходом по очереди
    try:
        res = handle_cleanup(JobContext(svc, job["id"], JobType.CLEANUP, None, json.loads(job["params"])))
    finally:
        svc.queue.release_account(a, 999999)
    assert res["purged"] == 0 and svc.db.get_account(a) is not None and res["problems"]


def test_single_purge_and_delete_refused_while_account_is_busy(admin):
    client, svc = admin
    a = _acc(svc, "A")
    assert svc.queue.acquire_account(a, 777)
    try:
        r = client.post(f"/api/accounts/{a}/purge", json={"confirm_name": "A"})
        assert r.status_code == 400 and "выполняется задание" in r.json()["message"]
        d = client.delete(f"/api/accounts/{a}")
        assert d.status_code == 400
    finally:
        svc.queue.release_account(a, 777)
    assert client.post(f"/api/accounts/{a}/purge", json={"confirm_name": "A"}).status_code == 200


def test_bulk_delete_skips_held_archives(admin):
    client, svc = admin
    a, held = _acc(svc, "A"), _acc(svc, "Held")
    svc.db.set_account_hold(held, "9999-12-31", "manual")
    pv = _bulk(client, "delete", [a, held]).json()
    assert pv["counts"] == {"ok": 1, "skip": 1, "fail": 0} and "удерживается" in next(iter(pv["reasons"]))


def test_schedule_spread_past_midnight_moves_weekdays(admin):
    client, svc = admin
    ids = [_acc(svc, f"A{i}") for i in range(2)]
    _bulk(client, "schedule_set", ids, {"when": "weekdays", "time": "23:30", "spread_minutes": 60}, preview=False)
    crons = sorted(r["cron_expr"] for r in svc.db.list_schedules())
    assert crons == ["0 0 * * 2-6", "30 23 * * 1-5"]


def test_create_list_never_echoes_line_content_and_keeps_password(admin):
    client, svc = admin
    text = ("ivanov@example.ru Secret-Pa55\n"                # разделитель — пробел: строка непонятна
            "Secret2;petrov@example.ru\n"                    # колонки перепутаны
            "tab@example.ru\t pass with spaces \tТабов Иван\n")
    body = {"text": text, "host": "mx.example.ru", "preview": False}
    r = client.post("/api/accounts/bulk-create", json=body).json()
    dump = json.dumps(r, ensure_ascii=False)
    assert "Secret-Pa55" not in dump and "Secret2" not in dump
    assert r["counts"] == {"ok": 1, "skip": 0, "fail": 2}
    acc = svc.db.get_account_by_username("tab@example.ru")
    assert acc.password == " pass with spaces " and acc.name == "Табов Иван"
    hist = client.get(f"/api/accounts/bulk/history/{r['op_id']}").json()
    assert "Secret" not in json.dumps(hist, ensure_ascii=False)


def test_export_list_survives_control_characters(admin):
    client, svc = admin
    _acc(svc, "A", notes="строка\x0bс вертикальной табуляцией\x01")
    r = client.post("/api/accounts/export-list", json={"format": "xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_validation_errors_do_not_echo_input(admin):
    client, _svc = admin
    r = client.post("/api/accounts/bulk-create", json={"text": "x" * 2_000_001, "host": "h"})
    assert r.status_code == 422 and "xxxx" not in r.text


def test_single_quarantine_delete_respects_archive_hold(admin):
    """Меню ящика → «Прежние копии»: при удержании архива карантин не удаляется (как и групповым действием)."""
    client, svc = admin
    a = _acc(svc, "Held")
    cur = os.path.join(svc.store.account_dir(a), "INBOX", "cur")
    os.makedirs(cur, exist_ok=True)
    with open(os.path.join(cur, "1.eml"), "wb") as fh:
        fh.write(b"Subject: old\r\n\r\nold")
    q, _files, _bytes = svc.store.quarantine_account_files(a)
    assert q and os.path.isdir(q)
    svc.db.set_account_hold(a, "2099-12-31", "manual")
    r = client.post(f"/api/accounts/{a}/quarantines/delete", json={"path": q})
    assert r.status_code == 400 and "удерживается" in r.json()["message"]
    assert os.path.isdir(q)
    svc.db.set_account_hold(a, "", "")
    # несравнённую прежнюю копию — только с явным подтверждением потери писем
    r = client.post(f"/api/accounts/{a}/quarantines/delete", json={"path": q})
    assert r.status_code == 400 and "не сравнивалась" in r.json()["message"]
    r = client.post(f"/api/accounts/{a}/quarantines/delete", json={"path": q, "force": True})
    assert r.status_code == 200
    from mailarchiver.queue.jobs import JobContext, handle_cleanup
    job = svc.db.get_job(r.json()["job_id"])
    handle_cleanup(JobContext(svc, job["id"], JobType.CLEANUP, None, json.loads(job["params"])))
    assert not os.path.exists(q)


def test_plain_delete_respects_archive_hold(admin):
    """«Удалить ящик» (без архива) при удержании запрещено — как групповое удаление и «вместе с архивом»."""
    client, svc = admin
    a = _acc(svc, "HeldDel")
    svc.db.set_account_hold(a, "2099-12-31", "manual")
    r = client.delete(f"/api/accounts/{a}")
    assert r.status_code == 400 and "удерживается" in r.json()["message"]
    assert svc.db.get_account(a) is not None
    svc.db.set_account_hold(a, "", "")
    assert client.delete(f"/api/accounts/{a}").status_code == 200
    assert svc.db.get_account(a) is None
