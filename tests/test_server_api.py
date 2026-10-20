"""REST API раздела «Ящики на сервере»: права, список, проверка и запуск, история, файлы.

Нужен настоящий FastAPI (фикстура client); сервер Axigen — учебный (tests/axigen_fake.py).
"""
import io
import time

import pytest

from axigen_fake import FakeAxigenCLI, FakeAxigenFTP, FtpWorld, World, message

PW = "Sw0rdfish!"
ADMIN_PW = "Adm1nPass"


def _login(client):
    client.post("/api/setup", json={"username": "admin", "password": PW})
    client.post("/api/login", json={"username": "admin", "password": PW})


@pytest.fixture()
def servers():
    world = World()
    fworld = FtpWorld({"domains": {"company.ru": {
        "ivanov": {"Inbox": {"1.eml": message("Первое")}}, "petrov": {"Inbox": {}}, "postmaster": {"Inbox": {}}}}})
    with FakeAxigenCLI(world) as cli, FakeAxigenFTP(fworld) as ftp:
        yield type("S", (), {"world": world, "fworld": fworld, "cli": cli, "ftp": ftp})


def _configure(client, s, **extra):
    values = {"axigen.enabled": True, "axigen.host": "127.0.0.1", "axigen.cli_port": s.cli.port,
              "axigen.ftp_port": s.ftp.port, "axigen.admin_user": "admin", "axigen.admin_password": ADMIN_PW,
              "axigen.timeout_s": 5, "axigen.imap_host": "mx.company.ru", "axigen.allow_changes": True}
    values.update(extra)
    r = client.put("/api/settings", json={"values": values})
    assert r.status_code == 200, r.text


def _wait_job(client, job_id, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("success", "partial", "failed", "cancelled"):
            return job
        time.sleep(0.3)
    raise AssertionError(f"задание {job_id} не завершилось")


def _sync(client):
    r = client.post("/api/axigen/sync")
    assert r.status_code == 200, r.text
    _wait_job(client, r.json()["job_id"])


def test_section_needs_admin(client):
    assert client.get("/api/server/targets").status_code == 401
    _login(client)
    data = client.get("/api/server/targets").json()
    assert data["rows"] == [] and data["axigen"]["enabled"] is False
    acts = client.get("/api/server/actions").json()
    keys = {a["key"]: a for a in acts["actions"]}
    assert "imap_trash" in keys and keys["ax_block"]["unavailable"]


def test_preview_run_history_and_report(client, servers):
    _login(client)
    _configure(client, servers)
    _sync(client)
    rows = client.get("/api/server/targets").json()["rows"]
    assert {r["key"] for r in rows} >= {"ax:ivanov@company.ru", "ax:petrov@company.ru"}
    body = {"action": "report", "keys": ["ax:ivanov@company.ru", "ax:nobody@company.ru"], "preview": True}
    pv = client.post("/api/server/ops", json=body).json()
    assert pv["counts"] == {"ok": 1, "skip": 0, "fail": 1}
    body["preview"] = False
    res = client.post("/api/server/ops", json=body).json()
    assert res["op_id"] and res["jobs"]
    job = _wait_job(client, res["jobs"][0])
    assert job["status"] == "success", job
    hist = client.get("/api/server/history").json()["items"]
    assert hist[0]["id"] == res["op_id"] and hist[0]["action"] == "report"
    assert all(i["id"] != res["op_id"] for i in client.get("/api/accounts/bulk/history").json()["items"])
    item = client.get(f"/api/server/history/{res['op_id']}").json()
    assert item["job"]["report"] is True and item["job"]["items"]
    r = client.post(f"/api/server/history/{res['op_id']}/report")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    assert client.get("/api/accounts/bulk/history/%d" % res["op_id"]).status_code == 404


def test_passwords_download_once(client, servers):
    _login(client)
    _configure(client, servers)
    _sync(client)
    body = {"action": "ax_password", "keys": ["ax:petrov@company.ru"], "params": {"mode": "random_file"},
            "preview": False}
    res = client.post("/api/server/ops", json=body).json()
    _wait_job(client, res["jobs"][0])
    r = client.post(f"/api/server/history/{res['op_id']}/passwords")
    assert r.status_code == 200
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(2, 1).value == "petrov@company.ru"
    assert ws.cell(2, 3).value == servers.world.domains["company.ru"].accounts["petrov"].password
    assert client.post(f"/api/server/history/{res['op_id']}/passwords").status_code == 400


def test_dangerous_action_needs_confirmation_over_api(client, servers):
    _login(client)
    _configure(client, servers)
    _sync(client)
    body = {"action": "ax_alias_remove", "keys": ["ax:ivanov@company.ru"],
            "params": {"which": "all"}, "preview": True}
    assert client.post("/api/server/ops", json=body).json()["counts"]["ok"] == 1
    # ящик без архива: очищать корзину нечем — пропуск, подтверждать нечего
    body = {"action": "imap_trash", "keys": ["ax:ivanov@company.ru"], "params": {"days": 30}, "preview": False}
    r = client.post("/api/server/ops", json=body)
    assert r.status_code == 200 and r.json()["counts"]["skip"] == 1 and not r.json()["jobs"]
    # ящик архива на другом сервере: опасное действие без числа ящиков не запускается
    aid = client.post("/api/accounts", json={"name": "Ящик", "host": "imap.example.com", "port": 993,
                                             "username": "u@example.com", "password": "secret", "security": "ssl",
                                             "auth_type": "password"}).json()["id"]
    body = {"action": "imap_trash", "keys": [f"mb:{aid}"], "params": {"days": 30}, "preview": False}
    r = client.post("/api/server/ops", json=body)
    assert r.status_code == 400 and "введите число ящиков" in r.text


def test_create_preview_and_export(client, servers):
    _login(client)
    _configure(client, servers)
    _sync(client)
    r = client.post("/api/server/create", json={"text": "new@company.ru; Анна; Смирнова\nivanov@company.ru",
                                                "preview": True}).json()
    assert r["counts"] == {"ok": 1, "skip": 1, "fail": 0}
    x = client.post("/api/server/targets/export", json={})
    assert x.status_code == 200 and x.headers["content-disposition"].startswith("attachment")
