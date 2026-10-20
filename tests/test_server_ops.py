"""Раздел «Ящики на сервере»: действия по IMAP (любой сервер) и с учётными записями Axigen.

IMAP-сервер — учебный, в памяти (tests/imap_fake.py), Axigen — учебные CLI и FTP Backup
(tests/axigen_fake.py). Копирование ящика перед удалением настоящее. Каждое изменение
проверяется по состоянию учебного сервера, а не по ответу клиента.
"""
import io

import pytest

from imap_fake import ImapWorld, raw_message
from mailarchiver import bulk, models
from mailarchiver.errors import ValidationError
from mailarchiver.imap import client as client_mod
from mailarchiver.models import AuthType, JobStatus, JobType
from mailarchiver.queue.jobs import JobContext
from mailarchiver.serverops import base, report, secrets
from mailarchiver.serverops import targets as tg

ADMIN = {"username": "tester", "role": "admin"}


@pytest.fixture()
def imap(services, monkeypatch):
    world = ImapWorld()
    monkeypatch.setattr(client_mod, "IMAPClient", world.client_class())
    return type("Imap", (), {"world": world, "svc": services})


def mailbox(svc, username, *, password="pw", **kw):
    data = dict(name=username, host="imap.test", port=993, username=username, password=password,
                auth_type=AuthType.PASSWORD, retention_days=0)
    data.update(kw)
    return svc.db.create_account(models.Account(**data))


def run_job(svc, action, keys, **options):
    params = {"action": action, "keys": keys, "options": options, "requested_by": "tester", "op_id": None}
    job_id = svc.db.enqueue_job(JobType.SERVER_ACTION, None, params, 5, 1, "tester")
    ctx = JobContext(svc, job_id, JobType.SERVER_ACTION, None, params)
    from mailarchiver.queue.jobs import HANDLERS
    res = HANDLERS[JobType.SERVER_ACTION](ctx)
    return res, [r["message"] for r in svc.db.list_job_events(job_id)]


def item(res):
    assert len(res["items"]) == 1, res
    return res["items"][0]


def preview(svc, action, keys, **params):
    return base.run(svc, ADMIN, action, keys, params, preview=True)


def archived(svc, acc_id, folder):
    return {r["uid"] for r in svc.db.query("SELECT uid FROM messages WHERE account_id=? AND folder=?",
                                           (acc_id, folder))}


# ---------------------------------------------------------------------------
#  Список ящиков
# ---------------------------------------------------------------------------
def test_targets_are_archive_mailboxes_with_imap_access(imap):
    a = mailbox(imap.svc, "u1@test.ru")
    b = mailbox(imap.svc, "u2@test.ru", password="")
    rows = {t.key: t for t in tg.load(imap.svc)}
    assert set(rows) == {f"mb:{a}", f"mb:{b}"}
    assert tg.imap_problem(imap.svc, rows[f"mb:{a}"].mailbox) == ""
    assert "нет сохранённого пароля" in tg.imap_problem(imap.svc, rows[f"mb:{b}"].mailbox)
    data = tg.overview(imap.svc)
    by_key = {r["key"]: r for r in data["rows"]}
    assert by_key[f"mb:{a}"]["mailbox"]["imap"] is True and by_key[f"mb:{a}"]["mailbox"]["keeps_all"] is True
    assert by_key[f"mb:{b}"]["mailbox"]["imap"] is False and data["axigen"]["enabled"] is False


def test_preview_skips_mailbox_without_password_and_unknown_keys(imap):
    a = mailbox(imap.svc, "u1@test.ru", password="")
    res = preview(imap.svc, "imap_seen", [f"mb:{a}", "mb:9999", "garbage"])
    assert res["counts"] == {"ok": 0, "skip": 1, "fail": 1}
    assert "нет сохранённого пароля" in res["results"][0]["detail"]


def test_axigen_actions_are_unavailable_without_axigen(imap):
    a = mailbox(imap.svc, "u1@test.ru")
    with pytest.raises(ValidationError, match="Axigen выключена"):
        preview(imap.svc, "ax_block", [f"mb:{a}"])


# ---------------------------------------------------------------------------
#  Очистка корзины и спама — только при полном архиве
# ---------------------------------------------------------------------------
def _trash_box(imap, **kw):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Trash": [b"\\Trash"], "Spam": [b"\\Junk"]})
    acc = mailbox(imap.svc, "u1@test.ru", **kw)
    w.add("u1@test.ru", "INBOX", raw_message("Входящее"), days_ago=100)
    old = [w.add("u1@test.ru", "Trash", raw_message(f"Старое {i}"), days_ago=40 + i) for i in range(2)]
    new = w.add("u1@test.ru", "Trash", raw_message("Недавнее"), days_ago=1)
    return acc, old, new


def test_trash_cleanup_removes_only_archived_old_messages(imap):
    acc, old, new = _trash_box(imap)
    res, events = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    it = item(res)
    assert it["status"] == "success", it
    assert "удалено писем" in it["detail"] and "2" in it["detail"]
    assert imap.world.uids("u1@test.ru", "Trash") == [new]              # недавнее осталось
    assert len(imap.world.uids("u1@test.ru", "INBOX")) == 1             # другие папки не тронуты
    assert set(old) <= archived(imap.svc, acc, "Trash")                 # удалённое — в архиве
    assert "UID EXPUNGE" in imap.world.commands                         # стёрты только эти письма


def test_trash_cleanup_skips_folder_that_is_not_copied(imap):
    acc, old, new = _trash_box(imap, folder_exclude=["Trash"])
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    it = item(res)
    assert it["status"] == "skipped" and "не копируется в архив" in it["detail"]
    assert imap.world.uids("u1@test.ru", "Trash") == old + [new]


def test_trash_cleanup_skips_folder_when_a_message_is_not_in_archive(imap, monkeypatch):
    acc, old, new = _trash_box(imap)
    from mailarchiver.serverops import imap_ops
    monkeypatch.setattr(imap_ops, "fresh_copy", lambda ctx, acc, emit: (True, ""))   # архив пуст
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    it = item(res)
    assert it["status"] == "skipped" and "2 из 2 писем нет в архиве" in it["detail"]
    assert len(imap.world.uids("u1@test.ru", "Trash")) == 3


def test_lost_archive_file_counts_as_missing(imap):
    import os
    acc, old, new = _trash_box(imap)
    from mailarchiver.serverops import imap_ops
    run_job(imap.svc, "imap_seen", [f"mb:{acc}"], folders="", days=0)        # ничего не удаляет
    from mailarchiver.queue.jobs import _backup_account
    params = {"action": "x", "keys": [], "options": {}}
    job_id = imap.svc.db.enqueue_job(JobType.SERVER_ACTION, None, params, 5, 1, "t")
    jctx = JobContext(imap.svc, job_id, JobType.SERVER_ACTION, None, params)
    _backup_account(jctx, imap.svc.db.get_account(acc))
    row = imap.svc.db.query_one("SELECT uidvalidity, stored_path FROM messages WHERE account_id=? AND folder='Trash' "
                                "AND uid=?", (acc, old[0]))
    os.remove(imap.svc.store.message_path(acc, row["stored_path"]))
    ctx = base.Ctx(imap.svc, "t", {}, preview=True)
    assert imap_ops.missing_in_archive(ctx, imap.svc.db.get_account(acc), "Trash", row["uidvalidity"],
                                       old) == [old[0]]


def test_without_uidplus_other_deleted_messages_block_the_cleanup(imap):
    acc, old, new = _trash_box(imap)
    imap.world.capabilities.discard("UIDPLUS")
    imap.world.boxes["u1@test.ru"]["Trash"].messages[new].flags.add(b"\\Deleted")
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    it = item(res)
    assert it["status"] == "skipped" and "UIDPLUS" in it["detail"]
    assert len(imap.world.uids("u1@test.ru", "Trash")) == 3
    # без чужих отметок «удалено» — обычный EXPUNGE
    imap.world.boxes["u1@test.ru"]["Trash"].messages[new].flags.discard(b"\\Deleted")
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    assert item(res)["status"] == "success"
    assert imap.world.uids("u1@test.ru", "Trash") == [new] and "EXPUNGE" in imap.world.commands


def test_without_uidplus_a_concurrent_deleted_mark_stops_the_expunge(imap, monkeypatch):
    acc, old, new = _trash_box(imap)
    imap.world.capabilities.discard("UIDPLUS")
    real = client_mod.ImapConnection.add_flags

    def racing(self, uids, flags):
        real(self, uids, flags)
        # другой клиент как раз пометил «удалено» письмо, которого нет в архиве
        imap.world.boxes["u1@test.ru"]["Trash"].messages[new].flags.add(b"\\Deleted")
    monkeypatch.setattr(client_mod.ImapConnection, "add_flags", racing)
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    assert item(res)["status"] == "skipped" and "отметка снята" in item(res)["detail"]
    box = imap.world.boxes["u1@test.ru"]["Trash"].messages
    assert len(box) == 3 and not any(b"\\Deleted" in box[u].flags for u in old)


def test_junk_cleanup_finds_folder_by_name(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Спам": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "Спам", raw_message("Реклама"), days_ago=5)
    res, _ = run_job(imap.svc, "imap_junk", [f"mb:{acc}"], days=0, folder="")
    assert item(res)["status"] == "success" and w.uids("u1@test.ru", "Спам") == []


def test_cleanup_reports_missing_trash(imap):
    imap.world.mailbox("u1@test.ru")
    acc = mailbox(imap.svc, "u1@test.ru")
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=0, folder="")
    assert item(res)["status"] == "skipped" and "не найдена" in item(res)["detail"]


def test_purge_old_requires_archive_that_keeps_everything(imap):
    imap.world.mailbox("u1@test.ru")
    acc = mailbox(imap.svc, "u1@test.ru", retention_days=30)
    res = preview(imap.svc, "imap_purge", [f"mb:{acc}"], days=365, folders="")
    assert res["counts"]["skip"] == 1 and "хранит только последние 30 дн." in res["results"][0]["detail"]


def test_purge_old_removes_archived_messages_in_all_folders(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Sent": [b"\\Sent"]})
    acc = mailbox(imap.svc, "u1@test.ru")
    keep = w.add("u1@test.ru", "INBOX", raw_message("Свежее"), days_ago=10)
    w.add("u1@test.ru", "INBOX", raw_message("Старое"), days_ago=400)
    w.add("u1@test.ru", "Sent", raw_message("Старое отправленное"), days_ago=500)
    res, _ = run_job(imap.svc, "imap_purge", [f"mb:{acc}"], days=365, folders="")
    assert item(res)["status"] == "success", item(res)
    assert w.uids("u1@test.ru", "INBOX") == [keep] and w.uids("u1@test.ru", "Sent") == []
    assert len(archived(imap.svc, acc, "INBOX")) == 2


def test_imap_deletion_refuses_ftp_backup_mailboxes(imap):
    imap.world.mailbox("u1@test.ru")
    acc = mailbox(imap.svc, "u1@test.ru", auth_type=AuthType.AXIGEN_FTP)
    res = preview(imap.svc, "imap_trash", [f"mb:{acc}"], days=0)
    assert res["counts"]["skip"] == 1 and "FTP Backup" in res["results"][0]["detail"]


def test_dangerous_action_needs_number_confirmation_and_goes_to_history(imap, monkeypatch):
    imap.world.mailbox("u1@test.ru", folders={"Trash": [b"\\Trash"]})
    acc = mailbox(imap.svc, "u1@test.ru")
    queued = []
    monkeypatch.setattr(imap.svc.queue, "enqueue", lambda *a, **k: queued.append((a, k)) or 77)
    with pytest.raises(ValidationError, match="введите число ящиков"):
        base.run(imap.svc, ADMIN, "imap_trash", [f"mb:{acc}"], {"days": 30}, preview=False)
    res = base.run(imap.svc, ADMIN, "imap_trash", [f"mb:{acc}"], {"days": 30}, preview=False, confirm="1")
    assert res["jobs"] == [77] and queued[0][0][0] == JobType.SERVER_ACTION
    assert queued[0][0][2]["keys"] == [f"mb:{acc}"] and queued[0][0][2]["options"]["days"] == 30
    hist = base.history(imap.svc)
    assert hist[0]["action"] == "imap_trash" and hist[0]["id"] == res["op_id"]
    assert all(not h["action"].startswith("srv_") for h in bulk.history(imap.svc))
    assert bulk.history_item(imap.svc, res["op_id"]) is None
    assert base.history_item(imap.svc, res["op_id"])["job_id"] == 77


# ---------------------------------------------------------------------------
#  Папки и письма
# ---------------------------------------------------------------------------
def test_move_old_messages_into_folder_by_year(imap):
    import datetime as dt
    w = imap.world
    w.mailbox("u1@test.ru")
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "INBOX", raw_message("Свежее"), days_ago=5)
    w.add("u1@test.ru", "INBOX", raw_message("Старое"), days_ago=400)
    w.add("u1@test.ru", "INBOX", raw_message("Очень старое"), days_ago=800)
    res, _ = run_job(imap.svc, "imap_move", [f"mb:{acc}"], folders="INBOX", days=365, target="Архив/{year}")
    assert item(res)["status"] == "success", item(res)
    now = dt.datetime.now(dt.timezone.utc)
    y1 = (now - dt.timedelta(days=400)).year
    y2 = (now - dt.timedelta(days=800)).year
    assert len(w.uids("u1@test.ru", "INBOX")) == 1
    moved = sum(len(w.uids("u1@test.ru", f"Архив/{y}")) for y in {y1, y2})
    assert moved == 2 and f"Архив/{y1}" in w.subscribed["u1@test.ru"]


def test_move_without_move_extension_copies_and_expunges(imap):
    w = imap.world
    w.capabilities.discard("MOVE")
    w.mailbox("u1@test.ru")
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "INBOX", raw_message("Старое"), days_ago=400)
    res, _ = run_job(imap.svc, "imap_move", [f"mb:{acc}"], folders="INBOX", days=365, target="Архив")
    assert item(res)["status"] == "success"
    assert w.uids("u1@test.ru", "INBOX") == [] and len(w.uids("u1@test.ru", "Архив")) == 1
    assert "COPY" in w.commands and "UID EXPUNGE" in w.commands


def test_mark_read(imap):
    w = imap.world
    w.mailbox("u1@test.ru")
    acc = mailbox(imap.svc, "u1@test.ru")
    a = w.add("u1@test.ru", "INBOX", raw_message("Старое"), days_ago=40)
    b = w.add("u1@test.ru", "INBOX", raw_message("Свежее"), days_ago=1)
    res, _ = run_job(imap.svc, "imap_seen", [f"mb:{acc}"], folders="", days=30)
    assert "помечено прочитанными: 1" in item(res)["detail"]
    msgs = w.boxes["u1@test.ru"]["INBOX"].messages
    assert b"\\Seen" in msgs[a].flags and b"\\Seen" not in msgs[b].flags


def test_create_rename_and_delete_folders(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Отчёты": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    res, _ = run_job(imap.svc, "imap_mkdir", [f"mb:{acc}"], folders="Проекты\nПроекты/2026", subscribe=True)
    assert item(res)["status"] == "success"
    assert {"Проекты", "Проекты/2026"} <= set(w.boxes["u1@test.ru"]) and "Проекты/2026" in w.subscribed["u1@test.ru"]
    res, _ = run_job(imap.svc, "imap_mkdir", [f"mb:{acc}"], folders="Проекты", subscribe=True)
    assert item(res)["status"] == "skipped"
    res, _ = run_job(imap.svc, "imap_rename", [f"mb:{acc}"], old="Отчёты", new="Отчёты 2025")
    assert item(res)["status"] == "success" and "Отчёты 2025" in w.boxes["u1@test.ru"]
    with pytest.raises(ValidationError, match="INBOX"):
        preview(imap.svc, "imap_rename", [f"mb:{acc}"], old="INBOX", new="X")
    # удаление: вложенные не трогаем, пустую удаляем, с письмами — только с отметкой и при полном архиве
    w.add("u1@test.ru", "Отчёты 2025", raw_message("Отчёт"), days_ago=3)
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="Проекты\nПроекты/2026\nОтчёты 2025",
                     with_mail=False)
    detail = item(res)["detail"]
    assert "Проекты/2026" not in w.boxes["u1@test.ru"] and "вложенные" in detail and "только пустые" in detail
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="Отчёты 2025", with_mail=True)
    assert item(res)["status"] == "success", item(res)
    assert "Отчёты 2025" not in w.boxes["u1@test.ru"] and archived(imap.svc, acc, "Отчёты 2025")


def test_delete_folder_refuses_special_folders(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Sent": [b"\\Sent"]})
    acc = mailbox(imap.svc, "u1@test.ru")
    with pytest.raises(ValidationError):
        preview(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="INBOX")
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="Sent", with_mail=False)
    assert item(res)["status"] == "skipped" and "особая папка" in item(res)["detail"]


def test_trash_cleanup_accepts_only_a_real_trash_folder(imap):
    """Очистка корзины не требует, чтобы архив хранил всё, — и не должна чистить INBOX или «Проекты/Trash»."""
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Проекты": [], "Проекты/Trash": [], "Trash": []})
    acc = mailbox(imap.svc, "u1@test.ru", retention_days=365)
    w.add("u1@test.ru", "INBOX", raw_message("Старое"), days_ago=60)
    w.add("u1@test.ru", "Проекты/Trash", raw_message("Проект"), days_ago=60)
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="INBOX")
    assert item(res)["status"] == "skipped" and "не похожа" in item(res)["detail"]
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="Проекты/Trash")
    assert item(res)["status"] == "skipped"
    assert len(w.uids("u1@test.ru", "INBOX")) == 1 and len(w.uids("u1@test.ru", "Проекты/Trash")) == 1
    # без отметки \Trash: по имени — только папка в корне
    w.add("u1@test.ru", "Trash", raw_message("Корзина"), days_ago=60)
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    assert item(res)["status"] == "success" and w.uids("u1@test.ru", "Trash") == []
    assert len(w.uids("u1@test.ru", "Проекты/Trash")) == 1


def test_zero_uidvalidity_means_nothing_can_be_matched(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Trash": [b"\\Trash"]})
    w.boxes["u1@test.ru"]["Trash"].uidvalidity = 0
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "Trash", raw_message("x"), days_ago=40)
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    assert item(res)["status"] == "skipped" and "нет в архиве" in item(res)["detail"]
    assert len(w.uids("u1@test.ru", "Trash")) == 1


def test_delete_folder_keeps_mail_that_arrives_after_the_check(imap, monkeypatch):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Old": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "Old", raw_message("a"), days_ago=10)
    real = client_mod.ImapConnection.messages_now

    def late(self, folder):
        w.add("u1@test.ru", folder, raw_message("пришло после сверки"))
        return real(self, folder)
    monkeypatch.setattr(client_mod.ImapConnection, "messages_now", late)
    res, events = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="Old", with_mail=True)
    it = item(res)
    # проверенное письмо стёрто — это сделано и сказано в итоге; папка с новым письмом оставлена
    assert it["status"] == "success" and "появились новые письма" in it["detail"]
    assert "удалено писем (все были в архиве): Old: 1" in it["detail"]
    # проверенное письмо стёрто (оно в архиве), новое — на месте, папка — тоже
    assert "Old" in w.boxes["u1@test.ru"] and len(w.uids("u1@test.ru", "Old")) == 1
    assert any("удалено писем (все были в архиве): 1" in e for e in events)


def test_delete_folder_sees_children_of_wildcard_names_and_special_names(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"100%": [], "100%/child": [], "Sent": [], "Отправленные": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="100%\nSent\nОтправленные", with_mail=False)
    detail = item(res)["detail"]
    assert {"100%", "Sent", "Отправленные"} <= set(w.boxes["u1@test.ru"])
    assert "вложенные" in detail and detail.count("особая папка") == 2


def test_delete_folder_copy_failure_deletes_nothing(imap, monkeypatch):
    from mailarchiver.serverops import imap_ops
    w = imap.world
    w.mailbox("u1@test.ru", folders={"A": [], "B": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "B", raw_message("b"))
    monkeypatch.setattr(imap_ops, "fresh_copy", lambda ctx, acc, emit: (False, "сбой"))
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="A\nB", with_mail=True)
    assert item(res)["status"] == "skipped" and "ничего не удалено" in item(res)["detail"]
    assert {"A", "B"} <= set(w.boxes["u1@test.ru"])


def test_partial_cleanup_is_reported_and_audited_when_a_later_folder_fails(imap, monkeypatch):
    from mailarchiver.errors import ImapProtocolError
    w = imap.world
    w.mailbox("u1@test.ru", folders={"A": [], "B": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "A", raw_message("a"), days_ago=400)
    w.add("u1@test.ru", "B", raw_message("b"), days_ago=400)
    real = client_mod.ImapConnection.expunge_uids
    calls = []

    def flaky(self, uids):
        calls.append(1)
        if len(calls) == 2:
            raise ImapProtocolError("связь оборвалась")
        return real(self, uids)
    monkeypatch.setattr(client_mod.ImapConnection, "expunge_uids", flaky)
    res, events = run_job(imap.svc, "imap_purge", [f"mb:{acc}"], days=365, folders="A\nB")
    it = item(res)
    assert it["status"] == "failed" and "A: 1" in it["detail"] and "затем ошибка" in it["detail"]
    audit = " ".join(r["detail"] for r in imap.svc.db.query("SELECT detail FROM audit"))
    assert "«A»: удалено писем" in audit and any("«A»: удалено писем" in e for e in events)


def test_folders_can_be_given_with_slash_on_dot_servers(imap):
    w = imap.world
    w.delimiter = "."
    w.mailbox("u1@test.ru", folders={"Архив": [], "Архив.2020": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "Архив.2020", raw_message("Старое"), days_ago=400)
    res, _ = run_job(imap.svc, "imap_purge", [f"mb:{acc}"], days=365, folders="Архив/2020")
    assert item(res)["status"] == "success", item(res)
    assert w.uids("u1@test.ru", "Архив.2020") == []


def test_delete_folder_with_dot_delimiter_noselect_and_unknown_status(imap, monkeypatch):
    w = imap.world
    w.delimiter = "."
    w.mailbox("u1@test.ru", folders={"Архив": [b"\\Noselect"], "Архив.2020": [], "Пусто": [], "Контейнер": [b"\\Noselect"]})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "Архив.2020", raw_message("Старое"), days_ago=400)
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="Архив/2020\nАрхив\nКонтейнер", with_mail=True)
    it = item(res)
    assert "Архив.2020" not in w.boxes["u1@test.ru"] and "Контейнер" not in w.boxes["u1@test.ru"]
    assert "Архив" in w.boxes["u1@test.ru"] and "вложенные" in it["detail"]     # порядок: родитель ещё с вложенной
    # сервер не ответил числом на STATUS — «неизвестно», а не «пусто»: папка остаётся
    monkeypatch.setattr(client_mod.ImapConnection, "messages_now", lambda self, folder: None)
    res, _ = run_job(imap.svc, "imap_rmdir", [f"mb:{acc}"], folders="Пусто", with_mail=False)
    assert "Пусто" in w.boxes["u1@test.ru"] and "не сообщил" in item(res)["detail"]


def test_has_children_trusts_haschildren_flag_and_bad_names():
    from mailarchiver.imap.client import FolderInfo
    from mailarchiver.serverops.imap_ops import has_children
    p = FolderInfo(name="P", delimiter="/", flags=["\\HasChildren"], selectable=True)
    assert has_children([p], p, "/")
    q = FolderInfo(name="Q", delimiter="/", flags=[], selectable=True)
    assert not has_children([q], q, "/") and has_children([q], q, "/", ["Q/R&D"])


def test_without_uidplus_aborting_keeps_the_users_own_deleted_marks(imap, monkeypatch):
    acc, old, new = _trash_box(imap)
    imap.world.capabilities.discard("UIDPLUS")
    box = imap.world.boxes["u1@test.ru"]["Trash"].messages
    box[old[0]].flags.add(b"\\Deleted")                         # сотрудник сам пометил «удалено»
    real = client_mod.ImapConnection.add_flags

    def racing(self, uids, flags):
        real(self, uids, flags)
        box[new].flags.add(b"\\Deleted")
    monkeypatch.setattr(client_mod.ImapConnection, "add_flags", racing)
    res, _ = run_job(imap.svc, "imap_trash", [f"mb:{acc}"], days=30, folder="")
    assert item(res)["status"] == "skipped"
    assert b"\\Deleted" in box[old[0]].flags and b"\\Deleted" not in box[old[1]].flags


def test_subscribe_all_folders(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"A": [], "B": []})
    acc = mailbox(imap.svc, "u1@test.ru")
    w.subscribed["u1@test.ru"].add("INBOX")
    res, _ = run_job(imap.svc, "imap_subscribe", [f"mb:{acc}"])
    assert item(res)["status"] == "success" and w.subscribed["u1@test.ru"] == {"INBOX", "A", "B"}


def test_login_check_records_status(imap):
    imap.world.mailbox("u1@test.ru")
    good = mailbox(imap.svc, "u1@test.ru")
    bad = mailbox(imap.svc, "u1@test.ru", password="wrong", name="второй")
    res, _ = run_job(imap.svc, "imap_login", [f"mb:{good}", f"mb:{bad}"])
    st = {i["key"]: i["status"] for i in res["items"]}
    assert st == {f"mb:{good}": "success", f"mb:{bad}": "failed"}
    assert imap.svc.db.get_account(good).login_status == "ok"
    assert imap.svc.db.get_account(bad).login_status == "auth_error"


def test_report_by_imap_and_excel(imap):
    w = imap.world
    w.mailbox("u1@test.ru", folders={"Sent": []})
    w.quota["u1@test.ru"] = (512, 2048)
    acc = mailbox(imap.svc, "u1@test.ru")
    w.add("u1@test.ru", "INBOX", raw_message("1"))
    w.add("u1@test.ru", "INBOX", raw_message("2"), flags=[b"\\Seen"])
    w.add("u1@test.ru", "Sent", raw_message("3"))
    res, _ = run_job(imap.svc, "report", [f"mb:{acc}"])
    rep = item(res)["report"]
    assert rep["messages"] == 3 and rep["unseen"] == 2 and rep["quota"] == 2048 * 1024 and rep["percent"] == 25.0
    assert res["report"] is True
    data, name, _ = report.targets_file(imap.svc)
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.cell(2, 1).value == "u1@test.ru" and name.endswith(".xlsx")
